# -*- coding: utf-8 -*-
"""Minutos de voz por cuenta y por direccion, para el reporte de cuentas AAA.

   POR QUE EXISTE
   --------------
   Daniel Martinez, 1 oct 2026: «me ayudas a sacar los clientes con mas consumo
   telefonico de nuestras cuentas AAA para compartir con Cacique y que nos ayude
   a revisar si no tienen temas de conectividad tipo Grupo 2711 o Finsus?
   Quizas el top 25. Y si fuera posible separarlo por entrante/saliente».

   DE DONDE SALE CADA COSA, Y POR QUE NO DE UN SOLO SITIO
   -----------------------------------------------------
   · El CONSUMO TOTAL sale del archivo de cortes (`data/cortes-facturacion.xlsx`),
     que es la fuente de facturacion y cubre las 210 cuentas AAA. El `% Consumo`
     de ese archivo NO se usa: en los planes por extensiones trae `Minutos
     Incluidos = 1` y publica porcentajes de seis cifras. La base se recalcula
     replicando `lib/plan-minutos.ts`, que es la regla de la casa.

   · La DIRECCION —entrante contra saliente— sale del Analisis de Llamadas. El
     modulo acumula minutos SOLO de las entrantes: `gen-analisis-llamadas.py`
     entra al bloque de destinos nada mas cuando `dire == 'ent'`. Los minutos
     salientes existen en los exports crudos (columna `total_minutes`) y es lo
     que este script recupera leyendolos.

   LO QUE NO SE PUEDE, Y NO SE INVENTA
   -----------------------------------
   Daniel pide ademas separar por SIP, celular y fijo. **Ninguna de las dos
   fuentes trae tipo de trafico.** El `destination_type` de los exports es la
   ATENCION de la llamada —Lost, Lost_by_agent, Self_service, Redirected,
   Voicemail—, y el `destination_data_1` es la etiqueta del destino DENTRO del
   cliente: nombres de agentes, «Ventas», «Bitacora 1».

   Y aunque se intentara deducir del numero marcado, no se puede: desde la
   unificacion a diez digitos de 2019 Mexico no distingue movil de fijo por el
   numero. Los prefijos 044/045 desaparecieron. Clasificar «celular» desde aqui
   seria inventarle un dato a un tercero.

   Lo que si es derivable del numero marcado, y se publica aparte y con ese
   nombre: extension interna, nacional con LADA, 800 y internacional.

   USO
   ---
       python scripts/consumo-aaa-cacique.py
       python scripts/consumo-aaa-cacique.py --top 25

   Escribe su resultado en JSON a D:/Archivos/salida/consumo-aaa-cacique.json
   para que el generador del reporte no vuelva a leer los 174 MB.
"""
import collections
import datetime
import glob
import io
import json
import os
import re
import sys

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', line_buffering=True)
import openpyxl

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CRUDOS = 'D:/Archivos/Análisis de Llamadas Dashboard'
SALIDA = 'D:/Archivos/salida'

MINUTOS_POR_EXTENSION = 1500
MIN_POR_EXT_PLAUSIBLE = 50
BOLSA_MINIMA = 3
RX_EXTENSIONES = re.compile(r'(\d+)\s*(?:extensi[oó]n(?:es)?|ext\b)', re.I)
RX_EXT_ABREVIADO = re.compile(r'^(\d+)\s+\S.*\bIL\b', re.I)
RX_SIN_VOZ = re.compile(r'\bchat\b|\bagentes?\s+cp\b|sin\s+saldo|n[uú]meros?\s+virtuales?|whatsapp', re.I)


def base_minutos(plan, incl):
    """Replica exacta de `baseMinutos` de lib/plan-minutos.ts."""
    nombre = str(plan or '')
    m = RX_EXTENSIONES.search(nombre) or RX_EXT_ABREVIADO.search(nombre)
    ext = int(m.group(1)) if m else None
    if RX_SIN_VOZ.search(nombre) and not ext and incl < BOLSA_MINIMA:
        return None, ext, 'sin_medicion'
    if ext and ext > 0:
        por_ext = incl / ext if incl > 0 else 0
        if por_ext >= MIN_POR_EXT_PLAUSIBLE:
            return incl, ext, 'bolsa'
        return ext * MINUTOS_POR_EXTENSION, ext, 'extensiones'
    if incl >= BOLSA_MINIMA:
        return incl, ext, 'bolsa'
    return None, ext, 'sin_medicion'


def num(v):
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(',', '').strip())
    except (TypeError, ValueError):
        return 0.0


# ── 1. Cortes: consumo facturado por cuenta AAA ──────────────────────────────
def cortes_aaa():
    """El ULTIMO corte de cada CID clasificado AAA, con su base recalculada."""
    ruta = os.path.join(RAIZ, 'data', 'cortes-facturacion.xlsx')
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    cab = list(next(it))
    ix = {c: i for i, c in enumerate(cab) if c}
    ultimo = {}
    for f in it:
        if f[0] is None:
            continue
        if str(f[ix['Clasificación de empresa']]) != 'AAA':
            continue
        cid = str(f[ix['CID']]).strip()
        fecha = str(f[ix['Fecha de corte']])[:10]
        if cid in ultimo and ultimo[cid]['fecha'] >= fecha:
            continue
        incl = num(f[ix['Minutos Incluidos']])
        cons = num(f[ix['Minutos Consumidos']])
        plan = str(f[ix['Nombre del Plan']] or '')
        base, ext, origen = base_minutos(plan, incl)
        ultimo[cid] = {
            'cid': cid,
            'empresa': str(f[ix['Nombre del Cliente']] or '').strip(),
            'fecha': fecha,
            'periodo': str(f[ix['Periodo']] or ''),
            'plan': plan,
            'inclArchivo': incl,
            'consumidos': cons,
            'base': base,
            'extensiones': ext,
            'origen': origen,
            'pct': (100.0 * cons / base) if base else None,
            'excedente': max(0.0, cons - base) if base else None,
            'pctEntArchivo': num(f[ix['% Llamadas entrantes']]) if '% Llamadas entrantes' in ix else None,
            'pctSalArchivo': num(f[ix['% Llamadas salientes']]) if '% Llamadas salientes' in ix else None,
        }
    wb.close()
    return ultimo


# ── 2. Analisis de Llamadas: direccion y minutos entrantes ───────────────────
def llamadas_modulo():
    ruta = os.path.join(RAIZ, 'data', 'analisis-llamadas.json')
    d = json.load(io.open(ruta, encoding='utf-8'))
    out = {}
    for cid, v in d['cuentas'].items():
        ent = v.get('ent') or {}
        sal = v.get('sal') or {}
        dest = ent.get('dest') or []
        out[str(cid).strip()] = {
            'entLlamadas': ent.get('total', 0),
            'salLlamadas': sal.get('total', 0),
            'entMinutos': sum((x.get('min') or 0) for x in dest),
            'entPerdidas': sum((x.get('l') or 0) for x in dest),
            'desde': ent.get('desde') or sal.get('desde'),
            'hasta': ent.get('hasta') or sal.get('hasta'),
        }
    return out, d['meta']


# ── 3. Exports crudos: minutos SALIENTES, que el modulo no acumula ───────────
def minutos_salientes():
    """Suma `total_minutes` de los exports de salientes, por CID.

       Devuelve tambien la clasificacion del numero marcado que SI es derivable.
       Nada de celular/fijo: en Mexico no se distingue por el numero desde 2019.
    """
    archivos = [a for a in glob.glob(os.path.join(CRUDOS, '*.xlsx'))
                if 'saliente' in os.path.basename(a).lower()]
    por_cid = collections.defaultdict(lambda: {'min': 0.0, 'llamadas': 0,
                                               'forma': collections.Counter()})
    for arch in sorted(archivos):
        print('  leyendo %s …' % os.path.basename(arch))
        wb = openpyxl.load_workbook(arch, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        it = ws.iter_rows(values_only=True)
        cab = [str(c or '') for c in next(it)]
        ix = {c: i for i, c in enumerate(cab)}
        n = 0
        for r in it:
            cid_i = ix.get('customer_id')
            if cid_i is None or r[cid_i] is None:
                continue
            cid = str(r[cid_i]).strip()
            e = por_cid[cid]
            e['llamadas'] += 1
            mi = ix.get('total_minutes')
            if mi is not None:
                e['min'] += num(r[mi])
            di = ix.get('destination_number')
            if di is not None:
                e['forma'][forma_del_numero(r[di])] += 1
            n += 1
        wb.close()
        print('    %s filas' % format(n, ','))
    return {k: {'min': round(v['min']), 'llamadas': v['llamadas'],
                'forma': dict(v['forma'])} for k, v in por_cid.items()}


def forma_del_numero(v):
    """Lo UNICO que el numero marcado permite afirmar en Mexico.

       NO devuelve «celular» ni «fijo»: desde la unificacion a diez digitos de
       2019 esa distincion no vive en el numero. Decirla seria inventarla.
    """
    s = re.sub(r'\D', '', str(v or ''))
    if not s:
        return 'sin número'
    if len(s) <= 6:
        return 'extensión interna'
    if s.startswith('00') or len(s) > 12:
        return 'internacional'
    if s.startswith('52') and len(s) == 12:
        return 'nacional (10 dígitos)'
    if len(s) == 10:
        if s.startswith('800') or s.startswith('900'):
            return 'número especial (800/900)'
        return 'nacional (10 dígitos)'
    if len(s) == 11 and s.startswith('1'):
        return 'nacional (10 dígitos)'
    return 'otra forma'


def main():
    top = 25
    if '--top' in sys.argv:
        top = int(sys.argv[sys.argv.index('--top') + 1])

    print('1) Cortes de facturación — cuentas AAA')
    aaa = cortes_aaa()
    print('   %d cuentas AAA con corte\n' % len(aaa))

    print('2) Análisis de Llamadas')
    mod, meta = llamadas_modulo()
    print('   %d cuentas con datos de llamadas; ventana %s a %s\n'
          % (len(mod), meta.get('desde'), meta.get('corte')))

    print('3) Exports crudos de salientes (minutos que el módulo no acumula)')
    sal = minutos_salientes()
    print('   %d CIDs con salientes\n' % len(sal))

    filas = []
    for cid, c in aaa.items():
        m = mod.get(cid) or {}
        s = sal.get(cid) or {}
        filas.append({**c,
                      'entLlamadas': m.get('entLlamadas'),
                      'salLlamadas': m.get('salLlamadas'),
                      'entMinutos': m.get('entMinutos'),
                      'entPerdidas': m.get('entPerdidas'),
                      'salMinutos': s.get('min'),
                      'salLlamadasCrudo': s.get('llamadas'),
                      'formaDestino': s.get('forma'),
                      'ventanaLlamadas': (m.get('desde'), m.get('hasta')) if m else None,
                      'medibleDireccion': bool(m)})

    conBase = [f for f in filas if f['base']]
    conBase.sort(key=lambda f: f['consumidos'], reverse=True)

    os.makedirs(SALIDA, exist_ok=True)
    destino = os.path.join(SALIDA, 'consumo-aaa-cacique.json')
    json.dump({
        'generado': datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=-6))).strftime('%Y-%m-%d %H:%M'),
        'meta': {
            'cuentasAAA': len(aaa),
            'conBaseMedible': len(conBase),
            'sinBaseMedible': len(filas) - len(conBase),
            'conDatosDeLlamadas': sum(1 for f in filas if f['medibleDireccion']),
            'ventanaLlamadas': [meta.get('desde'), meta.get('corte')],
            'corteFacturacion': max(c['fecha'] for c in aaa.values()),
        },
        'filas': conBase,
        'sinBase': [f for f in filas if not f['base']],
    }, io.open(destino, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    print('escrito: %s' % destino)

    print('\nTOP %d AAA POR MINUTOS CONSUMIDOS (corte de facturación)' % top)
    print('%-4s %-34s %10s %10s %7s %9s  %s' % ('CID', 'Empresa', 'Consum.', 'Base', '%', 'Exced.', 'dirección'))
    print('-' * 108)
    for f in conBase[:top]:
        d = ('ent %s / sal %s' % (format(f['entLlamadas'], ','), format(f['salLlamadas'], ','))
             if f['medibleDireccion'] else 'NO MEDIDO')
        print('%-4s %-34s %10s %10s %6.0f%% %9s  %s'
              % (f['cid'][:4], f['empresa'][:34], format(int(f['consumidos']), ','),
                 format(int(f['base']), ','), f['pct'], format(int(f['excedente']), ','), d))
    return 0


if __name__ == '__main__':
    sys.exit(main())
