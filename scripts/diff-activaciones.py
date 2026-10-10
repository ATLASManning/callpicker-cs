# -*- coding: utf-8 -*-
"""El .xlsx de Activaciones nuevo contra el que está en uso, antes de copiarlo.

   POR QUE EXISTE
   --------------
   `data/activaciones.xlsx` NO tiene generador: el archivo ES el dato, y
   actualizarlo es copiarlo encima. Eso quiere decir que un cambio de
   estructura en el origen rompe la pantalla **en silencio** — el módulo busca
   la hoja por nombre y las columnas por encabezado, así que un renombre deja
   la pantalla vacía sin un solo error.

   Las cuatro comprobaciones son las que pide la carga semanal:
     1. La hoja se sigue llamando igual.
     2. Las columnas, mismo nombre y mismo orden.
     3. Diff por `ID`: cuántos entran, cuántos salen, cuáles cambian.
     4. Que la calidad del dato NO empeore — y por eso se miden las anomalías
        en LOS DOS archivos, no sólo en el nuevo. Un número de anomalías sin
        su comparación no dice si vamos mejor o peor.

   Y una quinta que no es de estructura sino de alcance: cuántas filas pintaría
   de verdad el módulo con sus filtros (`ID` y `Cliente` no vacíos, `Año >=
   2020`), y cuántas son `demo`. Un demo no es una activación.

   USO
   ---
       python scripts/diff-activaciones.py
       python scripts/diff-activaciones.py "D:\\Archivos\\otro.xlsx"
"""
import io
import os
import sys
import unicodedata

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', line_buffering=True)


def norm(s):
    """Nombre de cliente comparable: sin acentos, sin espacios, minúsculas.

       Hace falta para no confundir dos cosas muy distintas. «Ambulancias
       Soporte Médico» y «AMBULANCIAS SOPORTE MEDICO» son el MISMO cliente
       escrito dos veces —quitar uno es dedupicar—; «promex» y
       «MULTIPOLIMEROS» son dos clientes que comparten ID por un dedazo, y
       quitar uno es perder una activación. Comparar en minúsculas a secas no
       distingue el primer caso del segundo.
    """
    s = unicodedata.normalize('NFD', str(s or ''))
    s = ''.join(c for c in s if unicodedata.category(c) != 'Mn')
    return ''.join(c for c in s.lower() if c.isalnum())


def plata(f):
    try:
        return float(f.get('1er Pago'))
    except (TypeError, ValueError):
        return 0.0

try:
    from openpyxl import load_workbook
except ImportError:
    print(u'  falta openpyxl: pip install openpyxl')
    sys.exit(1)

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EN_USO = os.path.join(RAIZ, 'data', 'activaciones.xlsx')
NUEVO = sys.argv[1] if len(sys.argv) > 1 else r'D:\Archivos\Tablero de Activaciones 2.0.xlsx'
HOJA = 'Hoja1'          # el módulo la busca por NOMBRE
ANIO_MINIMO = 2020      # el filtro de `app/activaciones/page.tsx`

fallos = []


def exige(ok, que, detalle=''):
    print(u'  %s  %s%s' % (u'ok  ' if ok else u'FALLA', que,
                           (u'   — ' + detalle) if detalle else u''))
    if not ok:
        fallos.append(que)


def lee(ruta):
    """(encabezados, filas como dicts). `data_only` para traer el VALOR de las
       fórmulas y no su texto; sin eso una columna calculada llega como
       '=A1+B1' y el filtro la descarta sin decir por qué."""
    wb = load_workbook(ruta, data_only=True, read_only=True)
    if HOJA not in wb.sheetnames:
        return None, None, wb.sheetnames
    ws = wb[HOJA]
    it = ws.iter_rows(values_only=True)
    cab = [str(c).strip() if c is not None else '' for c in next(it)]
    filas = []
    for r in it:
        if all(v is None or str(v).strip() == '' for v in r):
            continue
        filas.append({cab[i]: r[i] for i in range(min(len(cab), len(r)))})
    wb.close()
    return cab, filas, wb.sheetnames


print(u'  EN USO  %s' % EN_USO)
print(u'  NUEVO   %s\n' % NUEVO)
for p in (EN_USO, NUEVO):
    if not os.path.exists(p):
        print(u'  no existe: %s' % p)
        sys.exit(1)

cab_v, filas_v, hojas_v = lee(EN_USO)
cab_n, filas_n, hojas_n = lee(NUEVO)

# ── 1 · la hoja ──────────────────────────────────────────────────────────
exige(cab_n is not None, u'el archivo nuevo tiene la hoja «%s»' % HOJA,
      u'' if cab_n is not None else u'trae %s' % hojas_n)
if cab_n is None:
    sys.exit(1)
exige(cab_v is not None, u'el archivo en uso tiene la hoja «%s»' % HOJA)
if cab_v is None:
    sys.exit(1)

# ── 2 · las columnas ─────────────────────────────────────────────────────
print()
exige(cab_v == cab_n, u'las columnas son las mismas y en el mismo orden',
      u'%d en uso · %d nuevas' % (len(cab_v), len(cab_n)))
if cab_v != cab_n:
    faltan = [c for c in cab_v if c not in cab_n]
    sobran = [c for c in cab_n if c not in cab_v]
    if faltan:
        print(u'        DESAPARECEN: %s' % ', '.join(faltan))
    if sobran:
        print(u'        APARECEN:    %s' % ', '.join(sobran))
    movidas = [(i, a, b) for i, (a, b) in enumerate(zip(cab_v, cab_n)) if a != b]
    for i, a, b in movidas[:8]:
        print(u'        posicion %2d: «%s» -> «%s»' % (i, a, b))

print(u'\n  filas: %d en uso · %d nuevas  (%+d)'
      % (len(filas_v), len(filas_n), len(filas_n) - len(filas_v)))


# ── 3 · el diff por ID ───────────────────────────────────────────────────
def clave(f):
    v = f.get('ID')
    return '' if v is None else str(v).strip()


idx_v = {}
idx_n = {}
for f in filas_v:
    k = clave(f)
    if k:
        idx_v.setdefault(k, []).append(f)
for f in filas_n:
    k = clave(f)
    if k:
        idx_n.setdefault(k, []).append(f)

nuevos = sorted(set(idx_n) - set(idx_v))
idos = sorted(set(idx_v) - set(idx_n))
print(u'\n  ── DIFF POR ID ───────────────────────────────────────────────')
print(u'      %d IDs nuevos · %d que desaparecen' % (len(nuevos), len(idos)))
for k in nuevos[:12]:
    f = idx_n[k][0]
    print(u'        + %-10s %-30s %s' % (k, str(f.get('Cliente') or '')[:30],
                                         f.get('Mes 1er Pago') or ''))
if len(nuevos) > 12:
    print(u'        + … y %d más' % (len(nuevos) - 12))
for k in idos:
    f = idx_v[k][0]
    print(u'        - %-10s %-30s' % (k, str(f.get('Cliente') or '')[:30]))

# ── UN ID NUEVO NO ES NECESARIAMENTE NEGOCIO ────────────────────────────
#
# El corte del 9 oct 2026 trajo 29 IDs nuevos con sólo +6 filas netas, y leer
# eso como «29 activaciones» habría sido falso: 5 eran clientes que YA estaban
# y cambiaron de número —3 de ellos porque compartían ID con otro—. La
# diferencia entre negocio y limpieza no se ve en el conteo.
norms_v = {}
for f in filas_v:
    norms_v.setdefault(norm(f.get('Cliente')), []).append(clave(f))
limpieza, negocio = [], []
for kk in nuevos:
    f = idx_n[kk][0]
    previos = norms_v.get(norm(f.get('Cliente')), [])
    (limpieza if previos else negocio).append((kk, f, previos))

print(u'\n      de los %d IDs nuevos:' % len(nuevos))
print(u'        %d son clientes que YA estaban, renumerados (limpieza)' % len(limpieza))
for kk, f, previos in limpieza:
    print(u'          %-10s %-28s antes %s'
          % (kk, str(f.get('Cliente') or '')[:28], ', '.join(previos[:3])))
ndemo = [x for x in negocio if str(x[1].get('Tipo') or '').strip().lower() != 'demo']
print(u'        %d son clientes nuevos — %d de ellos no-demo, $%s de 1er pago'
      % (len(negocio), len(ndemo),
         format(int(sum(plata(f) for _, f, _ in ndemo)), ',d')))
for kk, f, _ in ndemo:
    print(u'          %-10s %-28s %-9s $%-8s %s'
          % (kk, str(f.get('Cliente') or '')[:28], str(f.get('Tipo') or ''),
             format(int(plata(f)), ',d'), f.get('Mes 1er Pago') or 'sin mes'))

# UN ID QUE DESAPARECE TAMPOCO ES NEUTRO. Puede ser un dedazo corregido —ya
# pasó con `189867]`, con un corchete pegado— o un registro perdido.
if idos:
    print(u'\n      ¿alguno de los idos sobrevive con otro ID?')
    norms_n = {}
    for f in filas_n:
        norms_n.setdefault(norm(f.get('Cliente')), []).append(clave(f))
    for k in idos:
        f = idx_v[k][0]
        gemelos = norms_n.get(norm(f.get('Cliente')), [])
        print(u'        %-10s %-26s %-9s $%-8s -> %s'
              % (k, str(f.get('Cliente') or '')[:26], str(f.get('Tipo') or ''),
                 format(int(plata(f)), ',d'),
                 ', '.join(gemelos) if gemelos else u'NO reaparece'))

# ── LOS DUPLICADOS RESUELTOS, UNO POR UNO ───────────────────────────────
#
# Que `ids_dup` baje a 0 se lee como una buena noticia, y a veces lo es. Pero
# un par puede ser el mismo cliente dos veces (quitar uno dedupica) o DOS
# clientes compartiendo ID (quitar uno borra una activación). El conteo no las
# distingue; el nombre normalizado, sí.
dup_v = sorted(kk for kk, lst in idx_v.items() if len(lst) > 1)
if dup_v:
    print(u'\n  ── LOS %d IDs DUPLICADOS DEL ARCHIVO EN USO ───────────────' % len(dup_v))
    mismos, distintos = [], []
    for kk in dup_v:
        a, b = idx_v[kk][0], idx_v[kk][1]
        (mismos if norm(a.get('Cliente')) == norm(b.get('Cliente'))
         else distintos).append((kk, a, b))
    print(u'      %d son el MISMO cliente escrito distinto; %d son dos clientes'
          % (len(mismos), len(distintos)))

    pierde_pago = []
    for kk, a, b in mismos:
        viejo = max(plata(a), plata(b))
        nuevo = max((plata(x) for x in idx_n.get(kk, [])), default=0.0)
        if nuevo < viejo - 0.5:
            pierde_pago.append((kk, str(a.get('Cliente') or ''), viejo, nuevo))
    exige(not pierde_pago, u'ningún duplicado del mismo cliente pierde su primer pago',
          u'' if not pierde_pago
          else u', '.join('%s $%d->$%d' % (x[0], x[2], x[3]) for x in pierde_pago))

    huerfanas = []
    for kk, a, b in distintos:
        for f in (a, b):
            if norm(f.get('Cliente')) not in {norm(x.get('Cliente')) for x in filas_n}:
                huerfanas.append((kk, f))
    if huerfanas:
        nd = [(kk, f) for kk, f in huerfanas
              if str(f.get('Tipo') or '').strip().lower() != 'demo']
        print(u'\n      %d cliente(s) de esos pares DESAPARECEN, %d no-demo:'
              % (len(huerfanas), len(nd)))
        for kk, f in huerfanas:
            # ¿El superviviente trae el MISMO importe y el mismo mes? Si sí, era
            # la misma activación con dos nombres y no se perdió nada.
            sup = idx_n.get(kk, [])
            igual = any(abs(plata(x) - plata(f)) < 0.5
                        and str(x.get('Mes 1er Pago') or '') == str(f.get('Mes 1er Pago') or '')
                        for x in sup)
            print(u'        %-10s %-26s %-11s $%-8s %-9s  %s'
                  % (kk, str(f.get('Cliente') or '')[:26], str(f.get('Tipo') or ''),
                     format(int(plata(f)), ',d'), f.get('Mes 1er Pago') or 'sin mes',
                     u'el que queda trae el mismo importe y mes: era la misma'
                     if igual else u'*** REVISAR: importe o mes distintos'))

# Los que cambian de contenido, campo por campo.
cambios = []
for k in sorted(set(idx_v) & set(idx_n)):
    a, b = idx_v[k][0], idx_n[k][0]
    difs = [c for c in cab_n
            if str(a.get(c) if a.get(c) is not None else '')
            != str(b.get(c) if b.get(c) is not None else '')]
    if difs:
        cambios.append((k, str(b.get('Cliente') or '')[:26], difs))
print(u'\n      %d ID(s) con algún campo distinto' % len(cambios))
porCampo = {}
for _, _, difs in cambios:
    for c in difs:
        porCampo[c] = porCampo.get(c, 0) + 1
for c, n in sorted(porCampo.items(), key=lambda kv: -kv[1])[:10]:
    print(u'        %-28s %4d' % (c[:28], n))
for k, nom, difs in cambios[:6]:
    print(u'        %-10s %-26s %s' % (k, nom, ', '.join(difs)[:60]))


# ── 4 · la calidad, en LOS DOS ───────────────────────────────────────────
def anomalias(filas, cab):
    a = {}
    a['filas'] = len(filas)
    a['anio_1899'] = sum(1 for f in filas if str(f.get('Año') or '').strip() == '1899')
    a['sin_id'] = sum(1 for f in filas if not clave(f))
    a['sin_cliente'] = sum(1 for f in filas
                           if not str(f.get('Cliente') or '').strip())
    ids = [clave(f) for f in filas if clave(f)]
    a['ids_dup'] = len(ids) - len(set(ids))
    a['nombre_sucio'] = sum(
        1 for f in filas
        if str(f.get('Cliente') or '') != str(f.get('Cliente') or '').strip())
    col = 'Fecha 1er Pago'
    if col in cab:
        a['pago_texto'] = sum(1 for f in filas
                              if isinstance(f.get(col), str) and f.get(col).strip())
    return a


av, an = anomalias(filas_v, cab_v), anomalias(filas_n, cab_n)
print(u'\n  ── CALIDAD DEL DATO, en los DOS archivos ─────────────────────')
print(u'      %-26s %8s %8s %8s' % (u'anomalía', u'en uso', u'nuevo', u'delta'))
peor = []
for k in ('anio_1899', 'sin_id', 'sin_cliente', 'ids_dup', 'nombre_sucio', 'pago_texto'):
    if k not in av and k not in an:
        continue
    x, y = av.get(k, 0), an.get(k, 0)
    d = y - x
    if d > 0:
        peor.append(u'%s (+%d)' % (k, d))
    print(u'      %-26s %8d %8d %+8d' % (k, x, y, d))
exige(not peor, u'la calidad del dato no empeora',
      u'' if not peor else u'empeora en: ' + ', '.join(peor))


# ── 5 · lo que el módulo PINTARÍA ────────────────────────────────────────
def pintables(filas):
    out = []
    for f in filas:
        if not clave(f):
            continue
        if not str(f.get('Cliente') or '').strip():
            continue
        try:
            anio = int(str(f.get('Año') or '0').strip() or 0)
        except ValueError:
            continue
        if anio < ANIO_MINIMO:
            continue
        out.append(f)
    return out


pv, pn = pintables(filas_v), pintables(filas_n)


def demos(filas):
    return [f for f in filas if str(f.get('Tipo') or '').strip().lower() == 'demo']


print(u'\n  ── LO QUE SE PINTARÍA (ID y Cliente, y Año >= %d) ────────────' % ANIO_MINIMO)
print(u'      %-26s %8s %8s' % (u'', u'en uso', u'nuevo'))
print(u'      %-26s %8d %8d' % (u'filas que pasan el filtro', len(pv), len(pn)))
print(u'      %-26s %8d %8d' % (u'de ellas, demo', len(demos(pv)), len(demos(pn))))
print(u'      %-26s %8d %8d' % (u'activaciones reales',
                                len(pv) - len(demos(pv)), len(pn) - len(demos(pn))))
descartadas = len(filas_n) - len(pn)
print(u'\n      %d fila(s) del archivo nuevo NO se pintan.' % descartadas)
if descartadas:
    print(u'      Por qué, una por una — un descarte silencioso es una')
    print(u'      activación real que nadie ve (ya pasó con GAMP, ID 189924):')
    for f in filas_n:
        if f in pn:
            continue
        por = []
        if not clave(f):
            por.append('sin ID')
        if not str(f.get('Cliente') or '').strip():
            por.append('sin Cliente')
        try:
            if int(str(f.get('Año') or '0').strip() or 0) < ANIO_MINIMO:
                por.append('Año=%s' % (f.get('Año')))
        except ValueError:
            por.append('Año no numérico (%s)' % (f.get('Año'),))
        print(u'        %-10s %-28s %-10s %s'
              % (clave(f) or '—', str(f.get('Cliente') or '')[:28],
                 str(f.get('Tipo') or ''), ', '.join(por)))

print()
if fallos:
    print(u'  *** %d comprobación(es) fallan. NO copiar todavía.' % len(fallos))
    sys.exit(1)
print(u'  estructura y calidad en orden: se puede copiar encima de data/activaciones.xlsx')
