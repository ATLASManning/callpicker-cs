"""Genera app/cuentas/llamadas-data.ts desde TODOS los archivos de llamadas.

   Dos entregas con criterio distinto, que NO se solapan (verificado: 0 CIDs en
   común) y que juntas cubren 148 de las 219 cuentas con asesor y CID:

     · corte «0-40»  — clientes con consumo de 0 a 40% de su plan.  86 CIDs.
     · corte «40+»   — el resto de la cartera medida.               62 CIDs.

   ── LAS COLUMNAS NO SON IGUALES ENTRE ARCHIVOS ─────────────────────────────
   Por eso todo se resuelve POR NOMBRE DE COLUMNA y nunca por posición:
     · la empresa viene como `empresa`, `Nombre Empresa` o `Nombre de Empresa`
     · «Entrantes Mayor consumo 40 Parte 1» NO trae `destination_data_1`: trae
       `src`. Son 799,999 filas del 20 may al 14 sep —el 45% de las entrantes
       del corte 40+— sin destino. Esas filas van a un bucket propio,
       «(destino no venía en el archivo)», y NO al de «(sin destino
       registrado)»: una cosa es que la llamada no llegara a ninguna extensión
       y otra que la columna no se haya exportado. Confundirlas inventaría un
       hallazgo de configuración que nadie midió.

   ── LO QUE NO SE USA ───────────────────────────────────────────────────────
   «Salientes Mayor consumo 40 Parte 2» es copia byte a byte de la Parte 1
   (mismas 900,000 filas, mismos 54 CIDs, mismo hash, filas idénticas a toda
   profundidad; ambas traen la hoja `xaa`, que es lo que deja `split`). Se lee
   UNA sola vez: sumarlas duplicaría cada llamada saliente de 54 cuentas.

   ── VENTANAS DISTINTAS POR FUENTE ──────────────────────────────────────────
   Las entrantes del corte 40+ cubren del 1 ene al 14 sep, pero sus salientes
   arrancan el 9 de abril: al archivo le falta el tramo anterior. Cada cuenta
   guarda la ventana real de cada dirección para que la pantalla pueda decirlo
   en vez de dar por hecho que ausencia es cero.
"""
import sys, io, os, re, json, unicodedata, datetime, collections
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', line_buffering=True)
import openpyxl
# Los nombres llegan con la codificación rota desde el origen («Bit√°cora»).
# El mismo helper lo usa gen-analisis-llamadas.py: los dos paneles tienen que
# decir lo mismo del mismo destino.
from _texto_roto import arregla, reporte as reporteTexto

ARCH = r"D:\Archivos"
SALIDA = r"D:\Windows\Projects\callpicker-cs\app\cuentas\llamadas-data.ts"

# (archivo, direccion, corte)
FUENTES = [
    ('Llamadas_entrantes_Clientes_AAA_Poco_Consumo_Actualizado.xlsx', 'ent', '0-40'),
    ('Llamadas Salientes Clientes AAA Poco Consumo.xlsx',             'sal', '0-40'),
    ('Entrantes Mayor consumo 40 Parte 1.xlsx',                       'ent', '40+'),
    ('Entrantes Mayor consumo 40 Parte 2.xlsx',                       'ent', '40+'),
    ('Entrantes Mayor consumo 40 Parte 3.xlsx',                       'ent', '40+'),
    ('Salientes Mayor consumo 40 Parte 1.xlsx',                       'sal', '40+'),
    # 'Salientes Mayor consumo 40 Parte 2.xlsx' — copia exacta de la Parte 1.
]

SIN_DESTINO = '(sin destino registrado)'
NO_EXPORTADO = '(destino no venía en el archivo)'
TOP_DESTINOS = 8
RAZON = r'\b(s\.?a\.?p\.?i\.?|s\.?a\.?|s\.?\s?de\s?r\.?l\.?|c\.?v\.?|de\s?c\.?v\.?|sc|sofom|e\.?n\.?r\.?|spr|rl)\b'

ALIAS = {
    'tipo':    ['destination_type'],
    'fecha':   ['date'],
    'cid':     ['customer_id'],
    'empresa': ['empresa', 'Nombre Empresa', 'Nombre de Empresa'],
    'destino': ['destination_data_1'],
    'caller':  ['caller_id'],
    'minutos': ['total_minutes'],
}


def norma(s):
    s = unicodedata.normalize('NFD', str(s or ''))
    s = ''.join(c for c in s if unicodedata.category(c) != 'Mn').lower()
    s = s.replace('&', ' y ')
    s = re.sub(r'[.,()\-_/]', ' ', s)
    s = re.sub(RAZON, ' ', s)
    s = re.sub(r'[^a-z0-9 ]', ' ', s)
    return re.sub(r'\s+', ' ', s).strip()


def cid_de(v):
    if v is None:
        return ''
    if isinstance(v, float) and float(v).is_integer():
        return str(int(v))
    s = str(v).strip()
    return '' if s.upper() == 'NULL' else s


def mins(v):
    if v is None:
        return 0.0
    s = str(v).strip()
    if s == '' or s.upper() == 'NULL':
        return 0.0
    try:
        return float(s)
    except Exception:
        return 0.0


def texto(v):
    s = str(v or '').strip()
    return '' if s.upper() == 'NULL' else s


def nueva():
    return {
        'empresa': collections.Counter(), 'cortes': set(),
        'ent': {'total': 0, 'lost': 0, 'lostSinNum': 0,
                'meses': collections.defaultdict(collections.Counter),
                'dow': collections.Counter(), 'dowL': collections.Counter(),
                'hora': collections.Counter(), 'horaL': collections.Counter(),
                'dest': collections.defaultdict(lambda: {'l': 0, 'c': 0, 'm': 0.0, 'n': set()}),
                'desde': None, 'hasta': None, 'sinCol': 0},
        'sal': {'total': 0, 'noCon': 0, 'meses': collections.defaultdict(collections.Counter),
                'desde': None, 'hasta': None},
    }


DATOS = collections.defaultdict(nueva)

for arch, dire, corte in FUENTES:
    ruta = os.path.join(ARCH, arch)
    if not os.path.exists(ruta):
        print('!! FALTA: %s' % arch)
        continue
    wb = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    cab = [str(c).strip() if c is not None else '' for c in next(it)]
    ix = {}
    for k, nombres in ALIAS.items():
        ix[k] = next((cab.index(n) for n in nombres if n in cab), None)
    falta = [k for k in ('tipo', 'fecha', 'cid') if ix[k] is None]
    if falta:
        raise SystemExit('%s: faltan columnas obligatorias %s' % (arch, falta))
    print('%-46s %s · %s · destino=%s' % (arch[:46], dire, corte,
                                          'si' if ix['destino'] is not None else 'NO'))
    n = 0
    for r in it:
        n += 1
        if n % 400000 == 0:
            print('    ... %s' % format(n, ','))
        cid = cid_de(r[ix['cid']])
        if not cid:
            continue
        v = DATOS[cid]
        v['cortes'].add(corte)
        if ix['empresa'] is not None:
            e = arregla(texto(r[ix['empresa']]))
            if e:
                v['empresa'][e] += 1
        t = texto(r[ix['tipo']]) or 'Desconocido'
        f = r[ix['fecha']]
        iso = f.strftime('%Y-%m-%d') if isinstance(f, datetime.datetime) else str(f)[:10]
        mes = iso[:7] if len(iso) >= 7 else ''

        if dire == 'ent':
            d = v['ent']
            d['total'] += 1
            if ix['destino'] is not None:
                # Se repara ANTES de usarlo como llave, para que «Bit√°cora 1»
                # y «Bitácora 1» dejen de ser dos destinos distintos.
                dest = arregla(texto(r[ix['destino']])) or SIN_DESTINO
            else:
                dest = NO_EXPORTADO
                d['sinCol'] += 1
            dd = d['dest'][dest]
            dd['m'] += mins(r[ix['minutos']]) if ix['minutos'] is not None else 0.0
            if t == 'Lost':
                d['lost'] += 1
                dd['l'] += 1
                num = texto(r[ix['caller']]) if ix['caller'] is not None else ''
                if num:
                    dd['n'].add(num)
                else:
                    d['lostSinNum'] += 1
            else:
                dd['c'] += 1
            if mes:
                d['meses'][mes][t] += 1
                d['meses'][mes]['total'] += 1
            if isinstance(f, datetime.datetime):
                d['dow'][f.weekday()] += 1
                d['hora'][f.hour] += 1
                if t == 'Lost':
                    d['dowL'][f.weekday()] += 1
                    d['horaL'][f.hour] += 1
            if iso and iso != 'None':
                if d['desde'] is None or iso < d['desde']:
                    d['desde'] = iso
                if d['hasta'] is None or iso > d['hasta']:
                    d['hasta'] = iso
        else:
            d = v['sal']
            d['total'] += 1
            # Lost_by_agent se suma a «no conecto»: el archivo no define que la
            # distingue de Lost y no se le inventa un significado propio.
            noCon = t in ('Lost', 'Lost_by_agent')
            if noCon:
                d['noCon'] += 1
            if mes:
                d['meses'][mes]['total'] += 1
                if noCon:
                    d['meses'][mes]['noCon'] += 1
            if iso and iso != 'None':
                if d['desde'] is None or iso < d['desde']:
                    d['desde'] = iso
                if d['hasta'] is None or iso > d['hasta']:
                    d['hasta'] = iso
    wb.close()
    print('    %s filas' % format(n, ','))

# ── Armado ──────────────────────────────────────────────────────────────────
todas = [x for v in DATOS.values() for x in (v['ent']['hasta'], v['sal']['hasta']) if x]
CORTE = max(todas)
MESES = sorted({m for v in DATOS.values() for m in list(v['ent']['meses']) + list(v['sal']['meses']) if m})

MES_CERRADO = None
for m in reversed(MESES):
    y, mm = int(m[:4]), int(m[5:7])
    ultimoDia = (datetime.date(y + (mm == 12), (mm % 12) + 1, 1) - datetime.timedelta(days=1)).isoformat()
    if CORTE >= ultimoDia:
        MES_CERRADO = m
        break
BASE_FIN = MESES[MESES.index(MES_CERRADO) - 1] if MES_CERRADO and MESES.index(MES_CERRADO) > 0 else None
print()
print('corte del archivo : %s' % CORTE)
print('meses             : %s → %s' % (MESES[0], MESES[-1]))
print('ultimo mes CERRADO: %s · base hasta: %s' % (MES_CERRADO, BASE_FIN))

salida = {}
for cid in sorted(DATOS):
    v = DATOS[cid]
    e, s = v['ent'], v['sal']
    emp = v['empresa'].most_common(1)[0][0] if v['empresa'] else ''

    dest = []
    if e['total']:
        ordenados = sorted(e['dest'].items(), key=lambda x: -x[1]['l'])
        for especial in (SIN_DESTINO, NO_EXPORTADO):
            x = e['dest'].get(especial)
            if x and x['l'] > 0:
                dest.append({'d': especial, 'l': x['l'], 'c': x['c'], 'min': round(x['m']), 'n': len(x['n'])})
        conNombre = [(d, x) for d, x in ordenados if d not in (SIN_DESTINO, NO_EXPORTADO) and x['l'] > 0]
        for d, x in conNombre[:TOP_DESTINOS]:
            dest.append({'d': d, 'l': x['l'], 'c': x['c'], 'min': round(x['m']), 'n': len(x['n'])})
        resto = conNombre[TOP_DESTINOS:]
        if resto:
            dest.append({'d': 'otros destinos', 'l': sum(x['l'] for _, x in resto),
                         'c': sum(x['c'] for _, x in resto),
                         'min': round(sum(x['m'] for _, x in resto)), 'n': -1, 'otros': len(resto)})
        assert sum(x['l'] for x in dest) == e['lost'], 'destinos no cierran en CID %s' % cid

    me = {m: {'t': mv['total'], 'l': mv.get('Lost', 0), 'r': mv.get('Redirected', 0),
              's': mv.get('Self_service', 0), 'v': mv.get('Voicemail', 0)}
          for m, mv in sorted(e['meses'].items())}
    ms = {m: {'t': mv['total'], 'n': mv.get('noCon', 0)} for m, mv in sorted(s['meses'].items())}

    cerrado = me.get(MES_CERRADO) if MES_CERRADO else None
    prevs = [x for m, x in me.items() if BASE_FIN and m <= BASE_FIN]

    salida[cid] = {
        'cid': cid, 'empresa': emp, 'norm': norma(emp),
        'corte': sorted(v['cortes'])[0] if v['cortes'] else '',
        'ent': ({'total': e['total'], 'lost': e['lost'], 'meses': me,
                 'dow':  [e['dow'].get(i, 0) for i in range(7)],
                 'dowL': [e['dowL'].get(i, 0) for i in range(7)],
                 'hora':  [e['hora'].get(h, 0) for h in range(24)],
                 'horaL': [e['horaL'].get(h, 0) for h in range(24)],
                 'dest': dest, 'primera': e['desde'], 'ultima': e['hasta'],
                 'lostSinNum': e['lostSinNum'], 'sinCol': e['sinCol']} if e['total'] else None),
        'sal': ({'total': s['total'], 'noCon': s['noCon'], 'meses': ms,
                 'desde': s['desde'], 'ultima': s['hasta']} if s['total'] else None),
        'cerrado': ({'t': cerrado['t'], 'l': cerrado['l']} if cerrado else None),
        'base': {'t': sum(x['t'] for x in prevs), 'l': sum(x['l'] for x in prevs)},
        'ultima': max([x for x in (e['hasta'], s['hasta']) if x], default=None),
    }

META = {
    'corte': CORTE, 'mesCerrado': MES_CERRADO, 'baseFin': BASE_FIN, 'meses': MESES,
    'cuentas': len(salida),
    'fuente': 'Llamadas entrantes y salientes de Callpicker — dos entregas: clientes con consumo de 0 a 40% de su plan y el resto de la cartera medida',
    'entTotal': sum(v['ent']['total'] for v in salida.values() if v['ent']),
    'entLost':  sum(v['ent']['lost']  for v in salida.values() if v['ent']),
    'salTotal': sum(v['sal']['total'] for v in salida.values() if v['sal']),
    'salNoCon': sum(v['sal']['noCon'] for v in salida.values() if v['sal']),
}
META['baseEnt'] = round(100 * META['entLost'] / max(META['entTotal'], 1), 1)
META['baseSalCon'] = round(100 * (META['salTotal'] - META['salNoCon']) / max(META['salTotal'], 1), 1)

os.makedirs(os.path.dirname(SALIDA), exist_ok=True)
with io.open(SALIDA, 'w', encoding='utf-8', newline='\n') as f:
    f.write('/* GENERADO por scripts/gen-llamadas-data.py — NO editar a mano.\n')
    f.write(' * %s\n' % META['fuente'])
    f.write(' * Corte: %s · ultimo mes cerrado: %s · %d cuentas.\n' % (CORTE, MES_CERRADO, len(salida)))
    f.write(' * %s entrantes / %s salientes.\n'
            % (format(META['entTotal'], ','), format(META['salTotal'], ',')))
    f.write(' */\n')
    f.write("import type { LlamadasCuenta, LlamadasMeta } from '@/lib/llamadas-cuenta'\n\n")
    f.write('export const LLAMADAS_META: LlamadasMeta = %s\n\n' % json.dumps(META, ensure_ascii=False))
    f.write('export const LLAMADAS: Record<string, LlamadasCuenta> = {\n')
    for cid, v in salida.items():
        f.write('  %s: %s,\n' % (json.dumps(cid), json.dumps(v, ensure_ascii=False, separators=(',', ':'))))
    f.write('}\n')

print()
print('escrito %s (%.0f KB)' % (SALIDA, os.path.getsize(SALIDA) / 1024))
print('  cuentas   : %d' % len(salida))
print('  entrantes : %s · no contestadas %s (%.1f%%)'
      % (format(META['entTotal'], ','), format(META['entLost'], ','), META['baseEnt']))
print('  salientes : %s · conectaron %.1f%%' % (format(META['salTotal'], ','), META['baseSalCon']))
sc = sum(v['ent']['sinCol'] for v in salida.values() if v['ent'])
print('  filas entrantes sin columna de destino: %s (%.1f%%)' % (format(sc, ','), 100 * sc / max(META['entTotal'], 1)))
por = collections.Counter(v['corte'] for v in salida.values())
print('  por corte de consumo: %s' % dict(por))
reporteTexto('nombres con codificación reparada')
