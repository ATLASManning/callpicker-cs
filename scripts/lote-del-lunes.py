# -*- coding: utf-8 -*-
"""Que actividades de foco se disparan el proximo lunes, cuenta por cuenta.

   POR QUE EXISTE
   --------------
   Jose Manuel, 25 sep 2026: «enlista cuales seran el tipo de actividades que
   vas a desatar a partir del lunes 28».

   `simula-focos.py` responde la pregunta agregada —cuantas por clase, si la
   rotacion se agota—. Esta responde la concreta: QUE CUENTA le toca a QUIEN,
   por que motivo y con que trabajo encima. Es la lista que se puede leer en una
   junta y contrastar contra la realidad.

   Replica la misma regla que `lib/focos-riesgo.ts`. Si las dos difieren, la
   que manda es el TypeScript: esto es un espejo para poder verlo sin Node, no
   una segunda fuente de verdad.

   Es de SOLO LECTURA.

   USO
   ---
       python scripts/lote-del-lunes.py
       python scripts/lote-del-lunes.py 2026-10-05    # otro lunes
"""
import glob
import io
import json
import os
import re
import sys
import unicodedata
import urllib.request
from collections import Counter, defaultdict
from datetime import datetime, timedelta

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', line_buffering=True)
import openpyxl

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LUNES = sys.argv[1] if len(sys.argv) > 1 else '2026-09-28'
HOY = datetime.strptime(LUNES, '%Y-%m-%d')

# Mismas constantes que el TS.
POR_SEMANA = 10
DIAS_LIMITE = 60
UMBRAL_USO = 40
TRABAJOS = ['relacion', 'decisores', 'tickets', 'llamadas', 'factura', 'datos', 'crecimiento']
ORDEN = ['nunca_tocada', 'sin_contacto', 'uso_bajo', 'auditoria', 'soporte', 'sin_corte', 'rotacion']
MOTIVO = {
    'nunca_tocada': 'nunca ha tenido un seguimiento registrado',
    'sin_contacto': 'pasa de los 60 dias sin contacto',
    'uso_bajo':     'usa menos del 40% de su plan',
    'auditoria':    'auditoria entregada y cuenta en riesgo',
    'soporte':      'volumen de tickets o fallas que conversar',
    'sin_corte':    'no aparece en el ultimo corte',
    'rotacion':     'le toca su turno',
}


def norm(s):
    s = ''.join(c for c in unicodedata.normalize('NFD', s or '')
                if unicodedata.category(c) != 'Mn').lower()
    return re.sub(r'[^a-z0-9]', '', s)


def num(v):
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(',', '').strip())
    except (TypeError, ValueError):
        return 0.0


def mes_de(v):
    if isinstance(v, datetime):
        return v.strftime('%Y-%m')
    if isinstance(v, (int, float)):
        return (datetime(1899, 12, 30) + timedelta(days=float(v))).strftime('%Y-%m')
    return str(v or '')[:7]


RX_EXT = re.compile(r'(\d+)\s*(?:extensi[oó]n(?:es)?|ext\b)', re.I)
RX_ABR = re.compile(r'^(\d+)\s+\S.*\bIL\b', re.I)
RX_SV = re.compile(r'\bchat\b|\bagentes?\s+cp\b|sin\s+saldo|n[uú]meros?\s+virtuales?|whatsapp', re.I)


def base_min(plan, incl):
    n = str(plan or '')
    m = RX_EXT.search(n) or RX_ABR.search(n)
    ext = int(m.group(1)) if m else None
    if RX_SV.search(n) and not ext and incl < 3:
        return None
    if ext and ext > 0:
        return incl if (incl > 0 and incl / ext >= 50) else ext * 1500
    return incl if incl >= 3 else None


# ── datos ────────────────────────────────────────────────────────────
E = {}
for ln in io.open(os.path.join(RAIZ, '.env.local'), encoding='utf-8'):
    if '=' in ln and not ln.strip().startswith('#'):
        k, v = ln.split('=', 1)
        E[k.strip()] = v.strip().strip('"').strip("'")
H = {'apikey': E['SUPABASE_SERVICE_ROLE_KEY'],
     'Authorization': 'Bearer ' + E['SUPABASE_SERVICE_ROLE_KEY']}
B = E['NEXT_PUBLIC_SUPABASE_URL'] + '/rest/v1/'


def sb(t, p):
    f, o = [], 0
    while True:
        d = json.load(urllib.request.urlopen(
            urllib.request.Request(B + t + '?' + p + '&offset=%d&limit=1000' % o, headers=H), timeout=60))
        f += d
        if len(d) < 1000:
            return f
        o += 1000


CU = sb('cuentas', 'select=id,consecutivo,empresa,cid,asesor,estado,facturacion,'
                   'contacto_nombre,contactos_json')
SEG = sb('seguimientos', 'select=cuenta_id,fecha')
ACT = sb('actividades', 'select=cuenta_id,tipo,completada')

ultimo = {}
for s in SEG:
    f = str(s.get('fecha') or '')[:10]
    c = s.get('cuenta_id')
    if f and c and (c not in ultimo or f > ultimo[c]):
        ultimo[c] = f

# vueltas previas de foco por cuenta (decide el trabajo)
vueltas = Counter(a['cuenta_id'] for a in ACT if a.get('tipo') == 'foco_riesgo' and a.get('cuenta_id'))
abiertas = {a['cuenta_id'] for a in ACT
            if a.get('tipo') == 'foco_riesgo' and a.get('cuenta_id') and not a.get('completada')}

# churn: la regla exacta de lib/elegibilidad.ts
src = io.open(os.path.join(RAIZ, 'app', 'churn', 'aaa-grc-data.ts'), encoding='utf-8').read()
rep = io.open(os.path.join(RAIZ, 'app', 'churn', 'grc-reporte.ts'), encoding='utf-8').read()
m = re.search(r"GRC_MES_EN_CURSO\s*(?::[^=]*)?=\s*(?:'([^']*)'|null)", rep)
if not m:
    sys.exit('*** No se encontro GRC_MES_EN_CURSO en app/churn/grc-reporte.ts')
MES_VIVO = (m.group(1) or '').lower()
_re = io.open(os.path.join(RAIZ, 'lib', 'elegibilidad.ts'), encoding='utf-8').read()
_bl = re.search(r'REACTIVADAS_FUERA_DEL_CHURN[^=]*=\s*new Map\(\[(.*?)\]\)', _re, re.S)
REACTIVADAS = {norm(x) for x in re.findall(r"\['([^']+)'", _bl.group(1))} if _bl else set()

CHURN = set()
for b in re.finditer(r"mes:\s*'([^']+)'([\s\S]*?)(?=\n\s*\{\s*mes:|\Z)", src):
    if MES_VIVO and b.group(1).lower() == MES_VIVO:
        continue
    for x in re.finditer(r"cliente:\s*'([^']+)'[^}]*?movimiento:\s*'([^']*)'", b.group(2)):
        if 'Churn confirmado' in x.group(2):
            n = norm(x.group(1))
            if n not in REACTIVADAS:
                CHURN.add(n)
CHURN.discard('')
CANCEL = {norm(x) for x in re.findall(
    r"cliente:\s*'([^']+)'",
    io.open(os.path.join(RAIZ, 'lib', 'churn-cancelados-data.ts'), encoding='utf-8').read())}

AUD = set()
for f in glob.glob(os.path.join(RAIZ, 'app', 'auditoria', '*-data.ts')):
    t = io.open(f, encoding='utf-8').read()
    e = re.search(r"^  estado:\s+'([^']+)'", t, re.M)
    n = re.search(r"^  nombre:\s+'([^']+)'", t, re.M)
    if e and n and e.group(1) in ('en_riesgo', 'rescatable', 'en_recuperacion'):
        AUD.add(norm(n.group(1)))

wb = openpyxl.load_workbook(os.path.join(RAIZ, 'data', 'cortes-facturacion.xlsx'),
                            data_only=True, read_only=True)
ws = wb[wb.sheetnames[0]]
it = ws.iter_rows(values_only=True)
cab = [str(c).strip() if c is not None else '' for c in next(it)]
idx = {c: i for i, c in enumerate(cab)}
cortes = defaultdict(list)
for r in it:
    f = {c: (r[i] if i < len(r) else None) for c, i in idx.items()}
    cid = str(f.get('CID') or '').strip()
    if cid:
        cortes[cid].append({'mes': mes_de(f.get('Fecha de corte')),
                            'cons': num(f.get('Minutos Consumidos')),
                            'base': base_min(f.get('Nombre del Plan'), num(f.get('Minutos Incluidos')))})
wb.close()
for v in cortes.values():
    v.sort(key=lambda x: x['mes'])
ULTIMO = max((c['mes'] for v in cortes.values() for c in v), default='')

TK = json.load(io.open(os.path.join(RAIZ, 'lib', 'tickets-data.json'), encoding='utf-8'))
tkc = defaultdict(lambda: {'t': 0, 'f': 0})
for t in TK:
    c = str(t.get('cid') or '').strip()
    if c:
        tkc[c]['t'] += 1
        if t.get('es_falla') == 'Si':
            tkc[c]['f'] += 1

LL = io.open(os.path.join(RAIZ, 'app', 'cuentas', 'llamadas-data.ts'), encoding='utf-8').read()
CIDS_LL = set(re.findall(r'\n\s*"(\d{2,8})":', LL))

# ── la regla ─────────────────────────────────────────────────────────
filas = []
for c in CU:
    n = norm(c['empresa'])
    if n in CHURN or n in CANCEL:
        continue
    if str(c.get('estado') or '').strip() not in ('activo', 'en_riesgo'):
        continue
    cid = str(c.get('cid') or '').strip()
    serie = cortes.get(cid, [])
    uc = serie[-1] if serie else None
    uso = 100.0 * uc['cons'] / uc['base'] if uc and uc['base'] else None
    u = ultimo.get(c['id'], '')
    dias = (HOY - datetime.strptime(u, '%Y-%m-%d')).days if u else None
    tk = tkc.get(cid, {'t': 0, 'f': 0})

    if not u:
        clase = 'nunca_tocada'
    elif dias is not None and dias > DIAS_LIMITE:
        clase = 'sin_contacto'
    elif uso is not None and uso < UMBRAL_USO:
        clase = 'uso_bajo'
    elif n in AUD:
        clase = 'auditoria'
    elif tk['t'] >= 15 or tk['f'] >= 2:
        clase = 'soporte'
    elif uc and ULTIMO and uc['mes'] < ULTIMO:
        clase = 'sin_corte'
    else:
        clase = 'rotacion'

    v = vueltas.get(c['id'], 0)
    trabajo = 'relacion' if clase in ('nunca_tocada', 'sin_contacto') else TRABAJOS[v % len(TRABAJOS)]
    if trabajo == 'llamadas' and cid not in CIDS_LL:
        trabajo = 'llamadas (sin lectura: sacarlas de Callpicker y entregarlas)'

    filas.append({'c': c, 'clase': clase, 'dias': dias, 'trabajo': trabajo,
                  'peso': float(c.get('facturacion') or 0),
                  'abierta': c['id'] in abiertas, 'll': cid in CIDS_LL})

print('=== FOCOS QUE SE DISPARAN EL LUNES %s ===' % LUNES)
print('    %d por asesor · churn confirmado NO entra · el trabajo lo decide la vuelta' % POR_SEMANA)
print()
resumen = Counter()
trabajos = Counter()
for a in sorted({f['c'].get('asesor') or '[sin asesor]' for f in filas}):
    mias = [f for f in filas if (f['c'].get('asesor') or '[sin asesor]') == a and not f['abierta']]
    mias.sort(key=lambda f: (ORDEN.index(f['clase']),
                             -(f['dias'] if f['dias'] is not None else 99999),
                             -f['peso']))
    vivas = len([f for f in filas if (f['c'].get('asesor') or '[sin asesor]') == a])
    print('  ── %s · %d cuentas vivas · salen %d ──' % (a, vivas, min(POR_SEMANA, len(mias))))
    print('   %-5s %-28s %-24s %-7s %-11s %s'
          % ('#', 'CUENTA', 'POR QUE SALE', 'DIAS', 'FACTURA', 'TRABAJO'))
    for f in mias[:POR_SEMANA]:
        c = f['c']
        resumen[f['clase']] += 1
        trabajos[f['trabajo']] += 1
        print('   %-5s %-28s %-24s %-7s $%-10s %s'
              % (c.get('consecutivo') or '?', c['empresa'][:28], MOTIVO[f['clase']][:24],
                 'nunca' if f['dias'] is None else f['dias'],
                 format(int(f['peso']), ','), f['trabajo']))
    print()

print('=== RESUMEN DEL LUNES ===')
print('  POR QUE SALEN:')
for k in ORDEN:
    if resumen[k]:
        print('    %-14s %2d   (%s)' % (k, resumen[k], MOTIVO[k]))
print('  QUE TRABAJO LLEVAN:')
for k, v in trabajos.most_common():
    print('    %-52s %2d' % (k, v))
print('    %-52s %2d focos' % ('TOTAL', sum(resumen.values())))
print()
sinll = sum(1 for f in filas if not f['ll'])
print('  Cartera viva: %d cuentas · sin lectura de llamadas: %d.' % (len(filas), sinll))
print('  A esas, cuando les toque el trabajo de llamadas, les sale sacar la')
print('  extraccion de Callpicker y entregarsela a Jose Manuel.')
print()
print('  Ademas de estos focos salen las 4 actividades rutinarias por asesor y')
print('  las Aclaraciones de baja, que van por su propia via y fuera de este tope.')
