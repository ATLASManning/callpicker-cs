"""Convierte el export del detalle de Zoho en data/mrr-zoho.json.

   FUENTE. El tablero «DASHBOARD GROSS REVENUE CHURN (%)» de Zoho Analytics:
   clic derecho sobre el MRR inicio del mes -> «Ver datos subyacentes» ->
   Mas -> «Exportar Vista». Septiembre 2026 trae 2,574 filas.

   POR QUE UN ARCHIVO Y NO EL API. El token que tiene el proyecto no alcanza:
   no puede ni listar las vistas del workspace (INVALID_OAUTHSCOPE) y Zoho no
   permite exportar un dashboard de forma sincrona (SYNC_EXPORT_NOT_ALLOWED).
   Raspar 2,574 filas de una lista virtualizada cada semana es fragil y se
   rompe con cualquier cambio de su interfaz. El export es un clic y no se
   rompe.

   EL MAPEO, que direccion fijo el 17 sep 2026:
     factura_mensual_zoho  <-  «MRR Inicio Contrato (BCY)»
     mrr_zoho              <-  «Importe Acumulado Recurrente»

   OJO CON EL SEGUNDO. El campo se sigue llamando `mrr_zoho` en la base porque
   asi lo pidio direccion, pero ya NO es un ingreso mensual: es lo que el
   cliente ha pagado en toda su vida. Por eso en pantalla se rotula «Acumulado
   recurrente» y NUNCA «MRR», y por eso quedo fuera de toda suma mensual. Son
   magnitudes distintas por 35 a 65 veces.

   Uso:  python scripts/gen-mrr-zoho.py "D:/ruta/al/export.xlsx"
"""
import sys, io, os, json, re, unicodedata, collections
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', line_buffering=True)

SALIDA = r'D:\Windows\Projects\callpicker-cs\data\mrr-zoho.json'

# Los nombres tal como vienen en el export, con alternativas por si Zoho los
# renombra. Se resuelven POR NOMBRE, nunca por posicion: una columna que se
# mueve no debe cambiar en silencio lo que alimenta a la cartera.
ALIAS = {
    'cliente':   ['Cliente', 'cliente'],
    'mes':       ['Mes Nombre', 'mes_nombre'],
    'fecha':     ['Año_Mes fecha', 'Ano_Mes fecha', 'ano_mes_fecha'],
    'clasif':    ['clasificacion_cliente', 'Clasificación Cliente'],
    'facturas':  ['Facturas_2026', 'facturas_2026'],
    'meses':     ['Meses Activo', 'meses_activo'],
    'acumulado': ['Importe Acumulado Recurrente', 'importe_acumulado_recurrente'],
    'mrrIni':    ['MRR Inicio Contrato (BCY)', 'MRR Inicio Contrato', 'mrr_inicio_contrato'],
    'mrrFin':    ['MRR Fin Contrato (BCY)', 'MRR Fin Contrato', 'mrr_fin_contrato'],
    'movimiento': ['Movimiento', 'movimiento'],
    'rango':     ['Rango', 'rango', 'Rango de facturación'],
}


def norm(s):
    """Misma normalizacion que lib/zoho-enrich.ts: sin acentos, sin signos."""
    s = ''.join(c for c in unicodedata.normalize('NFD', str(s or ''))
                if not (0x300 <= ord(c) <= 0x36f)).lower()
    return re.sub(r'[^a-z0-9\s]', '', s).strip()


def num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r'[^0-9.\-]', '', str(v))
    try:
        return float(s) if s not in ('', '-', '.') else None
    except ValueError:
        return None


if len(sys.argv) < 2:
    print(__doc__)
    raise SystemExit('Falta la ruta del export.')
ORIGEN = sys.argv[1]
if not os.path.exists(ORIGEN):
    raise SystemExit('No existe: %s' % ORIGEN)

import openpyxl
wb = openpyxl.load_workbook(ORIGEN, data_only=True, read_only=True)
ws = wb[wb.sheetnames[0]]
it = ws.iter_rows(values_only=True)
cab = [str(c).strip() if c is not None else '' for c in next(it)]
ix = {k: next((cab.index(n) for n in ns if n in cab), None) for k, ns in ALIAS.items()}

print('hoja: %s' % wb.sheetnames[0])
print('columnas resueltas:')
for k, v in ix.items():
    print('   %-11s %s' % (k, ('col %d · «%s»' % (v, cab[v])) if v is not None else '** NO ENCONTRADA'))
faltan = [k for k in ('cliente', 'acumulado', 'mrrIni') if ix[k] is None]
if faltan:
    raise SystemExit('Sin estas columnas no se puede continuar: %s' % faltan)

filas, porCliente = 0, {}
movs = collections.Counter()
for r in it:
    filas += 1
    nom = str(r[ix['cliente']] or '').strip()
    if not nom:
        continue
    k = norm(nom)
    if not k:
        continue
    mov = str(r[ix['movimiento']] or '').strip() if ix['movimiento'] is not None else ''
    movs[mov or '(vacío)'] += 1
    # Si un cliente aparece varias veces se suma: el export trae una fila por
    # contrato, y una cuenta puede tener mas de uno. Sumar es lo correcto —
    # quedarse con la ultima perderia contratos sin avisar.
    d = porCliente.setdefault(k, {
        'nombre': nom, 'acumulado': 0.0, 'mrrInicio': 0.0, 'contratos': 0,
        'cerrados': 0, 'clasif': None, 'mesesActivo': None, 'movimientos': [],
    })
    cerrado = 'churn' in mov.lower()
    d['acumulado'] += num(r[ix['acumulado']]) or 0.0
    # EL MENSUAL EXCLUYE LOS CONTRATOS CERRADOS. Un contrato en «Churn
    # confirmado» no paga nada este mes: sumarlo inflaria la factura mensual
    # con dinero que ya no entra. El ACUMULADO si los incluye, porque ese
    # dinero si se cobro en su momento — son dos preguntas distintas y por eso
    # se tratan distinto.
    if not cerrado:
        d['mrrInicio'] += num(r[ix['mrrIni']]) or 0.0
    else:
        d['cerrados'] += 1
    d['contratos'] += 1
    if ix['clasif'] is not None and d['clasif'] is None:
        d['clasif'] = str(r[ix['clasif']] or '').strip() or None
    if ix['meses'] is not None:
        m = num(r[ix['meses']])
        if m is not None:
            d['mesesActivo'] = max(d['mesesActivo'] or 0, int(m))
    if mov and mov not in d['movimientos']:
        d['movimientos'].append(mov)
wb.close()

META = {
    'origen': os.path.basename(ORIGEN),
    'filas': filas,
    'clientes': len(porCliente),
    'mapeo': {
        'factura_mensual_zoho': cab[ix['mrrIni']],
        'mrr_zoho': cab[ix['acumulado']],
    },
    'advertencia': ('`mrr_zoho` trae el ACUMULADO de toda la vida del cliente, no un '
                    'ingreso mensual. En pantalla se rotula «Acumulado recurrente» y '
                    'queda fuera de cualquier suma mensual.'),
}

os.makedirs(os.path.dirname(SALIDA), exist_ok=True)
with io.open(SALIDA, 'w', encoding='utf-8') as f:
    json.dump({'meta': META, 'clientes': porCliente}, f, ensure_ascii=False, separators=(',', ':'))

print()
print('filas leídas : %s' % format(filas, ','))
print('clientes     : %s' % format(len(porCliente), ','))
print('movimientos  : %s' % dict(movs.most_common(6)))
tot_a = sum(d['acumulado'] for d in porCliente.values())
tot_m = sum(d['mrrInicio'] for d in porCliente.values())
print('suma acumulado : $%s' % format(round(tot_a), ','))
print('suma MRR inicio: $%s' % format(round(tot_m), ','))
print('  -> el acumulado es %.0f veces el mensual, como debe ser' % (tot_a / tot_m if tot_m else 0))
print()
print('escrito %s (%.0f KB)' % (SALIDA, os.path.getsize(SALIDA) / 1024))
print()
print('los 6 de mayor acumulado:')
for k, d in sorted(porCliente.items(), key=lambda x: -x[1]['acumulado'])[:6]:
    print('  %-38s acumulado $%-14s mensual $%s'
          % (d['nombre'][:38], format(round(d['acumulado']), ','), format(round(d['mrrInicio']), ',')))

# ── EL RIESGO DE LA REESTRUCTURA ───────────────────────────────────────────
# Direccion confirmo (17 sep 2026) que el churn de $1.7M de septiembre NO es
# churn: es la reestructura de facturacion de GTC. Una reestructura crea DOS
# representaciones del mismo dinero — las subcuentas viejas cerrando y la
# fusionada abriendo— y ese fue exactamente el mecanismo que inflo el churn.
#
# El mismo mecanismo amenaza a la cartera por otro lado: lookupZoho() de
# lib/zoho-enrich.ts suma TODAS las claves cuya primera palabra es la sigla de
# la empresa. Para «GRUPO TORRES CORZO» (sigla gtc) eso alcanza a cada
# «GTC - ...» del archivo. Si conviven las subcuentas viejas y la fusionada,
# esa plata se suma dos veces y la cuenta aparece facturando el doble.
#
# Aqui no se decide por nadie: se NOMBRA el caso para poder mirarlo.
print()
print('=== POSIBLE DOBLE CONTEO POR REESTRUCTURA ===')
grupos = collections.defaultdict(list)
for k, d in porCliente.items():
    ini = d['nombre'].split('-')[0].strip()
    if len(ini) >= 3 and len(ini.split()) <= 2:
        grupos[ini.lower()].append(d)
sosp = {g: v for g, v in grupos.items() if len(v) > 1}
print('  familias de subcuentas con prefijo comun: %d' % len(sosp))
for g, v in sorted(sosp.items(), key=lambda x: -sum(y['acumulado'] for y in x[1]))[:5]:
    cerr = sum(1 for y in v if y['cerrados'])
    print('    %-14s %2d subcuentas · %2d con contrato cerrado · acumulado $%s'
          % (g.upper()[:14], len(v), cerr, format(round(sum(y['acumulado'] for y in v)), ',')))
    if cerr:
        print('      ojo: conviven cerradas y activas. Si la cerrada y la nueva son el')
        print('      mismo dinero reestructurado, sumarlas lo cuenta dos veces.')
