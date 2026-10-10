# -*- coding: utf-8 -*-
"""¿Un 0 en «Dias activacion» es «no medido» o es «se activó el mismo día»?

   POR QUE EXISTE
   --------------
   `app/activaciones/page.tsx` convierte la columna así:

       diasActivacion: typeof r['Dias activacion'] === 'number'
                       && r['Dias activacion'] > 0 ? r['Dias activacion'] : null

   O sea que un CERO se vuelve `null`, «no medido». Si el 0 significara «se
   activó el mismo día», eso borra justamente el mejor caso de la operación y
   lo saca del conteo de «en 7 días o menos». Y al revés: si el 0 es basura de
   captura, contarlo como día cero hunde el promedio.

   La diferencia no se puede decidir leyendo el código: hay que mirar QUÉ
   TRAEN esas filas. Una activación con 0 días que sí tiene fecha de primer
   pago, importe y ejecutivo se parece a una activación del mismo día; una con
   0 días y todo lo demás vacío se parece a un registro sin capturar.

   Es de SOLO LECTURA.

   USO
   ---
       python scripts/mide-dias-cero.py
"""
import io
import os
import re
import sys

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', line_buffering=True)
from openpyxl import load_workbook

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(RAIZ, 'data', 'activaciones.xlsx')
COL = 'Dias activacion'

wb = load_workbook(XLSX, data_only=True, read_only=True)
ws = wb['Hoja1']
it = ws.iter_rows(values_only=True)
cab = [str(c).strip() if c is not None else '' for c in next(it)]
filas = []
for r in it:
    if all(v is None or str(v).strip() == '' for v in r):
        continue
    filas.append({cab[i]: r[i] for i in range(min(len(cab), len(r)))})
wb.close()
print(u'  columna «%s»: %s\n' % (COL, u'existe' if COL in cab else u'*** NO EXISTE'))


def num(v):
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(re.sub(r'[$,\s]', '', str(v or '')))
    except ValueError:
        return 0.0


regs = [f for f in filas
        if str(f.get('ID') or '').strip() and str(f.get('Cliente') or '').strip()]
regs = [f for f in regs if num(f.get(u'Año')) >= 2020]
act = [f for f in regs if str(f.get('Tipo') or '').strip().lower() != 'demo']

# ── Qué formas toma la celda ────────────────────────────────────────────────
formas = {}
for f in regs:
    v = f.get(COL)
    if v is None:
        k = u'None (celda vacía)'
    elif isinstance(v, (int, float)):
        k = u'número 0' if v == 0 else (u'número < 0' if v < 0 else u'número > 0')
    else:
        s = str(v).strip()
        k = u'texto «%s»' % s if len(s) <= 6 else u'texto largo'
    formas[k] = formas.get(k, 0) + 1
print(u'  LAS FORMAS DE LA CELDA, en las %d filas que la pantalla pinta:' % len(regs))
for k, v in sorted(formas.items(), key=lambda kv: -kv[1]):
    print(u'      %-22s %5d' % (k, v))
assert sum(formas.values()) == len(regs), u'el reparto de formas no cierra'

# ── ¿Los ceros son demos o son activaciones? ────────────────────────────────
ceros_todos = [f for f in regs if isinstance(f.get(COL), (int, float)) and f.get(COL) == 0]
ceros_act = [f for f in act if isinstance(f.get(COL), (int, float)) and f.get(COL) == 0]
print(u'\n  CEROS: %d en total, %d de ellos en activaciones (no demo)'
      % (len(ceros_todos), len(ceros_act)))

# ── La pregunta que decide: ¿esas filas están capturadas o vacías? ──────────
# Si traen mes de primer pago, importe y ejecutivo, son activaciones reales
# que se resolvieron el mismo día. Si no traen nada, son registros a medias.
print(u'\n  LAS %d ACTIVACIONES CON 0 DÍAS, una por una:' % len(ceros_act))
print(u'      %-28s %-10s %-12s %-12s %s'
      % (u'cliente', u'1er pago', u'mes', u'ejecutivo', u'complejidad'))
completas = 0
for f in sorted(ceros_act, key=lambda x: -num(x.get('1er Pago'))):
    mes = str(f.get('Mes 1er Pago') or '').strip()
    pago = num(f.get('1er Pago'))
    ejec = str(f.get('Ejecutivo') or '').strip()
    comp = str(f.get('Complejidad') or '').strip()
    if mes and pago > 0 and ejec:
        completas += 1
    print(u'      %-28s $%-9s %-12s %-12s %s'
          % (str(f.get('Cliente') or '')[:28], format(int(pago), ',d'),
             mes[:12] or u'(sin mes)', ejec[:12] or u'(sin ejec)', comp or u'(sin)'))

print(u'\n  %d de %d traen mes, importe y ejecutivo: son activaciones capturadas.'
      % (completas, len(ceros_act)))

# ── Y qué cambia en las cifras según cómo se lea el 0 ──────────────────────
CORTE = 3500.0
sel = [f for f in act if num(f.get('1er Pago')) > CORTE]


def medidas(grupo, cero_cuenta):
    ds = []
    for f in grupo:
        v = f.get(COL)
        if not isinstance(v, (int, float)):
            continue
        if v == 0 and not cero_cuenta:
            continue
        if v < 0:
            continue
        ds.append(float(v))
    return ds


for nombre, grupo in ((u'LAS %d GRANDES (más de $%s)'
                       % (len(sel), format(int(CORTE), ',d')), sel),
                      (u'TODAS LAS %d ACTIVACIONES' % len(act), act)):
    print(u'\n  LO QUE CAMBIA EN %s:' % nombre)
    for etq, cc in ((u'el 0 es «no medido» (como hoy)', False),
                    (u'el 0 es «mismo día»', True)):
        ds = sorted(medidas(grupo, cc))
        rap = sum(1 for d in ds if d <= 7)
        print(u'      %-32s n=%d  promedio %.1f  mediana %.0f  en <=7d %d (%.0f%%)'
              % (etq, len(ds), sum(ds) / len(ds), ds[len(ds) // 2],
                 rap, 100.0 * rap / len(ds)))

# Los NEGATIVOS sí son basura: no existen los días de activación negativos.
# Se quedan en `null` en las dos lecturas, y se dicen con nombre.
negs = [f for f in act if isinstance(f.get(COL), (int, float)) and f.get(COL) < 0]
print(u'\n  NEGATIVOS (basura de captura, siguen siendo «sin medir»): %d' % len(negs))
for f in negs:
    print(u'      %-30s %s días · $%s'
          % (str(f.get('Cliente') or '')[:30], f.get(COL),
             format(int(num(f.get('1er Pago'))), ',d')))

# ── El otro hallazgo que hay que comprobar: ¿hay activaciones en $0? ───────
en_cero = [f for f in act if num(f.get('1er Pago')) == 0]
print(u'\n  ACTIVACIONES CON PRIMER PAGO EN $0: %d' % len(en_cero))
for f in en_cero[:10]:
    print(u'      %-30s tipo %-12s mes «%s»'
          % (str(f.get('Cliente') or '')[:30], str(f.get('Tipo') or ''),
             str(f.get('Mes 1er Pago') or '').strip()))
print(u'  — importa porque una escalera de bandas con `> inferior` las deja')
print(u'    FUERA de todas las bandas y el cierre no da el total.')
