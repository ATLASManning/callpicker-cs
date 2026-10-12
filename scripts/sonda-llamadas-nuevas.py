# -*- coding: utf-8 -*-
u"""Audita archivos de llamadas ANTES de meterlos al generador.

   `gen-llamadas-data.py` documenta en su cabecera las tres trampas que ya
   costaron, y las tres se comprueban aqui ANTES de regenerar nada:

   1. UNA PARTE QUE ES COPIA DE OTRA. «Salientes Mayor consumo 40 Parte 2»
      resulto ser copia byte a byte de la Parte 1 —mismas 900,000 filas, mismo
      hash—. Sumarlas habria duplicado cada llamada saliente de 54 cuentas. Un
      «(1)» al final del nombre es el sufijo que pone Windows al descargar dos
      veces, y es exactamente esa trampa esperando.

   2. LAS COLUMNAS NO SE LLAMAN IGUAL ENTRE ARCHIVOS. La empresa viene como
      `empresa`, `Nombre Empresa` o `Nombre de Empresa`; el destino como
      `destination_data_1`… o no viene. Todo se resuelve POR NOMBRE y nunca
      por posicion.

   3. LA VENTANA DE FECHAS ES DISTINTA POR FUENTE. Las entrantes de un corte
      cubren de enero a septiembre y sus salientes arrancan en abril. Si no se
      guarda la ventana real, la pantalla lee la ausencia como un cero.

   Y una cuarta, que es la que importa cuando llega material nuevo:

   4. ¿APORTA COBERTURA O LA DUPLICA? Un CID que ya trae lectura no se vuelve
      a sumar. Se dice cuales son NUEVOS y cuales ya estaban.

   NO escribe nada. Solo mira y reporta.

   USO
       python scripts/sonda-llamadas-nuevas.py <archivo.xlsx> [mas...]
"""
import io
import os
import re
import sys
import json
import hashlib
import datetime
import collections

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', line_buffering=True)
import openpyxl

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

ALIAS = {
    'tipo':    ['destination_type'],
    'fecha':   ['date'],
    'cid':     ['customer_id'],
    'empresa': ['empresa', 'Nombre Empresa', 'Nombre de Empresa'],
    'destino': ['destination_data_1'],
    'caller':  ['caller_id'],
    'minutos': ['total_minutes'],
}


def sha1(ruta):
    h = hashlib.sha1()
    with open(ruta, 'rb') as f:
        for trozo in iter(lambda: f.read(1 << 20), b''):
            h.update(trozo)
    return h.hexdigest()


def cids_ya_cubiertos():
    u"""Los CID que YA tienen lectura en `llamadas-data.ts`."""
    p = os.path.join(RAIZ, 'app', 'cuentas', 'llamadas-data.ts')
    if not os.path.exists(p):
        return set(), None
    txt = io.open(p, encoding='utf-8').read()
    # OJO: `export const LLAMADAS` tambien casa con `LLAMADAS_META`, y esa
    # confusion ya costo clasificar 222 cuentas contra 12 claves de metadatos.
    # Se ancla al nombre EXACTO.
    m = re.search(r'export const LLAMADAS\s*:\s*Record<[^>]*>\s*=\s*(\{)', txt)
    meta = re.search(r'export const LLAMADAS_META\s*:\s*\w+\s*=\s*(\{.*?\})\n', txt, re.S)
    if not m:
        return set(), None
    # Las claves de primer nivel del objeto.
    ini = m.end(1)
    nivel, i = 1, ini
    while i < len(txt) and nivel > 0:
        if txt[i] == '{':
            nivel += 1
        elif txt[i] == '}':
            nivel -= 1
        i += 1
    cuerpo = txt[ini:i - 1]
    claves = set(re.findall(r'"([^"]+)"\s*:\s*\{', cuerpo))
    try:
        m2 = json.loads(meta.group(1)) if meta else None
    except Exception:
        m2 = None
    return claves, m2


def mira(ruta):
    nombre = os.path.basename(ruta)
    print(u'\n  ' + u'─' * 74)
    print(u'  %s' % nombre)
    if not os.path.exists(ruta):
        print(u'    *** NO EXISTE'); return None
    tam = os.path.getsize(ruta)
    h = sha1(ruta)
    print(u'    %.1f MB · sha1 %s' % (tam / 1048576.0, h[:16]))

    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    print(u'    hojas: %s' % u', '.join(wb.sheetnames))
    hoja = wb[wb.sheetnames[0]]
    filas = hoja.iter_rows(values_only=True)
    try:
        cab = next(filas)
    except StopIteration:
        print(u'    *** hoja vacia'); return None
    cols = [u'' if c is None else unicode(c) if sys.version_info[0] < 3 else str(c)
            for c in cab]
    print(u'    %d columna(s): %s' % (len(cols), u', '.join(cols[:14])))

    # ¿Que alias resuelve cada campo que el generador necesita?
    idx = {}
    for campo, nombres in ALIAS.items():
        for n in nombres:
            if n in cols:
                idx[campo] = cols.index(n)
                break
    faltan = [c for c in ALIAS if c not in idx]
    for campo in sorted(ALIAS):
        print(u'      %-8s %s' % (campo,
              u'col «%s»' % cols[idx[campo]] if campo in idx else u'*** NO VIENE'))

    # Recorrido: CIDs, fechas, filas.
    n = 0
    cids = collections.Counter()
    fmin = fmax = None
    for r in filas:
        n += 1
        if 'cid' in idx and idx['cid'] < len(r) and r[idx['cid']] is not None:
            cids[str(r[idx['cid']]).strip()] += 1
        if 'fecha' in idx and idx['fecha'] < len(r):
            v = r[idx['fecha']]
            d = None
            if isinstance(v, datetime.datetime):
                d = v.date()
            elif isinstance(v, datetime.date):
                d = v
            elif v is not None:
                mm = re.match(r'(\d{4})-(\d{2})-(\d{2})', str(v))
                if mm:
                    d = datetime.date(*[int(x) for x in mm.groups()])
            if d:
                fmin = d if fmin is None or d < fmin else fmin
                fmax = d if fmax is None or d > fmax else fmax
    wb.close()

    print(u'    %s fila(s) · %d CID(s) distintos' % (u'{:,}'.format(n), len(cids)))
    print(u'    ventana: %s → %s' % (fmin or u'—', fmax or u'—'))
    if faltan:
        print(u'    *** sin %s' % u', '.join(sorted(faltan)))
    return {'nombre': nombre, 'sha1': h, 'filas': n, 'cids': set(cids),
            'conteo': cids, 'fmin': fmin, 'fmax': fmax, 'faltan': faltan}


if len(sys.argv) < 2:
    print(__doc__); sys.exit(2)

vistos = [mira(a) for a in sys.argv[1:]]
vistos = [v for v in vistos if v]

# ── 1 · ¿HAY COPIAS? ───────────────────────────────────────────────────────
print(u'\n  ' + u'═' * 74)
print(u'  COPIAS EXACTAS')
porHash = collections.defaultdict(list)
for v in vistos:
    porHash[v['sha1']].append(v['nombre'])
dup = {h: ns for h, ns in porHash.items() if len(ns) > 1}
if dup:
    for h, ns in dup.items():
        print(u'    *** MISMO ARCHIVO, %d veces: %s' % (len(ns), u' · '.join(ns)))
    print(u'    Sumarlos duplicaria cada llamada de esos CIDs. Se lee UNO.')
else:
    print(u'    ninguna: los %d son archivos distintos.' % len(vistos))

# Y copias por CONTENIDO aunque el hash difiera (reexport del mismo corte).
print(u'\n  SOLAPE ENTRE ARCHIVOS (mismos CIDs en dos fuentes)')
hubo = False
for i in range(len(vistos)):
    for j in range(i + 1, len(vistos)):
        a, b = vistos[i], vistos[j]
        comun = a['cids'] & b['cids']
        if not comun:
            continue
        hubo = True
        mismas = sum(1 for c in comun if a['conteo'][c] == b['conteo'][c])
        print(u'    %s ∩ %s = %d CID(s)%s'
              % (a['nombre'][:30], b['nombre'][:30], len(comun),
                 u'  *** y %d con EL MISMO numero de filas: sospecha de recorte repetido'
                 % mismas if mismas else u''))
if not hubo:
    print(u'    ninguno: no se solapan.')

# ── 2 · ¿APORTAN COBERTURA? ────────────────────────────────────────────────
ya, meta = cids_ya_cubiertos()
print(u'\n  COBERTURA')
print(u'    `llamadas-data.ts` trae hoy %d CID(s).' % len(ya))
if meta:
    print(u'    corte actual: %s · mes cerrado %s' % (meta.get('corte'), meta.get('mesCerrado')))
todos = set()
for v in vistos:
    todos |= v['cids']
nuevos = todos - ya
repetidos = todos & ya
print(u'    los archivos traen %d CID(s): %d NUEVOS y %d que ya estaban.'
      % (len(todos), len(nuevos), len(repetidos)))
if nuevos:
    print(u'    nuevos: %s%s' % (u', '.join(sorted(nuevos)[:18]),
                                 u' …' if len(nuevos) > 18 else u''))
print(u'\n  Lo que ya estaba NO se vuelve a sumar: el generador reescribe por')
print(u'  CID, pero si una fuente nueva cubre OTRO periodo del mismo CID hay')
print(u'  que decidir si extiende la ventana o la reemplaza. Eso se mira caso')
print(u'  por caso, no se asume.')
