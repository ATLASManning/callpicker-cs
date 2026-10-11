# -*- coding: utf-8 -*-
u"""Corrobora que un export «GRC - AAA» cuadre con lo YA PUBLICADO en el modulo.

   POR QUE EXISTE
   --------------
   El modulo Churn · GRC AAA se regenera desde el export INSTITUCIONAL
   (`Gross_Revenue_Facturación.xlsx`), por instruccion de direccion: una sola
   fuente para los dos modulos, «no puede haber dos cifras».

   Cuando llega aparte un export del tablero AAA, lo que toca NO es regenerar
   —eso crearia la segunda cifra que la instruccion prohibe— sino COMPROBAR que
   dice lo mismo. Si cuadra, no hay nada que hacer. Si no cuadra, la diferencia
   se reporta; no se elige un archivo a dedo.

   Es de SOLO LECTURA. No escribe nada.

   COMO COMPARA
   ------------
   El `.ts` publicado lleva TODAS las clasificaciones (AAA · AA · A · B · C) y
   solo las filas de PERDIDA. El export AAA puede traer un subconjunto. Asi que
   el filtro no se supone: se LEE del propio xlsx —que clasificaciones y que
   meses trae— y el `.ts` se recorta a ese mismo universo antes de comparar.
   Comparar un subconjunto contra un total da una diferencia que no es un
   error, y eso ya costo una conclusion falsa en este proyecto.

   USO
   ---
       python scripts/concilia-grc-aaa-xlsx.py "C:/ruta/GRC_-_AAA_-_2026.xlsx"
"""
import collections
import io
import os
import re
import sys
import unicodedata

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', line_buffering=True)
from openpyxl import load_workbook

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TS = os.path.join(RAIZ, 'app', 'churn', 'aaa-grc-data.ts')
CENT = 0.01

if len(sys.argv) < 2:
    print(u'  uso: python scripts/concilia-grc-aaa-xlsx.py "<ruta del xlsx>"')
    sys.exit(2)
XLSX = sys.argv[1]


def norma(s):
    s = unicodedata.normalize('NFD', str(s or ''))
    s = ''.join(c for c in s if unicodedata.category(c) != 'Mn').lower()
    return re.sub(r'\s+', ' ', s).strip()


def num(v):
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(re.sub(r'[$,\s]', '', str(v or '')))
    except ValueError:
        return 0.0


# ── 1 · El xlsx, tal como viene ─────────────────────────────────────────────
wb = load_workbook(XLSX, data_only=True, read_only=True)
ws = wb[wb.sheetnames[0]]
it = ws.iter_rows(values_only=True)
cab = [str(c).strip() if c is not None else '' for c in next(it)]
ix = {c: i for i, c in enumerate(cab)}


def col(*cand):
    for c in cab:
        for k in cand:
            if k in norma(c):
                return c
    return None


K_MES = col('mes nombre', 'mes')
K_CLI = col('cliente')
K_CLA = col('clasificacion')
K_MOV = col('movimiento')
K_PER = col('ingreso perdido contrato (bcy) real', 'perdido contrato (bcy) real')
K_PE2 = col('fraude')

print(u'  === EL EXPORT ===')
print(u'  %s' % os.path.basename(XLSX))
print(u'  hoja «%s» · %d columnas' % (wb.sheetnames[0], len(cab)))
for etq, k in ((u'mes', K_MES), (u'cliente', K_CLI), (u'clasificacion', K_CLA),
               (u'movimiento', K_MOV), (u'perdido real', K_PER), (u'fraude/reestr.', K_PE2)):
    print(u'     %-16s %s' % (etq, k if k else u'*** NO SE ENCUENTRA'))
if not all((K_MES, K_CLI, K_MOV, K_PER)):
    print(u'\n  *** faltan columnas imprescindibles: no se puede comparar.')
    sys.exit(1)

filas = []
for r in it:
    if all(v is None or str(v).strip() == '' for v in r):
        continue
    filas.append({
        'mes':  str(r[ix[K_MES]] or '').strip(),
        'cli':  str(r[ix[K_CLI]] or '').strip(),
        'clas': str(r[ix[K_CLA]] or '').strip() if K_CLA else '',
        'mov':  str(r[ix[K_MOV]] or '').strip(),
        'per':  num(r[ix[K_PER]]),
        'per2': num(r[ix[K_PE2]]) if K_PE2 else 0.0,
    })
wb.close()

meses_x = collections.Counter(f['mes'] for f in filas)
clas_x = collections.Counter(f['clas'] for f in filas)
print(u'\n  filas: %s' % format(len(filas), ',d'))
print(u'  meses: %s' % u' · '.join(u'%s %d' % (k, v) for k, v in sorted(meses_x.items())))
print(u'  clasificaciones: %s'
      % u' · '.join(u'%s %d' % (k or u'(vacía)', v) for k, v in clas_x.most_common()))

# El modulo guarda SOLO filas de perdida. Si el xlsx trae de todo, se recorta.
con_perdida = [f for f in filas if abs(f['per']) > CENT or abs(f['per2']) > CENT]
print(u'  de ellas con pérdida (real o fraude): %s' % format(len(con_perdida), ',d'))

# ── 2 · Lo publicado ────────────────────────────────────────────────────────
FILA = re.compile(
    r"\{ cliente: '(.*?)', clas: '([^']*)',.*?"
    r"movimiento: '([^']*)', perdido: ([-\d.]+), perdido2: ([-\d.]+)")
txt = io.open(TS, encoding='utf-8').read()
pub = collections.defaultdict(list)
for b in re.split(r"\n  \{\n    mes: '", txt)[1:]:
    mes = b.split("'")[0]
    for cli, clas, mov, per, per2 in FILA.findall(b):
        pub[mes].append({'cli': cli, 'clas': clas, 'mov': mov,
                         'per': float(per), 'per2': float(per2)})

print(u'\n  === LO PUBLICADO (app/churn/aaa-grc-data.ts) ===')
print(u'  %s filas en %d meses: %s'
      % (format(sum(len(v) for v in pub.values()), ',d'), len(pub),
         u', '.join(pub.keys())))

# ── 3 · El universo COMUN, que es lo unico comparable ────────────────────────
meses_com = [m for m in pub if m in meses_x]
clas_com = {c for c in clas_x if c}
print(u'\n  === EL UNIVERSO COMPARABLE ===')
print(u'  meses en los DOS: %s' % (u', '.join(meses_com) or u'ninguno'))
solo_x = sorted(set(meses_x) - set(pub))
solo_p = sorted(set(pub) - set(meses_x))
if solo_x:
    print(u'  meses SOLO en el export : %s' % u', '.join(solo_x))
if solo_p:
    print(u'  meses SOLO en lo publicado: %s' % u', '.join(solo_p))
if clas_com:
    print(u'  el export trae estas clasificaciones: %s' % u', '.join(sorted(clas_com)))
    print(u'  -> lo publicado se RECORTA a ellas antes de comparar.')

fallos = []
print(u'\n  === MES A MES ===')
print(u'  %-12s %10s %10s %16s %16s %12s'
      % (u'MES', u'filas xls', u'filas pub', u'perdida xls', u'perdida pub', u'delta'))
for m in meses_com:
    fx = [f for f in con_perdida if f['mes'] == m
          and (not clas_com or f['clas'] in clas_com)]
    fp = [f for f in pub[m] if not clas_com or f['clas'] in clas_com]
    px = sum(f['per'] for f in fx)
    pp = sum(f['per'] for f in fp)
    d = pp - px
    ok = abs(d) < CENT
    if not ok:
        fallos.append(m)
    print(u'  %-12s %10d %10d %16s %16s %12s %s'
          % (m, len(fx), len(fp), format(px, ',.2f'), format(pp, ',.2f'),
             format(d, ',.2f'), u'' if ok else u'  <<<'))

    if not ok or len(fx) != len(fp):
        # Quien difiere, por cliente+movimiento. Sin esto solo se sabe QUE no
        # cuadra, que no sirve para resolverlo.
        ax = collections.defaultdict(float)
        ap = collections.defaultdict(float)
        for f in fx:
            ax[(norma(f['cli']), f['mov'])] += f['per']
        for f in fp:
            ap[(norma(f['cli']), f['mov'])] += f['per']
        difs = []
        for k in set(ax) | set(ap):
            if abs(ap[k] - ax[k]) > CENT:
                difs.append((ap[k] - ax[k], k, ax[k], ap[k]))
        for d2, k, vx, vp in sorted(difs, key=lambda x: -abs(x[0]))[:10]:
            print(u'        %-34s %-20s xls %12s  pub %12s  %s'
                  % (k[0][:34], k[1][:20], format(vx, ',.2f'),
                     format(vp, ',.2f'), format(d2, ',.2f')))

print(u'\n  ═══════════════════════════════════════════════════════════')
if fallos:
    print(u'  *** NO CUADRA en: %s' % u', '.join(fallos))
    print(u'  No se elige un archivo a dedo: se reporta la diferencia.')
    sys.exit(1)
print(u'  CUADRA en los %d meses comparables.' % len(meses_com))
if solo_x or solo_p:
    print(u'  (los meses que solo estan en uno de los dos no se comparan,')
    print(u'   y eso se dice arriba en vez de contarse como diferencia)')
