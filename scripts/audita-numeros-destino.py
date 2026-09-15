"""Que es exactamente la columna «Numeros» de «A donde entran las que no se contestan».

   Cadena completa:
     gen-analisis-llamadas.py  dd['n'] = set() de caller_id, SOLO en filas perdidas
                               -> se emite como 'n': len(x['n'])  (POR CUENTA)
     api/analisis-llamadas     e.n += x.n                          (SUMA entre cuentas)
     page.tsx:246              columna «Numeros» = x.n

   Dos preguntas que este script responde con datos, no con teoria:
     1. Numeros ¿cuenta llamadas o numeros distintos? -> distintos, y solo de
        las NO contestadas.
     2. Al sumar entre cuentas un conteo de distintos, ¿se duplica? Un mismo
        telefono que llamo al «Ventas» de dos clientes cuenta dos veces. Aqui se
        mide el tamano real de esa duplicacion: suma-de-distintos vs union real.

   Solo lee los archivos de ENTRANTES: la tabla es de entrantes.
"""
import sys, io, os, collections
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', line_buffering=True)
import openpyxl

ARCH = r"D:\Archivos"
ENTRANTES = [
    'Llamadas_entrantes_Clientes_AAA_Poco_Consumo_Actualizado.xlsx',
    'Entrantes Mayor consumo 40 Parte 1.xlsx',
    'Entrantes Mayor consumo 40 Parte 2.xlsx',
    'Entrantes Mayor consumo 40 Parte 3.xlsx',
]
SIN_DESTINO = '(sin destino registrado)'
NO_EXPORTADO = '(destino no venía en el archivo)'

# Los destinos que se ven en la pantalla del usuario.
FOCO = ['Ventas', 'Ventas 2', 'IMPE', 'VANESSA MEDINA GUTIERREZ',
        'Alejandra Casiano Solano', 'Agente Virtual OOAPAS',
        'Edith Rodriguez M', 'Lezly Iohanna Carreón Loera',
        'Alexa Yael Zavala Hernandez', SIN_DESTINO, NO_EXPORTADO]
FOCOSET = set(FOCO)


def texto(v):
    s = str(v or '').strip()
    return '' if s.upper() == 'NULL' else s


def cid_de(v):
    if v is None:
        return ''
    if isinstance(v, float) and float(v).is_integer():
        return str(int(v))
    s = str(v).strip()
    return '' if s.upper() == 'NULL' else s


# destino -> cid -> set(caller_id) de las PERDIDAS
porCuenta = collections.defaultdict(lambda: collections.defaultdict(set))
# destino -> conteos crudos
perdidas = collections.Counter()
contestadas = collections.Counter()
sinCaller = collections.Counter()      # perdidas sin caller_id: no entran a ningun set
cidsPorDestino = collections.defaultdict(set)

for arch in ENTRANTES:
    ruta = os.path.join(ARCH, arch)
    if not os.path.exists(ruta):
        print('!! FALTA %s' % arch)
        continue
    wb = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    cab = [str(c).strip() if c is not None else '' for c in next(it)]
    iT = cab.index('destination_type') if 'destination_type' in cab else None
    iC = cab.index('customer_id') if 'customer_id' in cab else None
    iD = cab.index('destination_data_1') if 'destination_data_1' in cab else None
    iK = cab.index('caller_id') if 'caller_id' in cab else None
    print('%-52s destino=%s caller=%s' % (arch[:52], 'si' if iD is not None else 'NO',
                                          'si' if iK is not None else 'NO'))
    n = 0
    for r in it:
        n += 1
        if n % 500000 == 0:
            print('    ... %s' % format(n, ','))
        cid = cid_de(r[iC])
        if not cid:
            continue
        dest = (texto(r[iD]) or SIN_DESTINO) if iD is not None else NO_EXPORTADO
        if dest not in FOCOSET:
            continue
        cidsPorDestino[dest].add(cid)
        esPerdida = texto(r[iT]) == 'Lost'
        if not esPerdida:
            contestadas[dest] += 1
            continue
        perdidas[dest] += 1
        num = texto(r[iK]) if iK is not None else ''
        if num:
            porCuenta[dest][cid].add(num)
        else:
            sinCaller[dest] += 1
    wb.close()
    print('    %s filas' % format(n, ','))

print()
print('=' * 104)
print('%-34s %9s %9s %6s %11s %11s %7s' %
      ('DESTINO', 'perdidas', 'contest.', 'CIDs', 'SUMA dist.', 'UNION real', 'infla'))
print('-' * 104)
for d in FOCO:
    if d not in porCuenta and d not in perdidas:
        continue
    suma = sum(len(s) for s in porCuenta[d].values())
    union = len(set().union(*porCuenta[d].values())) if porCuenta[d] else 0
    infla = ('+%.1f%%' % (100 * (suma - union) / union)) if union else '—'
    print('%-34s %9s %9s %6d %11s %11s %7s' %
          (d[:34], format(perdidas[d], ','), format(contestadas[d], ','),
           len(cidsPorDestino[d]), format(suma, ','), format(union, ','), infla))

print()
print('Perdidas SIN caller_id (no entran a ningun conteo de numeros):')
for d in FOCO:
    if sinCaller[d]:
        print('  %-34s %s de %s perdidas' % (d[:34], format(sinCaller[d], ','), format(perdidas[d], ',')))
if not any(sinCaller[d] for d in FOCO):
    print('  ninguna — toda perdida traia caller_id')

print()
print('LECTURA: «SUMA dist.» es lo que publica la pantalla hoy. «UNION real» es')
print('el numero de telefonos distintos de verdad. Si difieren, la columna esta')
print('sumando conjuntos que se traslapan entre cuentas.')
