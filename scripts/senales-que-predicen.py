# -*- coding: utf-8 -*-
"""Que senales SI separan a las cuentas que se fueron de las que se quedaron.

   POR QUE EXISTE
   --------------
   `valida-riesgo.py` probo un puntaje compuesto contra el churn real y salio
   PEOR QUE EL AZAR: en el top 10 cazaba 1 de 52 cuando el azar daria 5%. O sea
   que las senales que yo supuse importantes —consumo bajo, sin seguimiento,
   health score bajo— no discriminan, por sensatas que suenen.

   La razon probable: son UBICUAS. 60 de las 70 cuentas de Claudia estan por
   debajo de HS 60. Una senal que casi todos tienen no puede distinguir a nadie.

   Asi que en vez de inventar otro puntaje, esto mide una cosa por senal:
   **de las cuentas QUE TIENEN la senal, cuantas se fueron; de las que NO la
   tienen, cuantas se fueron.** El cociente entre las dos es el «lift»: cuanto
   multiplica esa senal la probabilidad de baja.

   Lift 1.0 = la senal no dice nada. Lift 2.0 = duplica el riesgo. Lift 0.5 =
   la senal predice lo CONTRARIO de lo que uno creia.

   Corta los datos a una fecha y mide contra lo que paso DESPUES, igual que el
   otro script. Es de SOLO LECTURA.

   USO
   ---
       python scripts/senales-que-predicen.py            # corta en julio
       python scripts/senales-que-predicen.py 2026-06
"""
import io
import json
import os
import re
import sys
import unicodedata
import urllib.request
from collections import defaultdict
from datetime import datetime, timedelta

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', line_buffering=True)
import openpyxl

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CORTE = sys.argv[1] if len(sys.argv) > 1 and not sys.argv[1].startswith('--') else '2026-07'


def norm(s):
    s = ''.join(c for c in unicodedata.normalize('NFD', s or '')
                if unicodedata.category(c) != 'Mn').lower()
    return re.sub(r'[^a-z0-9]', '', s)


def num(v):
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(',', '').strip())
    except (TypeError, ValueError):
        return 0.0


def mes_de(v):
    if isinstance(v, datetime):
        return v.strftime('%Y-%m')
    if isinstance(v, (int, float)):
        return (datetime(1899, 12, 30) + timedelta(days=float(v))).strftime('%Y-%m')
    return str(v or '')[:7]


RX_EXT = re.compile(r'(\d+)\s*(?:extensi[oó]n(?:es)?|ext\b)', re.I)
RX_EXT_ABR = re.compile(r'^(\d+)\s+\S.*\bIL\b', re.I)
RX_SIN_VOZ = re.compile(r'\bchat\b|\bagentes?\s+cp\b|sin\s+saldo|n[uú]meros?\s+virtuales?|whatsapp', re.I)


def base_min(plan, incl):
    n = str(plan or '')
    m = RX_EXT.search(n) or RX_EXT_ABR.search(n)
    ext = int(m.group(1)) if m else None
    if RX_SIN_VOZ.search(n) and not ext and incl < 3:
        return None
    if ext and ext > 0:
        return incl if (incl > 0 and incl / ext >= 50) else ext * 1500
    return incl if incl >= 3 else None


E = {}
for ln in io.open(os.path.join(RAIZ, '.env.local'), encoding='utf-8'):
    if '=' in ln and not ln.strip().startswith('#'):
        k, v = ln.split('=', 1)
        E[k.strip()] = v.strip().strip('"').strip("'")
H = {'apikey': E['SUPABASE_SERVICE_ROLE_KEY'],
     'Authorization': 'Bearer ' + E['SUPABASE_SERVICE_ROLE_KEY']}
B = E['NEXT_PUBLIC_SUPABASE_URL'] + '/rest/v1/'


def sb(t, p):
    f, o = [], 0
    while True:
        d = json.load(urllib.request.urlopen(
            urllib.request.Request(B + t + '?' + p + '&offset=%d&limit=1000' % o, headers=H), timeout=60))
        f += d
        if len(d) < 1000:
            return f
        o += 1000


CU = sb('cuentas', 'select=id,consecutivo,empresa,cid,asesor,estado,health_score,facturacion,tiene_chat_activo,tiene_integracion_api,num_oficinas,contacto_email,contacto_tel')
SEG = sb('seguimientos', 'select=cuenta_id,fecha')
TK = json.load(io.open(os.path.join(RAIZ, 'lib', 'tickets-data.json'), encoding='utf-8'))

wb = openpyxl.load_workbook(os.path.join(RAIZ, 'data', 'cortes-facturacion.xlsx'),
                            data_only=True, read_only=True)
ws = wb[wb.sheetnames[0]]
it = ws.iter_rows(values_only=True)
cab = [str(c).strip() if c is not None else '' for c in next(it)]
idx = {c: i for i, c in enumerate(cab)}
cortes = defaultdict(list)
for r in it:
    f = {c: (r[i] if i < len(r) else None) for c, i in idx.items()}
    cid = str(f.get('CID') or '').strip()
    if cid:
        cortes[cid].append({'mes': mes_de(f.get('Fecha de corte')),
                            'cons': num(f.get('Minutos Consumidos')),
                            'base': base_min(f.get('Nombre del Plan'), num(f.get('Minutos Incluidos')))})
wb.close()
for v in cortes.values():
    v.sort(key=lambda x: x['mes'])

CHURN_MES = defaultdict(set)
s = io.open(os.path.join(RAIZ, 'app', 'churn', 'aaa-grc-data.ts'), encoding='utf-8').read()
for bloque in re.finditer(r"mes:\s*'([^']+)'([\s\S]*?)(?=\n\s*\{\s*mes:|\Z)", s):
    for m in re.finditer(r"cliente:\s*'([^']+)'[^}]*?movimiento:\s*'([^']*)'", bloque.group(2)):
        if 'Churn confirmado' in m.group(2):
            CHURN_MES[bloque.group(1).lower()].add(norm(m.group(1)))

MESES = ['enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio',
         'julio', 'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre']
ncorte = int(CORTE[5:7])
futuro = set()
for i, nm in enumerate(MESES, start=1):
    if i > ncorte:
        futuro |= CHURN_MES.get(nm, set())

meses_hasta = sorted({x['mes'] for v in cortes.values() for x in v if x['mes'] <= CORTE})
ULT = meses_hasta[-1] if meses_hasta else CORTE

# ── rasgos de cada cuenta, con los datos de la fecha de corte ────────
def rasgos(c):
    cid = str(c.get('cid') or '').strip()
    serie = [x for x in cortes.get(cid, []) if x['mes'] <= CORTE]
    r = {}
    uso = None
    if serie and serie[-1]['base']:
        uso = 100 * serie[-1]['cons'] / serie[-1]['base']
    r['sin corte propio'] = not serie
    r['desapareció del corte'] = bool(serie) and serie[-1]['mes'] < ULT
    r['consumo 0%'] = uso is not None and uso < 1
    r['uso < 20%'] = uso is not None and uso < 20
    r['uso < 40%'] = uso is not None and uso < 40
    if len(serie) >= 6:
        a = sum(x['cons'] for x in serie[:3]) / 3
        b = sum(x['cons'] for x in serie[-3:]) / 3
        r['consumo cayó >25%'] = a >= 100 and (b - a) / a <= -0.25
    else:
        r['consumo cayó >25%'] = False

    tot = fal = 0
    for t in TK:
        if (t['cid'] or '').strip() != cid:
            continue
        ap = (t.get('apertura') or '')[:7]
        if ap and ap <= CORTE:
            tot += 1
            if t['es_falla'] == 'Si':
                fal += 1
    r['sin tickets cruzados'] = tot == 0
    r['≥15 tickets'] = tot >= 15
    r['≥40 tickets'] = tot >= 40
    r['≥2 fallas'] = fal >= 2
    r['≥4 fallas'] = fal >= 4
    r['≥8 fallas'] = fal >= 8

    fechas = [str(x.get('fecha') or '')[:10] for x in SEG
              if x.get('cuenta_id') == c['id'] and str(x.get('fecha') or '')[:7] <= CORTE]
    r['sin seguimiento'] = not fechas
    if fechas:
        try:
            d = (datetime.strptime(CORTE + '-28', '%Y-%m-%d')
                 - datetime.strptime(max(fechas), '%Y-%m-%d')).days
            r['>60 días sin contacto'] = d > 60
        except ValueError:
            r['>60 días sin contacto'] = False
    else:
        r['>60 días sin contacto'] = True

    hs = c.get('health_score')
    r['HS < 45'] = isinstance(hs, (int, float)) and hs < 45
    r['HS < 60'] = isinstance(hs, (int, float)) and hs < 60
    fac = float(c.get('facturacion') or 0)
    r['factura < $3,000'] = fac < 3000
    r['factura > $20,000'] = fac > 20000
    r['sin correo ni teléfono'] = not (str(c.get('contacto_email') or '').strip()
                                       or str(c.get('contacto_tel') or '').strip())
    r['tiene chat'] = bool(c.get('tiene_chat_activo'))
    r['tiene API'] = bool(c.get('tiene_integracion_api'))
    return r


datos = [(c, rasgos(c), norm(c['empresa']) in futuro) for c in CU]
base = sum(1 for _, _, y in datos if y)
n = len(datos)
print('cartera %d · se fueron después de %s: %d (%.0f%%)' % (n, CORTE, base, 100.0 * base / n))
print()
print('Lift = cuánto multiplica esa señal la probabilidad de baja. 1.0 = no dice nada.')
print()
print('  %-26s %6s %9s %6s %9s %7s' % ('SEÑAL', 'CON', 'SE FUE', 'SIN', 'SE FUE', 'LIFT'))
filas = []
for k in datos[0][1]:
    con = [y for _, r, y in datos if r[k]]
    sin = [y for _, r, y in datos if not r[k]]
    if len(con) < 8 or len(sin) < 8:
        continue
    pc = sum(con) / len(con)
    ps = sum(sin) / len(sin)
    lift = (pc / ps) if ps > 0 else float('inf')
    filas.append((k, len(con), pc, len(sin), ps, lift))
for k, nc, pc, ns, ps, lift in sorted(filas, key=lambda x: -x[5]):
    marca = ''
    if lift >= 1.5:
        marca = '  <-- SÍ separa'
    elif lift <= 0.67:
        marca = '  <-- separa AL REVÉS'
    print('  %-26s %6d %8.0f%% %6d %8.0f%% %6.2f%s' % (k, nc, 100 * pc, ns, 100 * ps, lift, marca))
print()
print('  Lo que esté cerca de 1.00 no sirve para predecir, por sensato que suene.')
