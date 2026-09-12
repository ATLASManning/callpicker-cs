"""Verifica que destination_type signifique lo que creo, antes de construir nada.

   Hipotesis a probar:
     H1  Una entrante 'Lost' es una llamada que entro y NADIE contesto
         -> deberia tener total_minutes = 0 / NULL casi siempre.
     H2  'Redirected' si se contesto -> minutos > 0.
     H3  'Self_service' es el IVR resolviendo sin agente.
     H4  D2 (Grupo System ooapas) con 99.7% perdidas es real y no un artefacto.
"""
import sys, io, os, collections, datetime
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', line_buffering=True)
import openpyxl

ENT = r"C:\Users\manni\OneDrive\Escritorio\Llamadas_entrantes_Clientes_AAA_Poco_Consumo_Actualizado.xlsx"
SAL = r"C:\Users\manni\OneDrive\Escritorio\Llamadas Salientes Clientes AAA Poco Consumo.xlsx"


def num(v):
    if v is None:
        return None
    s = str(v).strip()
    if s == '' or s.upper() == 'NULL':
        return None
    try:
        return float(s)
    except Exception:
        return None


print('=== H1/H2/H3 · minutos por destination_type (ENTRANTES, muestra 250k) ===')
wb = openpyxl.load_workbook(ENT, data_only=True, read_only=True)
ws = wb[wb.sheetnames[0]]
it = ws.iter_rows(values_only=True)
cab = [str(c).strip() for c in next(it)]
ix = {c: k for k, c in enumerate(cab)}
st = collections.defaultdict(lambda: {'n': 0, 'cero': 0, 'nulo': 0, 'pos': 0, 'suma': 0.0, 'max': 0})
d2 = collections.defaultdict(lambda: collections.Counter())
d2ext = collections.Counter()
d2call = collections.Counter()
n = 0
for r in it:
    n += 1
    if n > 250000:
        break
    t = str(r[ix['destination_type']] or '').strip()
    m = num(r[ix['total_minutes']])
    s = st[t]
    s['n'] += 1
    if m is None:
        s['nulo'] += 1
    elif m == 0:
        s['cero'] += 1
    else:
        s['pos'] += 1; s['suma'] += m; s['max'] = max(s['max'], m)
wb.close()
for t, s in sorted(st.items(), key=lambda x: -x[1]['n']):
    print('  %-14s n=%-8s  NULL=%5.1f%%  cero=%5.1f%%  >0=%5.1f%%  prom(>0)=%.1f  max=%s'
          % (t, format(s['n'], ','), 100*s['nulo']/s['n'], 100*s['cero']/s['n'],
             100*s['pos']/s['n'], s['suma']/max(s['pos'], 1), s['max']))

print()
print('=== H4 · D2 Grupo System ooapas (CID de la cartera) — pasada completa ===')
# D2 -> hay que resolver su CID desde la cartera
import json
cu = json.load(io.open(r'D:\Proyectos\CP\cartera_asesor.json', encoding='utf-8'))
cidD2 = next((str(c['cid']).strip() for c in cu if c['consecutivo'] == 'D2'), None)
print('  CID de D2 = %s' % cidD2)

wb = openpyxl.load_workbook(ENT, data_only=True, read_only=True)
ws = wb[wb.sheetnames[0]]
it = ws.iter_rows(values_only=True)
next(it)
porMes = collections.defaultdict(lambda: collections.Counter())
ext = collections.Counter()
caller = collections.Counter()
minutos = collections.Counter()
for r in it:
    if str(r[ix['customer_id']] or '').strip() != cidD2:
        continue
    t = str(r[ix['destination_type']] or '').strip()
    d = r[ix['date']]
    mes = d.strftime('%Y-%m') if isinstance(d, datetime.datetime) else str(d)[:7]
    porMes[mes][t] += 1
    porMes[mes]['total'] += 1
    e = str(r[ix['destination_data_1']] or '').strip()
    ext[e if e and e.upper() != 'NULL' else '(sin ext)'] += 1
    caller[str(r[ix['caller_id']] or '')[:14]] += 1
    m = num(r[ix['total_minutes']])
    minutos['NULL' if m is None else ('cero' if m == 0 else 'positivo')] += 1
wb.close()
print('  por mes:')
for m in sorted(porMes):
    v = porMes[m]
    print('    %s  total=%-6s Lost=%-6s Redirected=%-5s Self=%-5s Voicemail=%-4s  (%.1f%% perdidas)'
          % (m, v['total'], v['Lost'], v['Redirected'], v['Self_service'], v['Voicemail'],
             100*v['Lost']/max(v['total'], 1)))
print('  minutos: %s' % dict(minutos))
print('  extensiones destino (top 6): %s' % ext.most_common(6))
print('  numeros que llaman distintos: %d · top: %s' % (len(caller), caller.most_common(4)))
