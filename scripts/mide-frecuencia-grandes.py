# -*- coding: utf-8 -*-
u"""La FRECUENCIA con la que se activan las cuentas de mas de $3,500.

Direccion, 14 oct 2026: «construye una grafica para ver su periodo de
frecuencia con el que se activan este tipo de cuentas».

Esto se mide ANTES de escribir el JSX, no despues: una grafica dibujada sobre
una serie que resulta plana, o con mas huecos que datos, es un indicador que
entrena a no mirarlo. Ver [[feedback-kpi-con-fuente]].

Lo que calcula, y por que cada cosa:

  · La serie MENSUAL completa, incluidos los meses en CERO que caen DENTRO de
    la ventana. Un mes sin activaciones grandes dentro del periodo es un cero
    MEDIDO y tiene que pintarse; un mes fuera de la ventana no se midio y no
    existe. Ver [[feedback-cero-sin-medicion]].
  · Las filas SIN MES, que no pueden colocarse en la serie. Se cuentan aparte
    y hay que decirlas con palabras, nunca repartirlas ni esconderlas.
  · El intervalo entre activaciones consecutivas, que es lo que de verdad
    pregunto direccion: cada cuanto cae una de estas.
  · La racha mas larga sin ninguna.

Uso:  python scripts/mide-frecuencia-grandes.py [umbral]
"""
import collections
import io
import os
import sys

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', line_buffering=True)
RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

try:
    from openpyxl import load_workbook
except ImportError:
    sys.stderr.write('Falta openpyxl:  python -m pip install --user openpyxl\n')
    sys.exit(2)

UMBRAL = float(sys.argv[1]) if len(sys.argv) > 1 else 3500.0

MES_NUM = {'January': 1, 'February': 2, 'March': 3, 'April': 4, 'May': 5,
           'June': 6, 'July': 7, 'August': 8, 'September': 9, 'October': 10,
           'November': 11, 'December': 12}
MES_ES = ['', 'Ene', 'Feb', 'Mar', 'Abr', 'May', 'Jun',
          'Jul', 'Ago', 'Sep', 'Oct', 'Nov', 'Dic']


def columna(cab, *nombres):
    """Busca una columna por varios nombres posibles, sin distinguir mayusculas."""
    bajo = {str(c).strip().lower(): i for i, c in enumerate(cab) if c is not None}
    for n in nombres:
        if n.lower() in bajo:
            return bajo[n.lower()]
    for n in nombres:
        for k, i in bajo.items():
            if n.lower() in k:
                return i
    return None


def main():
    p = os.path.join(RAIZ, 'data', 'activaciones.xlsx')
    wb = load_workbook(p, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    filas = ws.iter_rows(values_only=True)
    cab = list(next(filas))
    print(u'  hoja «%s» · %d columnas' % (wb.sheetnames[0], len(cab)))

    c_pago = columna(cab, 'Primer pago', 'primer_pago', 'primerpago', 'pago')
    c_mes = columna(cab, 'Mes', 'mes')
    c_ano = columna(cab, 'Año', 'Ano', 'year', 'anio')
    c_tipo = columna(cab, 'Tipo', 'tipo')
    c_cli = columna(cab, 'Cliente', 'cliente', 'empresa')
    for nom, i in (('pago', c_pago), ('mes', c_mes), ('año', c_ano), ('tipo', c_tipo)):
        if i is None:
            print(u'  *** no encuentro la columna %s. Cabeceras: %s'
                  % (nom, [str(c) for c in cab][:20]))
            return 2

    grandes, sin_mes, base = [], [], 0
    for f in filas:
        if f is None or all(v is None for v in f):
            continue
        tipo = str(f[c_tipo] or '').strip().lower()
        if tipo == 'demo':          # un demo no es una activacion
            continue
        base += 1
        try:
            pago = float(f[c_pago] or 0)
        except (TypeError, ValueError):
            pago = 0.0
        if pago <= UMBRAL:
            continue
        mes = str(f[c_mes] or '').strip()
        try:
            ano = int(f[c_ano] or 0)
        except (TypeError, ValueError):
            ano = 0
        reg = {'cliente': str(f[c_cli] or '') if c_cli is not None else '',
               'pago': pago, 'mes': mes, 'ano': ano}
        if mes in MES_NUM and ano:
            reg['k'] = ano * 12 + MES_NUM[mes]
            grandes.append(reg)
        else:
            sin_mes.append(reg)

    n = len(grandes) + len(sin_mes)
    print(u'\n  universo: %d activaciones (sin demos) · %d pasan de $%s'
          % (base, n, '{:,.0f}'.format(UMBRAL)))
    print(u'  con mes y año: %d · SIN MES O SIN AÑO: %d  <- no caben en la serie'
          % (len(grandes), len(sin_mes)))
    if sin_mes:
        for r in sin_mes[:6]:
            print(u'      %-38s $%-10s mes=%r año=%r'
                  % (r['cliente'][:38], '{:,.0f}'.format(r['pago']), r['mes'], r['ano']))

    if not grandes:
        print(u'  *** ninguna con fecha: no hay serie que dibujar')
        return 1

    por_mes = collections.Counter(r['k'] for r in grandes)
    dinero = collections.defaultdict(float)
    for r in grandes:
        dinero[r['k']] += r['pago']

    k0, k1 = min(por_mes), max(por_mes)
    ventana = list(range(k0, k1 + 1))
    conteos = [por_mes.get(k, 0) for k in ventana]

    def etq(k):
        return u'%s %d' % (MES_ES[(k - 1) % 12 + 1], (k - 1) // 12)

    print(u'\n  VENTANA MEDIDA: %s → %s · %d meses' % (etq(k0), etq(k1), len(ventana)))
    print(u'  (los meses en CERO dentro de esta ventana son ceros MEDIDOS y se pintan;\n'
          u'   fuera de ella no se midio nada y no existen)')

    vacios = sum(1 for c in conteos if c == 0)
    print(u'\n  meses con al menos una: %d de %d (%.0f%%) · meses en cero: %d'
          % (len(ventana) - vacios, len(ventana),
             100.0 * (len(ventana) - vacios) / len(ventana), vacios))
    print(u'  media: %.2f al mes · mediana: %s · maximo en un mes: %d'
          % (sum(conteos) / float(len(ventana)),
             sorted(conteos)[len(conteos) // 2], max(conteos)))

    # ── CADA CUANTO CAE UNA: el intervalo entre meses CON activacion ─────────
    meses_con = sorted(por_mes)
    saltos = [meses_con[i + 1] - meses_con[i] for i in range(len(meses_con) - 1)]
    if saltos:
        ss = sorted(saltos)
        print(u'\n  CADA CUANTO: intervalo entre meses consecutivos CON alguna')
        print(u'    media %.1f meses · mediana %d · maximo %d · minimo %d'
              % (sum(saltos) / float(len(saltos)), ss[len(ss) // 2], max(saltos), min(saltos)))
    print(u'    ritmo global: una cada %.1f dias (%d meses / %d activaciones)'
          % (len(ventana) * 30.44 / len(grandes), len(ventana), len(grandes)))

    # racha mas larga de ceros
    mejor = act = 0
    ini = fin = None
    arr_ini = None
    for i, c in enumerate(conteos):
        if c == 0:
            if act == 0:
                arr_ini = i
            act += 1
            if act > mejor:
                mejor, ini, fin = act, arr_ini, i
        else:
            act = 0
    if mejor:
        print(u'    racha mas larga SIN ninguna: %d meses (%s → %s)'
              % (mejor, etq(ventana[ini]), etq(ventana[fin])))

    print(u'\n  LA SERIE, mes a mes (para saber si hay grafica que dibujar):')
    tope = max(conteos)
    for k, c in zip(ventana, conteos):
        barra = u'█' * int(round(14.0 * c / tope)) if tope else ''
        print(u'    %-9s %2d  %-14s $%s'
              % (etq(k), c, barra, '{:>11,.0f}'.format(dinero.get(k, 0))))

    print(u'\n  VEREDICTO DE DISPERSION: %s'
          % (u'hay serie — varia lo suficiente para que la grafica diga algo'
             if max(conteos) - min(conteos) >= 2 else
             u'*** la serie es casi plana: una grafica de barras no aporta'))
    return 0


if __name__ == '__main__':
    sys.exit(main())
