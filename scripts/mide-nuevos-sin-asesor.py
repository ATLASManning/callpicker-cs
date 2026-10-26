# -*- coding: utf-8 -*-
u"""CLIENTES DE MAS DE $3,500 QUE NADIE ESTA ATENDIENDO.

   Encargo de direccion, 16 oct 2026: «cada que identifiques un nuevo cliente
   que facture arriba de $3,500, sea en Activaciones 2.0 o en facturacion,
   genera una alerta en el Dashboard… con la leyenda: se pide revisar si a
   esta cuenta se le asigna un Asesor».

   DOS FUENTES, Y NO MIDEN LO MISMO
   --------------------------------
   ACTIVACIONES (`data/activaciones.xlsx`) trae el PRIMER PAGO. Es lo que el
   cliente pago al entrar, no lo que factura hoy: de las activaciones grandes
   vivas, treinta facturan hoy mas que su primer pago y diecisiete menos.
   Sirve para detectar al que ACABA de entrar. [[activaciones-fuente]]

   FACTURACION (`data/grc-zoho.json`) trae el MRR institucional, que es la
   cifra buena --`mrrFin`, o `mrrIni` si la linea esta desmentida--. Sirve
   para detectar al que CRECIO hasta pasar el umbral sin que nadie lo note.
   [[grc-fuente-zoho]]

   Las dos cuentan, porque un cliente puede llegar al umbral por cualquiera
   de los dos caminos y en ambos casos hace falta un asesor.

   QUE ES «SIN ASESOR»
   -------------------
   Que no cruza con ninguna cuenta de la cartera, o que cruza y su campo
   `asesor` esta vacio. Las dos cosas significan lo mismo para el negocio:
   nadie responde por ese cliente.

   NO ESCRIBE NADA.

   USO
       python scripts/mide-nuevos-sin-asesor.py
"""
import io
import json
import os
import re
import sys
import urllib.error
import urllib.request

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', line_buffering=True)
RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

UMBRAL = 3500


def entorno():
    d = {}
    with io.open(os.path.join(RAIZ, '.env.local'), encoding='utf-8') as fh:
        for ln in fh:
            ln = ln.strip()
            if ln and not ln.startswith('#') and '=' in ln:
                k, v = ln.split('=', 1)
                d[k.strip()] = v.strip().strip('"').strip("'")
    return d


def supabase(tabla, params):
    E = entorno()
    url = E['NEXT_PUBLIC_SUPABASE_URL'].rstrip('/') + '/rest/v1/' + tabla
    key = E.get('SUPABASE_SERVICE_ROLE_KEY') or E['NEXT_PUBLIC_SUPABASE_ANON_KEY']
    out, desde = [], 0
    while True:
        q = dict(params); q['offset'] = str(desde); q['limit'] = '1000'
        pide = urllib.request.Request(
            url + '?' + '&'.join('%s=%s' % (k, v) for k, v in q.items()),
            headers={'apikey': key, 'Authorization': 'Bearer ' + key})
        lote = json.loads(urllib.request.urlopen(pide, timeout=60).read().decode('utf-8'))
        out.extend(lote)
        if len(lote) < 1000:
            return out
        desde += 1000


def norm(s):
    u"""Para cruzar nombres: sin acentos, sin razon social, sin puntuacion."""
    import unicodedata
    t = unicodedata.normalize('NFD', str(s or ''))
    t = u''.join(c for c in t if unicodedata.category(c) != 'Mn').upper()
    for ruido in (' SA DE CV', ' S A DE C V', ' SAPI', ' SC', ' SRL', ' S DE RL',
                  ' DE CV', ' SA', '.', ',', '-', '  '):
        t = t.replace(ruido, ' ')
    return u' '.join(t.split())


def mrr_institucional(f):
    fin = float(f.get('mrrFin') or 0)
    if fin > 0:
        return fin
    if f.get('verificacion') == 'sigue_viva':
        ini = float(f.get('mrrIni') or 0)
        if ini > 0:
            return ini
    return 0.0


def lineas_de_cuenta():
    u"""La tabla `LINEAS_DE_CUENTA` de lib/nuevos-sin-asesor.ts, LEIDA DE ALLI.

       No se copia: una segunda copia de una tabla curada se desalinea sola, y
       entonces este medidor acusa de huerfano a un CID que el modulo ya
       resolvio. Ya paso con la sonda de activaciones y con la de despliegue:
       la simulacion tiene que seguir al modulo, no correr en paralelo.
    """
    ts = os.path.join(RAIZ, 'lib', 'nuevos-sin-asesor.ts')
    try:
        txt = io.open(ts, encoding='utf-8').read()
    except OSError:
        return {}
    m = re.search(r'LINEAS_DE_CUENTA[^=]*=\s*\{(.*?)\n\}', txt, re.S)
    if not m:
        return {}
    out = {}
    for mm in re.finditer(r"'(\d+)':\s*\{(.*?)\}", m.group(1), re.S):
        cid_padre = re.search(r"cidCuenta:\s*'(\d+)'", mm.group(2))
        if cid_padre:
            out[mm.group(1)] = cid_padre.group(1)
    return out


def main():
    cuentas = supabase('cuentas', {'select': 'id,consecutivo,empresa,cid,asesor,estado'})
    por_cid = {}
    por_nombre = {}
    for c in cuentas:
        k = str(c.get('cid') or '').strip()
        if k:
            por_cid[k] = c
        por_nombre[norm(c.get('empresa'))] = c

    #: Palabras que no distinguen a nadie: casi todas las razones sociales las
    #: llevan, así que casar por ellas junta clientes que no tienen que ver.
    VACIAS = {u'GRUPO', u'SERVICIOS', u'COMERCIALIZADORA', u'CORPORATIVO',
              u'DE', u'DEL', u'LA', u'EL', u'LOS', u'LAS', u'Y', u'MEXICO',
              u'MX', u'INTERNACIONAL', u'NACIONAL', u'SOLUCIONES'}

    def tokens(s):
        return [t for t in norm(s).split() if len(t) > 3 and t not in VACIAS]

    #: Índice por palabra distintiva, para el segundo intento.
    por_token = {}
    for c in cuentas:
        for t in tokens(c.get('empresa')):
            por_token.setdefault(t, []).append(c)


    # La tabla curada, leida de lib/nuevos-sin-asesor.ts (una sola fuente).
    ALIAS = lineas_de_cuenta()

    def asesor_de(cid, nombre):
        u"""Devuelve (cuenta, veredicto).

        TRES INTENTOS, y el tercero existe por FINSUS: en activaciones se
        llama «FINSUS CRÉDITO AL CONSUMO» y en la cartera «Finsus Growth».
        Cruzando el nombre entero sale «no está» y SÍ está, con asesora. Un
        aviso que acusa de huérfana a una cuenta atendida se aprende a
        ignorar, y entonces no sirve para la que sí lo está.

        Por eso el tercer intento no concluye «está»: concluye «se parece a»,
        que es una cosa distinta y la pantalla lo dice así.
        """
        k = str(cid or '').strip()
        # Antes de nada: ¿es una LINEA de una cuenta que ya existe?
        padre = ALIAS.get(k)
        if padre and padre in por_cid:
            c = por_cid[padre]
            a = str(c.get('asesor') or '').strip()
            return c, (u'sin asesor' if not a else a + u'  (linea de %s)' % str(c.get('consecutivo')))
        if k and k in por_cid:
            c = por_cid[k]
            a = str(c.get('asesor') or '').strip()
            return c, (u'sin asesor' if not a else a)
        c = por_nombre.get(norm(nombre))
        if c:
            a = str(c.get('asesor') or '').strip()
            return c, (u'sin asesor' if not a else a)
        for t in tokens(nombre):
            cand = por_token.get(t)
            if cand:
                return cand[0], u'¿es %s?' % str(cand[0].get('empresa'))[:26]
        return None, u'no está en la cartera'

    print(u'\n  CLIENTES DE MAS DE $%s SIN ASESOR' % format(UMBRAL, ',d'))
    print(u'  ' + u'=' * 72)

    # ── FACTURACION (GRC) ──────────────────────────────────────────────────
    J = json.loads(io.open(os.path.join(RAIZ, 'data', 'grc-zoho.json'),
                           encoding='utf-8').read())
    porCid = {}
    for f in J.get('filas', []):
        cid = str(f.get('cid') or '').strip()
        if not cid:
            continue
        r = porCid.setdefault(cid, {'mrr': 0.0, 'cliente': f.get('cliente') or f.get('empresa')})
        r['mrr'] += mrr_institucional(f)

    huerfanos_grc = []
    for cid, r in porCid.items():
        if r['mrr'] < UMBRAL:
            continue
        c, quien = asesor_de(cid, r['cliente'])
        if quien in (u'sin asesor', u'no está en la cartera'):
            huerfanos_grc.append((cid, r['cliente'], r['mrr'], quien))
    huerfanos_grc.sort(key=lambda x: -x[2])

    print(u'\n  (1) FACTURACION — MRR institucional del GRC')
    print(u'      %d CID(s) por encima del umbral · %d sin asesor'
          % (sum(1 for r in porCid.values() if r['mrr'] >= UMBRAL), len(huerfanos_grc)))
    for cid, cli, mrr, quien in huerfanos_grc[:14]:
        print(u'        %-9s %-38s $%9s  %s'
              % (cid, str(cli)[:38], format(int(round(mrr)), ',d'), quien))
    if len(huerfanos_grc) > 14:
        print(u'        … y %d más' % (len(huerfanos_grc) - 14))

    # ── ACTIVACIONES ───────────────────────────────────────────────────────
    print(u'\n  (2) ACTIVACIONES 2.0 — primer pago')
    ruta = os.path.join(RAIZ, 'data', 'activaciones.xlsx')
    if not os.path.exists(ruta):
        print(u'      (no está data/activaciones.xlsx)')
    else:
        try:
            from openpyxl import load_workbook
        except ImportError:
            print(u'      (falta openpyxl para leer el .xlsx)')
            load_workbook = None
        if load_workbook:
            wb = load_workbook(ruta, data_only=True, read_only=True)
            ws = wb[wb.sheetnames[0]]
            filas = ws.iter_rows(values_only=True)
            cab = [str(x or '').strip() for x in next(filas)]
            def col(*nombres):
                for n in nombres:
                    if n in cab:
                        return cab.index(n)
                return None
            i_cli = col('Cliente', 'CLIENTE', 'Empresa')
            # EN ESTE ARCHIVO LA COLUMNA SE LLAMA «ID», Y ES EL CID. Buscando
            # solo «CID» no la encontraba, pasaba None y cruzaba por nombre —
            # el mismo fallo que tenia el modulo en TS. [[activaciones-fuente]]
            i_cid = col('ID', 'CID', 'Cid', 'cid')
            i_pago = col('1er Pago', 'Primer Pago', 'Facturan')
            if i_cli is None or i_pago is None:
                print(u'      *** no encuentro las columnas: %s' % cab[:12])
            else:
                huer, dudosos = [], []
                n_grandes = 0
                for r in filas:
                    try:
                        monto = float(r[i_pago] or 0)
                    except (TypeError, ValueError):
                        continue
                    if monto < UMBRAL:
                        continue
                    n_grandes += 1
                    cli = r[i_cli]
                    cid = r[i_cid] if i_cid is not None else None
                    c, quien = asesor_de(cid, cli)
                    if quien in (u'sin asesor', u'no está en la cartera'):
                        huer.append((cid, cli, monto, quien))
                    elif quien.startswith(u'¿es'):
                        dudosos.append((cid, cli, monto, quien))
                huer.sort(key=lambda x: -x[2])
                print(u'      %d activación(es) por encima del umbral · %d SIN CUENTA · %d por confirmar'
                      % (n_grandes, len(huer), len(dudosos)))
                for cid, cli, monto, quien in huer[:14]:
                    print(u'        %-9s %-38s $%9s  %s'
                          % (str(cid or '—'), str(cli)[:38],
                             format(int(round(monto)), ',d'), quien))
                if len(huer) > 14:
                    print(u'        … y %d más' % (len(huer) - 14))

    print(u'\n  ' + u'=' * 72)
    print(u'  Esto es lo que la alerta del Dashboard tendría que mostrar hoy.\n')


if __name__ == '__main__':
    main()
