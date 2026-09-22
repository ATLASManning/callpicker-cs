# -*- coding: utf-8 -*-
"""Concilia el estado de la cartera contra el estado REAL de Callpicker.

   INSTRUCCION DE DIRECCION (22 sep 2026)
   --------------------------------------
   «En el apartado Cuentas existe el combo activas y dormidas. Aplica una
   conciliacion con Churn hasta el mes de julio, y si hay cuentas como Holton
   que esten en activas e incluso tengan actividad SAC, colocalas en dormidas.
   Todo lo que tenga etiqueta 9 o diga cancelado deberas, si existe en cuentas,
   colocarlo en Dormida.»

   QUE HACE
   --------
   Lee el export de estado de Callpicker —una fila por cliente con
   `customer_state_id` y `Estatus`— y mueve a DORMIDA (`hibernacion`) toda
   cuenta de la cartera que alla este CANCELADA y aqui siga viva.

   QUE NO HACE, Y ES A PROPOSITO
   -----------------------------
   1. NO toca las que estan en Soft o Hard Suspend. Direccion pidio etiqueta 9
      o «cancelado»; una suspension no es una baja y moverla seria decidir por
      ella. Se listan aparte para que alguien lo decida.
   2. NO le quita la cuenta al asesor. A una cuenta muerta se le MARCA, no se
      le saca de la cartera: si desaparece, nadie responde por ella.
   3. NO borra sus actividades SAC abiertas. Holton tiene tres. Cerrarlas en
      silencio escondería que se trabajo sobre una cuenta ya cancelada, que es
      justo lo que esta conciliacion viene a destapar.

   Cada cambio deja nota fechada en Observaciones KAM, con el mismo formato que
   `notaReclasificacion` de lib/conciliacion.ts.

   USO
   ---
       python scripts/concilia-estado-callpicker.py <archivo.xlsx>
       python scripts/concilia-estado-callpicker.py <archivo.xlsx> --aplicar
"""
import io
import json
import os
import re
import sys
import urllib.request
from collections import Counter, defaultdict
from datetime import date

import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', line_buffering=True)

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(RAIZ, 'scripts'))
from observaciones_kam import anteponer_entrada   # noqa: E402  (helper local)

ESTADO_DORMIDA = 'hibernacion'
ESTADOS_VIVOS = ('activo', 'en_riesgo')
# customer_state_id de Callpicker. 9 es la baja; 6 y 7 son suspensiones.
STATE_CANCELADO = '9'
STATE_SUSPEND = ('6', '7')

APLICAR = '--aplicar' in sys.argv
ARCHIVOS = [a for a in sys.argv[1:] if not a.startswith('--')]
if not ARCHIVOS:
    raise SystemExit('Uso: python scripts/concilia-estado-callpicker.py <archivo.xlsx> [--aplicar]')
SRC = ARCHIVOS[0]

_env = {}
for _ln in io.open(os.path.join(RAIZ, '.env.local'), encoding='utf-8'):
    _ln = _ln.strip()
    if _ln and not _ln.startswith('#') and '=' in _ln:
        _k, _v = _ln.split('=', 1)
        _env[_k.strip()] = _v.strip().strip('"').strip("'")
URL = _env['NEXT_PUBLIC_SUPABASE_URL'].rstrip('/')
KEY = _env['SUPABASE_SERVICE_ROLE_KEY']
H = {'apikey': KEY, 'Authorization': 'Bearer %s' % KEY,
     'Content-Type': 'application/json', 'Prefer': 'return=representation'}


def get(p):
    with urllib.request.urlopen(urllib.request.Request(URL + p, headers=H), timeout=60) as r:
        return json.loads(r.read().decode('utf-8'))


def patch(p, body):
    req = urllib.request.Request(URL + p, data=json.dumps(body).encode('utf-8'),
                                 headers=H, method='PATCH')
    with urllib.request.urlopen(req, timeout=60) as r:
        return json.loads(r.read().decode('utf-8'))


def norm(v):
    s = str(v if v is not None else '').strip()
    if s.endswith('.0'):
        s = s[:-2]
    return '' if s in ('0', 'None', 'nan') else s


# ══════════════════════════════════════════════════════════════════════════
#  El estado real
# ══════════════════════════════════════════════════════════════════════════
wb = openpyxl.load_workbook(SRC, data_only=True, read_only=True)
hoja = wb.sheetnames[0]
ws = wb[hoja]
_it = ws.iter_rows(values_only=True)
cab = [str(c).strip() if c is not None else '' for c in next(_it)]
for _c in ('id', 'company_name', 'customer_state_id', 'Estatus'):
    if _c not in cab:
        wb.close()
        raise SystemExit('El archivo no trae la columna «%s». Trae: %s' % (_c, cab))
I = dict((c, i) for i, c in enumerate(cab))
real = {}
for r in _it:
    cid = norm(r[I['id']])
    if cid:
        real[cid] = {'nombre': str(r[I['company_name']] or '').strip(),
                     'state': norm(r[I['customer_state_id']]),
                     'estatus': str(r[I['Estatus']] or '').strip()}
wb.close()

print('=== ESTADO REAL DE CALLPICKER ===')
print('  archivo: %s (hoja «%s»)' % (os.path.basename(SRC), hoja))
print('  clientes: %d · %s' % (len(real),
                               ', '.join('%s=%d' % kv for kv in
                                         Counter(v['estatus'] for v in real.values()).most_common())))

cuentas = get('/rest/v1/cuentas?select=id,consecutivo,cid,empresa,estado,asesor,facturacion,observaciones_kam&limit=500')
por_cid = dict((norm(c['cid']), c) for c in cuentas if norm(c['cid']))
print('\n  cartera: %d cuentas · %d con CID · %d cruzan con el archivo'
      % (len(cuentas), len(por_cid), sum(1 for k in real if k in por_cid)))

# ══════════════════════════════════════════════════════════════════════════
#  Quien se mueve
# ══════════════════════════════════════════════════════════════════════════
mover, suspend, ya_dormidas = [], [], []
for cid, r in real.items():
    c = por_cid.get(cid)
    if not c:
        continue
    cancelada = r['state'] == STATE_CANCELADO or r['estatus'].lower() == 'cancelado'
    if cancelada and c['estado'] in ESTADOS_VIVOS:
        mover.append((c, r))
    elif cancelada:
        ya_dormidas.append((c, r))
    elif r['state'] in STATE_SUSPEND and c['estado'] in ESTADOS_VIVOS:
        suspend.append((c, r))

print('\n' + '=' * 74)
print('A DORMIDA — canceladas en Callpicker y vivas aqui: %d' % len(mover))
print('=' * 74)
print('  %-34s %-8s %-10s %-9s %9s' % ('cuenta', 'CID', 'hoy', 'asesor', 'factura'))
for c, r in sorted(mover, key=lambda x: -(x[0].get('facturacion') or 0)):
    print('  %-34s %-8s %-10s %-9s %9s'
          % (str(c['empresa'])[:34], norm(c['cid']), c['estado'],
             (c.get('asesor') or '—')[:9], format(int(c.get('facturacion') or 0), ',')))
print('\n  (ya estaban en Dormida o cancelado: %d)' % len(ya_dormidas))

print('\n' + '=' * 74)
print('NO SE TOCAN — suspendidas, que no es lo mismo que canceladas: %d' % len(suspend))
print('=' * 74)
for c, r in suspend:
    print('  %-34s %-8s Callpicker: %-14s aqui: %s'
          % (str(c['empresa'])[:34], norm(c['cid']), r['estatus'], c['estado']))
if suspend:
    print('\n  Direccion pidio etiqueta 9 o «cancelado». Una suspension no es una')
    print('  baja: se listan para que alguien decida, no se mueven solas.')

# ── Actividades SAC abiertas sobre lo que se va a mover ──────────────────
print('\n=== ACTIVIDADES SAC SOBRE LAS QUE SE MUEVEN ===')
for c, r in mover:
    act = get('/rest/v1/actividades?select=id,tipo,estado&cuenta_id=eq.%s&limit=100' % c['id'])
    abiertas = [a for a in act if str(a.get('estado') or '').lower() not in ('completada', 'cancelada')]
    print('  %-34s %2d actividad(es), %d abierta(s)%s'
          % (str(c['empresa'])[:34], len(act), len(abiertas),
             '  <-- se trabajo sobre una cuenta ya cancelada' if abiertas else ''))
print('\n  NO se cierran ni se borran: cerrarlas en silencio escondería el hallazgo.')

if not APLICAR:
    print('\n  (diagnostico; no se toco nada)')
    print('  Para aplicar:  python scripts/concilia-estado-callpicker.py "%s" --aplicar'
          % os.path.basename(SRC))
    raise SystemExit(0)

# ══════════════════════════════════════════════════════════════════════════
#  Aplicar
# ══════════════════════════════════════════════════════════════════════════
print('\n=== APLICANDO ===')
hoy = date.today().isoformat()
for c, r in mover:
    motivo = ('Callpicker la reporta como %s (customer_state_id %s) en el corte del %s, '
              'y aqui seguia como «%s». Pasa a Dormida. El asesor la conserva en cartera: '
              'a una cuenta muerta se le marca, no se le quita.'
              % (r['estatus'], r['state'], hoy, c['estado']))
    nota = '[Conciliación %s] %s' % (hoy, motivo)
    res = patch('/rest/v1/cuentas?id=eq.%s' % c['id'], {
        'estado': ESTADO_DORMIDA,
        'observaciones_kam': anteponer_entrada(c.get('observaciones_kam'), nota, 'sistema'),
    })
    assert len(res) == 1, 'se tocaron %d filas y esperaba 1' % len(res)
    print('  %-34s %s -> %s' % (str(c['empresa'])[:34], c['estado'], ESTADO_DORMIDA))

# ── Releer y comprobar ───────────────────────────────────────────────────
print('\n=== COMPROBANDO, releyendo de la base ===')
ok = True
for c, r in mover:
    f = get('/rest/v1/cuentas?id=eq.%s&select=empresa,estado,asesor' % c['id'])[0]
    bien = f['estado'] == ESTADO_DORMIDA and f.get('asesor') == c.get('asesor')
    ok = ok and bien
    print('  %-34s %-12s asesor %-9s %s'
          % (str(f['empresa'])[:34], f['estado'], f.get('asesor') or '—',
             'OK' if bien else '*** NO QUEDO'))

todas = get('/rest/v1/cuentas?select=estado&limit=500')
print('\n=== LA CARTERA DESPUES ===')
for k, v in Counter(x['estado'] for x in todas).most_common():
    print('  %-14s %3d' % (k, v))
print('  %-14s %3d' % ('TOTAL', len(todas)))
assert len(todas) == len(cuentas), 'cambio el numero de cuentas'
if not ok:
    raise SystemExit('  *** Algo no quedo bien.')
print('\n  Conciliacion aplicada. Ninguna cuenta perdio a su asesor.')
