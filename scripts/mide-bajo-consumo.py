# -*- coding: utf-8 -*-
"""Cuantas cuentas de la cartera estan en bajo consumo, y de quien son.

   POR QUE EXISTE
   --------------
   Direccion amplio la instruccion del 25 sep 2026: ademas de las cuentas con
   auditoria en riesgo, tambien las de BAJO CONSUMO deben generar actividad SAC.

   Antes de escribir esa regla hay que saber el TAMANO. Las auditadas en riesgo
   son 22 y por eso caben como acervo de una sola vez. Si las de bajo consumo
   fueran cien, el mismo diseno las convertiria en una carga que nadie puede
   trabajar — y una regla que no se puede cumplir se ignora entera, incluida la
   parte que si importaba.

   Es la regla de la casa: medir la dispersion ANTES de escribir el codigo.

   Y la otra mitad, que es donde ya se tropezo este tablero: un corte sin base
   de minutos medible NO es un corte al 0%. «Sin medicion» y «sin consumo» son
   cosas distintas y aqui se separan. Ver lib/plan-minutos.ts.

   Es de SOLO LECTURA.
"""
import io
import json
import os
import re
import sys
import unicodedata
import urllib.request
from collections import Counter, defaultdict
from datetime import datetime, timedelta

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', line_buffering=True)
import openpyxl

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# Replica de lib/plan-minutos.ts — mismo orden, mismos umbrales.
MINUTOS_POR_EXTENSION = 1500
MIN_POR_EXT_PLAUSIBLE = 50
BOLSA_MINIMA = 3
RX_EXT = re.compile(r'(\d+)\s*(?:extensi[oó]n(?:es)?|ext\b)', re.I)
RX_EXT_ABR = re.compile(r'^(\d+)\s+\S.*\bIL\b', re.I)
RX_SIN_VOZ = re.compile(r'\bchat\b|\bagentes?\s+cp\b|sin\s+saldo|n[uú]meros?\s+virtuales?|whatsapp', re.I)


def base_minutos(plan, incl):
    nombre = str(plan or '')
    m = RX_EXT.search(nombre) or RX_EXT_ABR.search(nombre)
    ext = int(m.group(1)) if m else None
    if RX_SIN_VOZ.search(nombre) and not ext and incl < BOLSA_MINIMA:
        return None
    if ext and ext > 0:
        return incl if (incl > 0 and incl / ext >= MIN_POR_EXT_PLAUSIBLE) else ext * MINUTOS_POR_EXTENSION
    return incl if incl >= BOLSA_MINIMA else None


def num(v):
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(',', '').strip())
    except (TypeError, ValueError):
        return 0.0


def mes_de(v):
    if isinstance(v, datetime):
        return v.strftime('%Y-%m')
    if isinstance(v, (int, float)):
        return (datetime(1899, 12, 30) + timedelta(days=float(v))).strftime('%Y-%m')
    return str(v or '')[:7]


def norm(s):
    s = ''.join(c for c in unicodedata.normalize('NFD', s or '')
                if unicodedata.category(c) != 'Mn').lower()
    return re.sub(r'[^a-z0-9]', '', s)


# ── cartera ──────────────────────────────────────────────────────────
E = {}
for ln in io.open(os.path.join(RAIZ, '.env.local'), encoding='utf-8'):
    if '=' in ln and not ln.strip().startswith('#'):
        k, v = ln.split('=', 1)
        E[k.strip()] = v.strip().strip('"').strip("'")
req = urllib.request.Request(
    E['NEXT_PUBLIC_SUPABASE_URL'] +
    '/rest/v1/cuentas?select=id,consecutivo,empresa,cid,asesor,estado,health_score,facturacion&limit=2000',
    headers={'apikey': E['SUPABASE_SERVICE_ROLE_KEY'],
             'Authorization': 'Bearer ' + E['SUPABASE_SERVICE_ROLE_KEY']})
CUENTAS = json.load(urllib.request.urlopen(req, timeout=40))
POR_CID = {str(c.get('cid') or '').strip(): c for c in CUENTAS if c.get('cid')}

# ── cortes ───────────────────────────────────────────────────────────
wb = openpyxl.load_workbook(os.path.join(RAIZ, 'data', 'cortes-facturacion.xlsx'),
                            data_only=True, read_only=True)
ws = wb[wb.sheetnames[0]]
it = ws.iter_rows(values_only=True)
cab = [str(c).strip() if c is not None else '' for c in next(it)]
idx = {c: i for i, c in enumerate(cab)}
filas = [{c: (r[i] if i < len(r) else None) for c, i in idx.items()}
         for r in it if any(x is not None and str(x).strip() for x in r)]
wb.close()

meses = sorted({mes_de(f.get('Fecha de corte')) for f in filas if mes_de(f.get('Fecha de corte'))})
ULT = meses[-1]
print('corte más reciente: %s  ·  %d filas en ese mes'
      % (ULT, sum(1 for f in filas if mes_de(f.get('Fecha de corte')) == ULT)))
print()

# ── evaluar SOLO las cuentas de la cartera ───────────────────────────
estado = {}
for f in filas:
    if mes_de(f.get('Fecha de corte')) != ULT:
        continue
    cid = str(f.get('CID') or '').strip()
    c = POR_CID.get(cid)
    if not c:
        continue
    plan = str(f.get('Nombre del Plan') or '').strip()
    incl = num(f.get('Minutos Incluidos'))
    cons = num(f.get('Minutos Consumidos'))
    base = base_minutos(plan, incl)
    pct = (100.0 * cons / base) if base else None
    prev = estado.get(c['id'])
    # Una cuenta puede traer varias lineas de servicio: se queda la de MAYOR
    # consumo. Quedarse con la primera acusaria de «sin consumo» a una cuenta
    # que si usa el servicio por otra linea.
    if prev is None or (pct is not None and (prev['pct'] is None or pct > prev['pct'])):
        estado[c['id']] = {'cuenta': c, 'plan': plan, 'incl': incl, 'cons': cons,
                           'base': base, 'pct': pct}

sin_corte = [c for c in CUENTAS if c['id'] not in estado]
medibles = [v for v in estado.values() if v['pct'] is not None]
sin_med = [v for v in estado.values() if v['pct'] is None]

print('=== LA CARTERA CONTRA EL CORTE DE %s ===' % ULT)
print('  cuentas en cartera         %d' % len(CUENTAS))
print('  con corte este mes         %d' % len(estado))
print('  sin corte este mes         %d   (no se puede afirmar nada de su consumo)' % len(sin_corte))
print('  con base medible           %d' % len(medibles))
print('  SIN medición               %d   (plan sin bolsa ni extensiones: NO son 0%%)' % len(sin_med))
print()

BANDAS = [(0, 0.0001, 'cero absoluto'), (0.0001, 10, 'bajo 10%'),
          (10, 20, '10–20%'), (20, 40, '20–40%'), (40, 80, '40–80%'), (80, 1e9, 'más de 80%')]
print('=== REPARTO POR CONSUMO (solo las %d medibles) ===' % len(medibles))
for lo, hi, etq in BANDAS:
    g = [v for v in medibles if lo <= (v['pct'] or 0) < hi]
    print('  %-16s %4d  %s' % (etq, len(g), '█' * min(60, len(g))))

print()
for umbral in (0.0001, 10, 20):
    g = [v for v in medibles if (v['pct'] or 0) < umbral]
    porases = Counter((v['cuenta'].get('asesor') or '[sin asesor]') for v in g)
    etq = 'exactamente 0%' if umbral < 1 else 'bajo %d%%' % umbral
    print('  Si el umbral fuera «%s»: %d cuenta(s) -> %s'
          % (etq, len(g), dict(porases.most_common())))

# ── solapamiento con las auditadas en riesgo ─────────────────────────
aud = set()
import glob
for f in glob.glob(os.path.join(RAIZ, 'app', 'auditoria', '*-data.ts')):
    s = io.open(f, encoding='utf-8').read()
    e = re.search(r"^  estado:\s+'([^']+)'", s, re.M)
    n = re.search(r"^  nombre:\s+'([^']+)'", s, re.M)
    if e and n and e.group(1) in ('en_riesgo', 'rescatable', 'en_recuperacion'):
        aud.add(norm(n.group(1)))

print()
print('=== SOLAPAMIENTO con las auditadas en riesgo ===')
for umbral in (0.0001, 10, 20):
    g = [v for v in medibles if (v['pct'] or 0) < umbral]
    ya = sum(1 for v in g if norm(v['cuenta']['empresa']) in aud)
    etq = 'exactamente 0%' if umbral < 1 else 'bajo %d%%' % umbral
    print('  «%s»: %d de %d ya tienen auditoría (no habría que duplicarles tarea)'
          % (etq, ya, len(g)))

print()
print('=== LAS DE CONSUMO CERO, CON NOMBRE ===')
cero = sorted([v for v in medibles if (v['pct'] or 0) < 0.0001],
              key=lambda v: -float(v['cuenta'].get('facturacion') or 0))
print('  %-5s %-30s %-9s %-12s %5s %10s  %s'
      % ('#', 'CUENTA', 'ASESOR', 'ESTADO', 'HS', 'FACT', 'PLAN'))
for v in cero[:30]:
    c = v['cuenta']
    print('  %-5s %-30s %-9s %-12s %5s %10s  %s'
          % (c.get('consecutivo') or '?', c['empresa'][:30], str(c.get('asesor'))[:9],
             c.get('estado'), c.get('health_score'),
             format(int(float(c.get('facturacion') or 0)), ','), v['plan'][:34]))
if len(cero) > 30:
    print('  ... y %d más' % (len(cero) - 30))
