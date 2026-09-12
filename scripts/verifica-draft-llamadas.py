"""Verifica las afirmaciones del draft contra los archivos. Nada se presenta sin esto.

   La compuerta a probar (reemplaza al filtro de palabras clave que marco a
   «CLAUDIA ROMAN» como agente virtual):
       un destino queda POR CONFIRMAR  <=>  contestadas == 0  Y  minutos == 0
                                            Y concentra >= 60% de lo no contestado
                                            Y tiene >= 200 no contestadas
   Es estructural: no mira el texto del nombre, mira si ese destino sostuvo
   alguna vez una conversacion.
"""
import sys, io, json, collections, datetime
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', line_buffering=True)
import openpyxl

ENT = r"C:\Users\manni\OneDrive\Escritorio\Llamadas_entrantes_Clientes_AAA_Poco_Consumo_Actualizado.xlsx"
SAL = r"C:\Users\manni\OneDrive\Escritorio\Llamadas Salientes Clientes AAA Poco Consumo.xlsx"
cu = json.load(io.open(r'D:\Proyectos\CP\cartera_asesor.json', encoding='utf-8'))
NOM = {str(c.get('cid') or '').strip(): c for c in cu if str(c.get('cid') or '').strip()}


def mins(v):
    if v is None:
        return 0.0
    s = str(v).strip()
    if s == '' or s.upper() == 'NULL':
        return 0.0
    try:
        return float(s)
    except Exception:
        return 0.0


dest = collections.defaultdict(lambda: {'lost': 0, 'cont': 0, 'min': 0.0, 'nums': set()})
porCid = collections.defaultdict(lambda: {'lost': 0, 'tot': 0, 'meses': collections.defaultdict(collections.Counter),
                                          'ultima': None, 'sinNum': 0})
wb = openpyxl.load_workbook(ENT, data_only=True, read_only=True)
ws = wb[wb.sheetnames[0]]
it = ws.iter_rows(values_only=True)
cab = [str(c).strip() for c in next(it)]
ix = {c: k for k, c in enumerate(cab)}
for r in it:
    cid = str(r[ix['customer_id']] or '').strip()
    t = str(r[ix['destination_type']] or '').strip()
    d = str(r[ix['destination_data_1']] or '').strip()
    d = d if d and d.upper() != 'NULL' else '(sin destino registrado)'
    m = mins(r[ix['total_minutes']])
    f = r[ix['date']]
    k = (cid, d)
    dest[k]['min'] += m
    if t == 'Lost':
        dest[k]['lost'] += 1
        n = str(r[ix['caller_id']] or '').strip()
        if n and n.upper() != 'NULL':
            dest[k]['nums'].add(n)
        else:
            porCid[cid]['sinNum'] += 1
    else:
        dest[k]['cont'] += 1
    c = porCid[cid]
    c['tot'] += 1
    if t == 'Lost':
        c['lost'] += 1
    if isinstance(f, datetime.datetime):
        mes = f.strftime('%Y-%m')
        c['meses'][mes][t] += 1
        c['meses'][mes]['total'] += 1
        iso = f.strftime('%Y-%m-%d')
        if c['ultima'] is None or iso > c['ultima']:
            c['ultima'] = iso
wb.close()

lostPorCid = collections.Counter()
for (cid, d), v in dest.items():
    lostPorCid[cid] += v['lost']

print('=' * 84)
print('PRUEBA 1 · LA COMPUERTA vs EL FILTRO POR CONCENTRACION SOLA')
print('=' * 84)
concentra, compuerta = [], []
for (cid, d), v in dest.items():
    tot = lostPorCid[cid]
    if v['lost'] < 200 or tot == 0:
        continue
    pct = 100 * v['lost'] / tot
    if pct >= 50:
        concentra.append((cid, d, v, pct))
    if d != '(sin destino registrado)' and pct >= 60 and v['cont'] == 0 and v['min'] == 0:
        compuerta.append((cid, d, v, pct))

print('\nA) Solo por concentracion (>=50%% de lo no contestado, >=200): %d destinos' % len(concentra))
print('   %-5s %-30s %-8s %-8s %-9s %s' % ('CONS', 'DESTINO', 'sin cont', 'CONTESTA', 'minutos', '%conc'))
for cid, d, v, pct in sorted(concentra, key=lambda x: -x[3]):
    c = NOM.get(cid)
    print('   %-5s %-30s %-8s %-8s %-9s %5.1f%%  %s' % (
        (c['consecutivo'] if c else '?'), d[:30], v['lost'], v['cont'], round(v['min']), pct,
        '<-- FALSO POSITIVO' if v['cont'] > 0 else ''))
fp = sum(1 for _, _, v, _ in concentra if v['cont'] > 0)
print('   => %d de %d serian FALSOS POSITIVOS (el destino si contesta llamadas)' % (fp, len(concentra)))

print('\nB) Con la compuerta (contestadas==0 Y minutos==0 Y >=60%% Y >=200): %d destinos' % len(compuerta))
for cid, d, v, pct in compuerta:
    c = NOM.get(cid)
    print('   %-5s %-30s sin contestar=%-6s contestadas=%-4s minutos=%-4s conc=%.1f%% numeros=%d' % (
        (c['consecutivo'] if c else '?'), d[:30], v['lost'], v['cont'], round(v['min']), pct, len(v['nums'])))
print('   => falsos positivos: %d' % sum(1 for _, _, v, _ in compuerta if v['cont'] > 0))

print()
print('=' * 84)
print('PRUEBA 2 · CIFRAS CONCRETAS QUE EL DRAFT PUBLICA')
print('=' * 84)


def datos(cons):
    cid = next((k for k, v in NOM.items() if v['consecutivo'] == cons), None)
    return cid, porCid.get(cid)


for cons in ['C68', 'D2', 'F59', 'C32', 'D16', 'C46']:
    cid, c = datos(cons)
    if not c:
        print('  %-5s sin registro' % cons); continue
    emp = NOM[cid]['empresa'][:28]
    ago = c['meses'].get('2026-08', {})
    at, al = ago.get('total', 0), ago.get('Lost', 0)
    prev = [m for m in c['meses'] if '2026-01' <= m <= '2026-07']
    bt = sum(c['meses'][m]['total'] for m in prev)
    bl = sum(c['meses'][m]['Lost'] for m in prev)
    print('  %-5s %-28s ago: %s de %s = %s · base ene-jul: %s de %s = %s · ultima: %s · meses con dato: %d' % (
        cons, emp, al, at, ('%.1f%%' % (100*al/at)) if at else '—',
        bl, bt, ('%.1f%%' % (100*bl/bt)) if bt else '—', c['ultima'], len(c['meses'])))

print()
print('  cuentas con <30 entrantes en TODA la ventana: %s'
      % [NOM[c]['consecutivo'] for c in porCid if c in NOM and porCid[c]['tot'] < 30])
pares = sum(1 for c in porCid for m in porCid[c]['meses'] if porCid[c]['meses'][m]['total'] < 30)
print('  pares cuenta-mes con <30 entrantes: %d' % pares)
print('  cuentas SIN los 9 meses: %d' % sum(1 for c in porCid if len(porCid[c]['meses']) < 9))

print()
print('=' * 84)
print('PRUEBA 3 · LA MARCA DE PRIORIDAD DE AGOSTO (cuantas cuentas prende)')
print('=' * 84)
print('  regla: agosto >=300 entrantes Y base ene-jul >=300 Y agosto sube >=5 pts sobre su base')
marca = []
for cid, c in porCid.items():
    if cid not in NOM:
        continue
    ago = c['meses'].get('2026-08', {})
    at, al = ago.get('total', 0), ago.get('Lost', 0)
    prev = [m for m in c['meses'] if '2026-01' <= m <= '2026-07']
    bt = sum(c['meses'][m]['total'] for m in prev)
    bl = sum(c['meses'][m]['Lost'] for m in prev)
    if at < 300 or bt < 300:
        continue
    delta = 100*al/at - 100*bl/bt
    if delta >= 5:
        marca.append((NOM[cid]['consecutivo'], NOM[cid]['empresa'][:26], NOM[cid]['asesor'],
                      at, al, 100*al/at, 100*bl/bt, delta))
for f in sorted(marca, key=lambda x: -x[7]):
    print('  %-5s %-26s %-8s ago %s/%s = %.1f%% · base %.1f%% · %+.1f pts' % (
        f[0], f[1], f[2], f[4], f[3], f[5], f[6], f[7]))
print('  => prenden %d cuentas' % len(marca))
porAsesor = collections.Counter(f[2] for f in marca)
print('  => por asesor: %s' % dict(porAsesor))

print()
print('=' * 84)
print('PRUEBA 4 · SALIENTES: cuantas cuentas tienen entrantes pero NO salientes')
print('=' * 84)
wb = openpyxl.load_workbook(SAL, data_only=True, read_only=True)
ws = wb[wb.sheetnames[0]]
it = ws.iter_rows(values_only=True)
cab2 = [str(c).strip() for c in next(it)]
ix2 = {c: k for k, c in enumerate(cab2)}
cidsSal = set()
for r in it:
    cidsSal.add(str(r[ix2['customer_id']] or '').strip())
wb.close()
solo = [c for c in porCid if c not in cidsSal]
print('  CIDs entrantes: %d · CIDs salientes: %d' % (len(porCid), len(cidsSal)))
print('  con entrantes y SIN salientes: %d -> %s' % (
    len(solo), [NOM[c]['consecutivo'] for c in solo if c in NOM]))
print('  con salientes y SIN entrantes: %d' % len([c for c in cidsSal if c not in porCid]))

print()
print('  llenado de caller_id: cuentas con >5%% de Lost sin numero: %s'
      % [NOM[c]['consecutivo'] for c in porCid
         if c in NOM and porCid[c]['lost'] > 0 and porCid[c]['sinNum']/porCid[c]['lost'] > 0.05])
