# -*- coding: utf-8 -*-
"""Comprueba si las correcciones pendientes ya se hicieron en la hoja de origen.

   POR QUE EXISTE
   --------------
   Jose Manuel, 25 sep 2026: «En la conciliacion del siguiente viernes debe estar
   corregido esto».

   Un pendiente que vive en la cabeza de alguien se pierde. Y uno que vive en un
   documento que nadie abre, tambien. Este lo comprueba contra el ARCHIVO, dice
   si ya quedo, y `carga-semanal.py` lo corre al final — asi que es imposible
   hacer la carga del viernes sin ver el estado.

   LAS CORRECCIONES VAN EN LA HOJA DE ORIGEN, NO AQUI. Parchear el .xlsx del
   repo no sirve para nada: el siguiente export lo sobreescribe. Y corregir la
   captura de otra persona sin que lo autorice no es nuestro trabajo.

   Sale con codigo 1 si hay alguna VENCIDA (pasada su fecha comprometida).

   USO
   ---
       python scripts/revisa-correcciones.py
"""
import io
import json
import os
import sys
from datetime import datetime, timedelta, timezone

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', line_buffering=True)
import openpyxl

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ORIGEN = r'D:\Archivos'
FICHERO = os.path.join(RAIZ, 'data', 'correcciones-pendientes.json')

# Se puede apuntar a otro fichero para probarlo. Un comprobador que solo se ha
# visto decir «pendiente» no esta probado: hay que verlo decir «corregida»
# tambien, o no se sabe si esa rama funciona.
if len(sys.argv) > 1 and not sys.argv[1].startswith('--'):
    FICHERO = sys.argv[1]


def hoy_en_mexico():
    return (datetime.now(timezone.utc) - timedelta(hours=6)).date()


_cache = {}


def filas_de(nombre):
    """Lee el archivo de origen. Si no esta, cae al que se cargo en el repo."""
    if nombre in _cache:
        return _cache[nombre]
    rutas = [os.path.join(ORIGEN, nombre),
             os.path.join(RAIZ, 'data', 'activaciones.xlsx')]
    for ruta in rutas:
        if not os.path.exists(ruta):
            continue
        wb = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
        ws = wb[wb.sheetnames[0]]
        it = ws.iter_rows(values_only=True)
        try:
            cab = [str(c).strip() if c is not None else '' for c in next(it)]
        except StopIteration:
            wb.close()
            continue
        idx = {c: i for i, c in enumerate(cab)}
        filas = [{c: (r[i] if i < len(r) else None) for c, i in idx.items()}
                 for r in it if any(x is not None and str(x).strip() for x in r)]
        wb.close()
        _cache[nombre] = (filas, ruta)
        return _cache[nombre]
    _cache[nombre] = (None, None)
    return _cache[nombre]


def texto(v):
    return str(v).strip() if v is not None else ''


def comprueba(p):
    """Devuelve (estado, detalle). Estado: corregida | pendiente | pregunta | sin_archivo."""
    c = p.get('comprobacion') or {}
    tipo = c.get('tipo')
    if tipo == 'pregunta':
        return 'pregunta', 'necesita una respuesta, no se puede comprobar sola'

    filas, ruta = filas_de(p['fuente'])
    if filas is None:
        return 'sin_archivo', 'no encontre «%s» ni en %s ni en data/' % (p['fuente'], ORIGEN)

    col_id = c.get('columna_id', 'ID')

    if tipo == 'id_exacto':
        ids = {texto(f.get(col_id)) for f in filas}
        bien = c['debe_existir'] in ids
        mal = c['no_debe_existir'] in ids
        if bien and not mal:
            return 'corregida', 'el ID %s existe y ya no está «%s»' % (c['debe_existir'], c['no_debe_existir'])
        partes = []
        if not bien:
            partes.append('no encuentro el ID %s' % c['debe_existir'])
        if mal:
            partes.append('sigue estando «%s»' % c['no_debe_existir'])
        return 'pendiente', '; '.join(partes)

    if tipo == 'campo_en_rango':
        objetivo = [f for f in filas if texto(f.get(col_id)) == c['valor_id']]
        if not objetivo:
            return 'pendiente', 'no encuentro la fila con %s = %s' % (col_id, c['valor_id'])
        v = objetivo[0].get(c['columna'])
        try:
            n = float(v)
        except (TypeError, ValueError):
            return 'pendiente', '«%s» trae %r, que no es un número' % (c['columna'], v)
        if c['min'] <= n <= c['max']:
            return 'corregida', '«%s» = %s' % (c['columna'], int(n))
        return 'pendiente', '«%s» sigue en %s (se espera entre %s y %s)' % (
            c['columna'], int(n), c['min'], c['max'])

    return 'pendiente', 'tipo de comprobación desconocido: %r' % tipo


def main():
    if not os.path.exists(FICHERO):
        print('  no hay correcciones pendientes registradas.')
        return 0
    d = json.load(io.open(FICHERO, encoding='utf-8'))
    pend = d.get('pendientes') or []
    if not pend:
        print('  no hay correcciones pendientes.')
        return 0

    hoy = hoy_en_mexico()
    print('=== CORRECCIONES PENDIENTES EN LA HOJA DE ORIGEN ===')
    print()
    vencidas = corregidas = 0
    for p in pend:
        estado, detalle = comprueba(p)
        try:
            lim = datetime.strptime(p['comprometida'], '%Y-%m-%d').date()
            dias = (hoy - lim).days
        except (KeyError, ValueError):
            lim, dias = None, None

        if estado == 'corregida':
            marca = 'YA QUEDÓ'
            corregidas += 1
        elif estado == 'pregunta':
            marca = 'ESPERA RESPUESTA'
        elif estado == 'sin_archivo':
            marca = '*** NO PUEDO COMPROBARLO'
        elif dias is not None and dias > 0:
            marca = '*** VENCIDA hace %d día(s)' % dias
            vencidas += 1
        else:
            marca = 'pendiente%s' % ('' if dias is None else ' (para el %s)' % p['comprometida'])

        print('  [%s]  %s' % (marca, p['id']))
        print('      fuente : %s · captura: %s' % (p['fuente'], p.get('captura', '—')))
        print('      qué    : %s' % p['que_pasa'])
        print('      espera : %s' % p['que_se_espera'])
        print('      ahora  : %s' % detalle)
        print()

    print('  %d de %d corregida(s).' % (corregidas, len(pend)))
    if corregidas:
        print('  Las que ya quedaron se quitan de data/correcciones-pendientes.json.')
    if vencidas:
        print('  *** %d vencida(s): pasó la fecha y siguen igual.' % vencidas)
        return 1
    return 0


if __name__ == '__main__':
    sys.exit(main())
