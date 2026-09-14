"""Genera data/analisis-llamadas.json para el apartado «Análisis de Llamadas».

   Es una agregación MÁS RICA que la de las fichas de cuenta. Allá se guardan
   solo los marginales de día y de hora —y se descartó el mapa de calor a
   propósito, porque con una cuenta sola la rejilla de 168 celdas repetía la
   misma silueta gastando 156px—. Aquí sí se cruza: con 148 cuentas y 4.2
   millones de llamadas el patrón día×hora es el hallazgo, no el adorno.

   Qué emite, por cuenta y en global:
     · serie por mes con los cuatro desenlaces
     · matriz día de la semana × hora (total y no contestadas)
     · serie diaria, para ver la tendencia fina y los huecos del archivo
     · destinos con contestadas, minutos y números distintos
     · ventana real por dirección

   Vive en data/ y lo lee el API del lado del servidor: 1.5 MB no tienen por
   qué viajar al navegador. Mismo patrón que data/cortes-facturacion.xlsx.

   Las seis fuentes y sus trampas están documentadas en
   scripts/gen-llamadas-data.py; aquí se reusan las mismas decisiones:
   la Parte 2 de salientes es copia de la Parte 1 y no se lee, las columnas se
   resuelven por NOMBRE, y las filas sin columna de destino van a un bucket
   propio que jamás se mezcla con «(sin destino registrado)».
"""
import sys, io, os, re, json, datetime, collections
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', line_buffering=True)
import openpyxl

ARCH = r"D:\Archivos"
SALIDA = r"D:\Windows\Projects\callpicker-cs\data\analisis-llamadas.json"

FUENTES = [
    ('Llamadas_entrantes_Clientes_AAA_Poco_Consumo_Actualizado.xlsx', 'ent', '0-40'),
    ('Llamadas Salientes Clientes AAA Poco Consumo.xlsx',             'sal', '0-40'),
    ('Entrantes Mayor consumo 40 Parte 1.xlsx',                       'ent', '40+'),
    ('Entrantes Mayor consumo 40 Parte 2.xlsx',                       'ent', '40+'),
    ('Entrantes Mayor consumo 40 Parte 3.xlsx',                       'ent', '40+'),
    ('Salientes Mayor consumo 40 Parte 1.xlsx',                       'sal', '40+'),
]

SIN_DESTINO = '(sin destino registrado)'
NO_EXPORTADO = '(destino no venía en el archivo)'
ALIAS = {
    'tipo': ['destination_type'], 'fecha': ['date'], 'cid': ['customer_id'],
    'empresa': ['empresa', 'Nombre Empresa', 'Nombre de Empresa'],
    'destino': ['destination_data_1'], 'caller': ['caller_id'], 'minutos': ['total_minutes'],
}
# Los cuatro desenlaces de una entrante, en el orden en que se apilan.
TIPOS_ENT = ['Redirected', 'Self_service', 'Voicemail', 'Lost']


def cid_de(v):
    if v is None:
        return ''
    if isinstance(v, float) and float(v).is_integer():
        return str(int(v))
    s = str(v).strip()
    return '' if s.upper() == 'NULL' else s


def texto(v):
    s = str(v or '').strip()
    return '' if s.upper() == 'NULL' else s


def mins(v):
    s = str(v or '').strip()
    if not s or s.upper() == 'NULL':
        return 0.0
    try:
        return float(s)
    except Exception:
        return 0.0


def nueva():
    return {
        'empresa': collections.Counter(), 'corte': '',
        'ent': {'tipos': collections.Counter(),
                'meses': collections.defaultdict(collections.Counter),
                'dh': collections.Counter(), 'dhL': collections.Counter(),
                'dia': collections.Counter(), 'diaL': collections.Counter(),
                'dest': collections.defaultdict(lambda: {'l': 0, 'c': 0, 'm': 0.0, 'n': set()}),
                'desde': None, 'hasta': None, 'sinCol': 0},
        'sal': {'tipos': collections.Counter(),
                'meses': collections.defaultdict(collections.Counter),
                'dh': collections.Counter(), 'dhL': collections.Counter(),
                'dia': collections.Counter(), 'diaL': collections.Counter(),
                'desde': None, 'hasta': None},
    }


D = collections.defaultdict(nueva)

for arch, dire, corte in FUENTES:
    ruta = os.path.join(ARCH, arch)
    if not os.path.exists(ruta):
        print('!! FALTA %s' % arch)
        continue
    if arch == 'Salientes Mayor consumo 40 Parte 2.xlsx':
        print('-- omitida (copia exacta de la Parte 1) %s' % arch)
        continue
    wb = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    cab = [str(c).strip() if c is not None else '' for c in next(it)]
    ix = {k: next((cab.index(n) for n in ns if n in cab), None) for k, ns in ALIAS.items()}
    print('%-46s %s %s destino=%s' % (arch[:46], dire, corte, 'si' if ix['destino'] is not None else 'NO'))
    n = 0
    for r in it:
        n += 1
        if n % 500000 == 0:
            print('    ... %s' % format(n, ','))
        cid = cid_de(r[ix['cid']])
        if not cid:
            continue
        v = D[cid]
        v['corte'] = v['corte'] or corte
        if ix['empresa'] is not None:
            e = texto(r[ix['empresa']])
            if e:
                v['empresa'][e] += 1
        t = texto(r[ix['tipo']]) or 'Desconocido'
        f = r[ix['fecha']]
        d = v[dire]
        d['tipos'][t] += 1
        esPerdida = (t == 'Lost') if dire == 'ent' else (t in ('Lost', 'Lost_by_agent'))

        if isinstance(f, datetime.datetime):
            iso, mes = f.strftime('%Y-%m-%d'), f.strftime('%Y-%m')
            d['meses'][mes][t] += 1
            d['meses'][mes]['total'] += 1
            k = f.weekday() * 24 + f.hour
            d['dh'][k] += 1
            d['dia'][iso] += 1
            if esPerdida:
                d['dhL'][k] += 1
                d['diaL'][iso] += 1
            if d['desde'] is None or iso < d['desde']:
                d['desde'] = iso
            if d['hasta'] is None or iso > d['hasta']:
                d['hasta'] = iso

        if dire == 'ent':
            if ix['destino'] is not None:
                dest = texto(r[ix['destino']]) or SIN_DESTINO
            else:
                dest = NO_EXPORTADO
                d['sinCol'] += 1
            dd = d['dest'][dest]
            dd['m'] += mins(r[ix['minutos']]) if ix['minutos'] is not None else 0.0
            if esPerdida:
                dd['l'] += 1
                num = texto(r[ix['caller']]) if ix['caller'] is not None else ''
                if num:
                    dd['n'].add(num)
            else:
                dd['c'] += 1
    wb.close()
    print('    %s filas' % format(n, ','))

# ── Volcado ─────────────────────────────────────────────────────────────────
def matriz(cnt):
    """168 celdas (7 días × 24 horas) en una lista plana. 0 donde no hubo nada."""
    return [cnt.get(i, 0) for i in range(168)]


cuentas = {}
for cid, v in D.items():
    e, s = v['ent'], v['sal']
    dest = sorted(v['ent']['dest'].items(), key=lambda x: -x[1]['l'])[:14]
    cuentas[cid] = {
        'cid': cid,
        'empresa': v['empresa'].most_common(1)[0][0] if v['empresa'] else '',
        'corte': v['corte'],
        'ent': {
            'total': sum(e['tipos'].values()), 'tipos': dict(e['tipos']),
            'meses': {m: dict(c) for m, c in sorted(e['meses'].items())},
            'dh': matriz(e['dh']), 'dhL': matriz(e['dhL']),
            'dia': dict(sorted(e['dia'].items())), 'diaL': dict(sorted(e['diaL'].items())),
            'dest': [{'d': d, 'l': x['l'], 'c': x['c'], 'min': round(x['m']), 'n': len(x['n'])}
                     for d, x in dest],
            'desde': e['desde'], 'hasta': e['hasta'], 'sinCol': e['sinCol'],
        } if e['tipos'] else None,
        'sal': {
            'total': sum(s['tipos'].values()), 'tipos': dict(s['tipos']),
            'meses': {m: dict(c) for m, c in sorted(s['meses'].items())},
            'dh': matriz(s['dh']), 'dhL': matriz(s['dhL']),
            'dia': dict(sorted(s['dia'].items())), 'diaL': dict(sorted(s['diaL'].items())),
            'desde': s['desde'], 'hasta': s['hasta'],
        } if s['tipos'] else None,
    }

todas = [x for c in cuentas.values() for d in (c['ent'], c['sal']) if d for x in (d['desde'], d['hasta']) if x]
meses = sorted({m for c in cuentas.values() for d in (c['ent'], c['sal']) if d for m in d['meses']})
META = {
    'corte': max(todas), 'desde': min(todas), 'meses': meses, 'cuentas': len(cuentas),
    'entTotal': sum(c['ent']['total'] for c in cuentas.values() if c['ent']),
    'salTotal': sum(c['sal']['total'] for c in cuentas.values() if c['sal']),
    'entLost': sum(c['ent']['tipos'].get('Lost', 0) for c in cuentas.values() if c['ent']),
    'salNoCon': sum(c['sal']['tipos'].get('Lost', 0) + c['sal']['tipos'].get('Lost_by_agent', 0)
                    for c in cuentas.values() if c['sal']),
    'sinCol': sum(c['ent']['sinCol'] for c in cuentas.values() if c['ent']),
    'tiposEnt': TIPOS_ENT,
    'generado': max(todas),
}

os.makedirs(os.path.dirname(SALIDA), exist_ok=True)
with io.open(SALIDA, 'w', encoding='utf-8') as f:
    json.dump({'meta': META, 'cuentas': cuentas}, f, ensure_ascii=False, separators=(',', ':'))

print()
print('escrito %s (%.1f MB)' % (SALIDA, os.path.getsize(SALIDA) / 1048576))
print('  cuentas   : %d' % META['cuentas'])
print('  ventana   : %s → %s (%d meses)' % (META['desde'], META['corte'], len(meses)))
print('  entrantes : %s · sin contestar %s (%.1f%%)'
      % (format(META['entTotal'], ','), format(META['entLost'], ','),
         100 * META['entLost'] / max(META['entTotal'], 1)))
print('  salientes : %s · no conecto %s (%.1f%%)'
      % (format(META['salTotal'], ','), format(META['salNoCon'], ','),
         100 * META['salNoCon'] / max(META['salTotal'], 1)))
print('  entrantes sin columna de destino: %s (%.1f%%)'
      % (format(META['sinCol'], ','), 100 * META['sinCol'] / max(META['entTotal'], 1)))
