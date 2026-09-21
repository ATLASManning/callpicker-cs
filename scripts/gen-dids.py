# -*- coding: utf-8 -*-
"""Genera `data/dids.json` desde el export «DIDs en Callpicker.xlsx».

   QUE ES UN DID
   -------------
   Los numeros telefonicos que Callpicker le entrega a cada cliente. Una cuenta
   suele tener VARIOS: de las 183 cuentas de la cartera que traen numero, 146
   (el 80%) tienen mas de uno.

   LAS TRES COLUMNAS DEL EXPORT
   ----------------------------
     NUM CALLPICKER  el numero
     CUENTA          la ETIQUETA de ese numero, no el nombre del cliente
     CID             el identificador del cliente

   OJO CON «CUENTA». No es la razon social: es como el cliente bautizo ESE
   numero —una sede, un area, una persona, una campania—. El CID 7 trae nueve
   etiquetas distintas («Citi Banamex», «Concilia», «Queretaro», «eTicket»…) y
   es un solo cliente. Pasa en 660 de los 3,717 CIDs.

   Instruccion de direccion (21 sep 2026): **el CID es el que dice a que
   cliente corresponde, y la etiqueta se muestra TAL CUAL viene** — el cliente
   puede nombrar sus numeros como le convenga por estrategia. Aqui no se
   renombra, no se normaliza y no se deduplica por nombre.

   LO UNICO QUE SI SE TOCA: LA CODIFICACION
   ----------------------------------------
   883 etiquetas (5.3%) llegan con el texto roto —«Lourdes Villase√±or»,
   «Cr√©dito y Cobranza»—: son UTF-8 leidos como Mac Roman en algun punto de la
   cadena de exportacion. Eso NO es un nombre distinto, es un defecto de
   codificacion, y se repara restaurando lo que la etiqueta de verdad dice.

   El arreglo es conservador por construccion: solo se aplica cuando el viaje
   de ida y vuelta cuadra. Comprobado contra las 16,534 etiquetas reales —883
   reparadas, 0 restos, y **0 etiquetas sanas tocadas**, que es la direccion
   que de verdad importa.

   USO
   ---
       python scripts/gen-dids.py                       # usa D:\\Archivos
       python scripts/gen-dids.py "ruta\\al\\archivo.xlsx"
"""
import io
import json
import os
import sys
from collections import Counter, defaultdict

import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', line_buffering=True)

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ORIGEN = sys.argv[1] if len(sys.argv) > 1 else r'D:\Archivos\DIDs en Callpicker.xlsx'
DESTINO = os.path.join(RAIZ, 'data', 'dids.json')

COL_NUM, COL_ETQ, COL_CID = 'NUM CALLPICKER', 'CUENTA', 'CID'


def repara(s):
    """UTF-8 leido como Mac Roman -> texto real. Solo si el viaje cuadra."""
    if not s:
        return s
    try:
        arreglado = s.encode('mac_roman').decode('utf-8')
    except (UnicodeEncodeError, UnicodeDecodeError):
        return s          # no era este mojibake: NO se toca
    return arreglado


def norm(v):
    s = str(v if v is not None else '').strip()
    return s[:-2] if s.endswith('.0') else s


if not os.path.exists(ORIGEN):
    raise SystemExit('No existe el archivo: %s' % ORIGEN)

wb = openpyxl.load_workbook(ORIGEN, data_only=True, read_only=True)
ws = wb[wb.sheetnames[0]]
it = ws.iter_rows(values_only=True)
cab = [str(c).strip() if c is not None else '' for c in next(it)]
for c in (COL_NUM, COL_ETQ, COL_CID):
    if c not in cab:
        wb.close()
        raise SystemExit('Falta la columna «%s». El archivo trae: %s' % (c, cab))
i_num, i_etq, i_cid = cab.index(COL_NUM), cab.index(COL_ETQ), cab.index(COL_CID)

por_cid = defaultdict(list)
total = reparadas = sin_cid = 0
vistos = set()
duplicados = 0
largos = Counter()

for r in it:
    num = norm(r[i_num])
    cid = norm(r[i_cid])
    etq_cruda = str(r[i_etq] or '').strip()
    etq = repara(etq_cruda)
    if etq != etq_cruda:
        reparadas += 1
    if not num:
        continue
    total += 1
    largos[len(num)] += 1
    if num in vistos:
        duplicados += 1
    vistos.add(num)
    # Un CID vacio o «0» no identifica a nadie: se cuenta y se deja fuera en
    # vez de colgarlo de una cuenta inventada.
    if not cid or cid == '0':
        sin_cid += 1
        continue
    por_cid[cid].append({'n': num, 'e': etq})

wb.close()

# Las tablas tienen que cerrar: lo agrupado + lo sin CID == lo leido.
agrupados = sum(len(v) for v in por_cid.values())
assert agrupados + sin_cid == total, \
    'no cierra: %d agrupados + %d sin CID != %d leidos' % (agrupados, sin_cid, total)

salida = {
    'totalNumeros': total,
    'totalCids': len(por_cid),
    'sinCid': sin_cid,
    'etiquetasReparadas': reparadas,
    'porCid': dict(por_cid),
}
os.makedirs(os.path.dirname(DESTINO), exist_ok=True)
io.open(DESTINO, 'w', encoding='utf-8').write(
    json.dumps(salida, ensure_ascii=False, separators=(',', ':')))

print('=== DIDs ===')
print('  origen:   %s' % ORIGEN)
print('  destino:  %s  (%s KB)'
      % (os.path.relpath(DESTINO, RAIZ), format(os.path.getsize(DESTINO) // 1024, ',')))
print()
print('  numeros leidos:        %s' % format(total, ','))
print('  agrupados por CID:     %s  en %s CIDs' % (format(agrupados, ','), format(len(por_cid), ',')))
print('  sin CID utilizable:    %s' % format(sin_cid, ','))
print('  (agrupados + sin CID == leidos  OK)')
print('  etiquetas reparadas:   %s' % format(reparadas, ','))
print('  numeros repetidos:     %s' % format(duplicados, ','))
print()
print('  largo del numero:')
for L in sorted(largos):
    print('    %2d digitos: %6s%s' % (L, format(largos[L], ','),
                                      '' if L == 12 else '   <- no es mexicano de 10 digitos + 52'))
