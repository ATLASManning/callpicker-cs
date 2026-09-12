"""Pasada completa sobre los dos archivos de llamadas para conocer el terreno
   ANTES de disenar el modulo. No escribe nada en la app: solo agrega y reporta.

   Salida: D:\\Proyectos\\CP\\llamadas_agregado.json
"""
import sys, io, os, json, datetime, collections
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', line_buffering=True)
import openpyxl

ARCH = {
    'salientes': r"C:\Users\manni\OneDrive\Escritorio\Llamadas Salientes Clientes AAA Poco Consumo.xlsx",
    'entrantes': r"C:\Users\manni\OneDrive\Escritorio\Llamadas_entrantes_Clientes_AAA_Poco_Consumo_Actualizado.xlsx",
}
DIAS = ['lunes', 'martes', 'miércoles', 'jueves', 'viernes', 'sábado', 'domingo']


def fecha(v):
    """Las celdas vienen como datetime o como texto ISO. Devuelve datetime o None."""
    if isinstance(v, datetime.datetime):
        return v
    if isinstance(v, datetime.date):
        return datetime.datetime(v.year, v.month, v.day)
    if isinstance(v, (int, float)):
        try:
            return datetime.datetime(1899, 12, 30) + datetime.timedelta(days=float(v))
        except Exception:
            return None
    s = str(v or '').strip()
    if not s or s == 'NULL':
        return None
    for f in ('%Y-%m-%d %H:%M:%S', '%Y-%m-%d %H:%M', '%Y-%m-%d'):
        try:
            return datetime.datetime.strptime(s[:19], f)
        except Exception:
            pass
    return None


def numero(v):
    if v is None:
        return 0.0
    s = str(v).strip()
    if s == '' or s.upper() == 'NULL':
        return 0.0
    try:
        return float(s)
    except Exception:
        return 0.0


def cid_de(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    s = str(v).strip()
    return '' if s.upper() == 'NULL' else s


def nuevo():
    return {
        'total': 0, 'minutos': 0.0,
        'tipos': collections.Counter(),          # destination_type crudo
        'meses': collections.defaultdict(lambda: collections.Counter()),
        'dow': collections.Counter(),            # dia de la semana (solo perdidas + total)
        'dowPerdidas': collections.Counter(),
        'hora': collections.Counter(),
        'horaPerdidas': collections.Counter(),
        'ext': collections.Counter(),
        'extPerdidas': collections.Counter(),
        'primera': None, 'ultima': None,
        'empresa': None,
    }


resumen = {}
for etq, ruta in ARCH.items():
    print('=' * 78)
    print('%s · %s MB' % (etq.upper(), round(os.path.getsize(ruta) / 1048576, 1)))
    wb = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    cab = [str(c).strip() if c is not None else '' for c in next(it)]
    ix = {c: k for k, c in enumerate(cab)}
    print('  columnas: %s' % ', '.join(cab))

    porCid = collections.defaultdict(nuevo)
    tiposGlobal = collections.Counter()
    n, sinFecha, sinCid = 0, 0, 0
    iDate, iCid, iTipo = ix.get('date'), ix.get('customer_id'), ix.get('destination_type')
    iMin, iExt = ix.get('total_minutes'), ix.get('destination_data_1')
    iEmp = ix.get('empresa')

    for r in it:
        n += 1
        if n % 200000 == 0:
            print('    ... %s filas' % format(n, ','))
        cid = cid_de(r[iCid]) if iCid is not None else ''
        if not cid:
            sinCid += 1
            continue
        d = fecha(r[iDate]) if iDate is not None else None
        tipo = str(r[iTipo] or '').strip() if iTipo is not None else ''
        tiposGlobal[tipo or '(vacío)'] += 1
        c = porCid[cid]
        c['total'] += 1
        c['minutos'] += numero(r[iMin]) if iMin is not None else 0.0
        c['tipos'][tipo or '(vacío)'] += 1
        if iEmp is not None and r[iEmp] and not c['empresa']:
            c['empresa'] = str(r[iEmp]).strip()
        if iExt is not None and r[iExt]:
            e = str(r[iExt]).strip()
            if e and e.upper() != 'NULL':
                c['ext'][e] += 1
                if tipo == 'Lost':
                    c['extPerdidas'][e] += 1
        if d is None:
            sinFecha += 1
            continue
        mes = d.strftime('%Y-%m')
        c['meses'][mes]['total'] += 1
        c['meses'][mes][tipo or '(vacío)'] += 1
        c['meses'][mes]['minutos'] += numero(r[iMin]) if iMin is not None else 0
        c['dow'][d.weekday()] += 1
        c['hora'][d.hour] += 1
        if tipo == 'Lost':
            c['dowPerdidas'][d.weekday()] += 1
            c['horaPerdidas'][d.hour] += 1
        iso = d.strftime('%Y-%m-%d %H:%M')
        if c['primera'] is None or iso < c['primera']:
            c['primera'] = iso
        if c['ultima'] is None or iso > c['ultima']:
            c['ultima'] = iso
    wb.close()

    print('  filas: %s · sin CID: %s · sin fecha: %s · CIDs distintos: %s'
          % (format(n, ','), format(sinCid, ','), format(sinFecha, ','), format(len(porCid), ',')))
    print('  destination_type:')
    for t, k in tiposGlobal.most_common():
        print('    %-16s %10s  (%.1f%%)' % (t, format(k, ','), 100 * k / max(n, 1)))

    # serializar
    ser = {}
    for cid, c in porCid.items():
        ser[cid] = {
            'total': c['total'], 'minutos': round(c['minutos'], 1),
            'tipos': dict(c['tipos']),
            'meses': {m: dict(v) for m, v in c['meses'].items()},
            'dow': {str(k): v for k, v in c['dow'].items()},
            'dowPerdidas': {str(k): v for k, v in c['dowPerdidas'].items()},
            'hora': {str(k): v for k, v in c['hora'].items()},
            'horaPerdidas': {str(k): v for k, v in c['horaPerdidas'].items()},
            'ext': dict(c['ext'].most_common(15)),
            'extPerdidas': dict(c['extPerdidas'].most_common(15)),
            'primera': c['primera'], 'ultima': c['ultima'], 'empresa': c['empresa'],
        }
    resumen[etq] = {'filas': n, 'tipos': dict(tiposGlobal), 'porCid': ser}
    print()

os.makedirs(r"D:\Proyectos\CP", exist_ok=True)
salida = r"D:\Proyectos\CP\llamadas_agregado.json"
with io.open(salida, 'w', encoding='utf-8') as f:
    json.dump(resumen, f, ensure_ascii=False)
print('guardado: %s (%.1f MB)' % (salida, os.path.getsize(salida) / 1048576))
