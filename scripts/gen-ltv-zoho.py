"""Convierte el export del detalle de Zoho en data/ltv-zoho.json.

   SUSTITUYE LA FUENTE DE LTV. La vista que alimentaba el modulo traia datos
   que direccion considera incorrectos; esta la reemplaza. El origen es el
   detalle del tablero GRC: clic derecho sobre el MRR inicio del mes ->
   «Ver datos subyacentes» -> Mas -> «Exportar Vista».

   LAS DOCE COLUMNAS que fijo direccion, y que este archivo exige:
     Cliente · clasificacion_cliente · Facturas_2026 · Meses Activo ·
     Importe Acumulado Recurrente · MRR Inicio Contrato (BCY) ·
     MRR Fin Contrato (BCY) · Ingreso Ganado Contrato (BCY) · Movimiento MRR ·
     Ingreso Perdido Contrato (BCY) Real ·
     Ingreso Perdido Contrato (BCY) Fraude-Reestructura · Rango MRR Fin Contrato

   ── LO QUE ESTE GENERADOR NO PUEDE DEJAR PASAR ─────────────────────────────
   El corte de septiembre 2026 marca 1,085 clientes como «Churn confirmado», y
   las 1,085 filas comparten un patron que las delata: TODAS traen MRR Fin = 0
   y en TODAS la perdida es exactamente igual al MRR inicio. Ninguna parcial.
   El export se tomo el 17 de septiembre, con el mes a la mitad.
   No es churn: es «todavia no se ha facturado este mes».

   La prueba: de esas 1,085, sesenta y cuatro son cuentas de la cartera, y
   sesenta y una siguen ACTIVAS o EN RIESGO en la propia base del tablero. La
   antiguedad promedio de las 64 es de 53 meses; una de ellas lleva 155.

   Por eso cada fila sale marcada con `provisional`: es «Churn confirmado» y su
   cuenta sigue viva en la base. Publicar eso como baja confirmada seria una
   alarma falsa de $546,811 sobre cuentas que no se han ido.

   Uso:  python scripts/gen-ltv-zoho.py "C:/ruta/Septiembre.xlsx"
"""
import sys, io, os, re, json, unicodedata, collections, urllib.request
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', line_buffering=True)
import openpyxl

SALIDA = r'D:\Windows\Projects\callpicker-cs\data\ltv-zoho.json'

COLS = {
    'cliente':    'Cliente',
    'clasif':     'clasificacion_cliente',
    'facturas':   'Facturas_2026',
    'meses':      'Meses Activo',
    'acumulado':  'Importe Acumulado Recurrente',
    'mrrIni':     'MRR Inicio Contrato (BCY)',
    'mrrFin':     'MRR Fin Contrato (BCY)',
    'ganado':     'Ingreso Ganado Contrato (BCY)',
    'movimiento': 'Movimiento MRR',
    'perdidaReal': 'Ingreso Perdido Contrato (BCY) Real',
    'perdidaFraude': 'Ingreso Perdido Contrato (BCY) Fraude-Reestructura',
    'rango':      'Rango MRR Fin Contrato',
    'mes':        'Mes Nombre',
}


def num(v):
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r'[^0-9.\-]', '', str(v or ''))
    try:
        return float(s) if s not in ('', '-', '.') else 0.0
    except ValueError:
        return 0.0


def norm(s):
    s = ''.join(c for c in unicodedata.normalize('NFD', str(s or '')).lower()
                if not (0x300 <= ord(c) <= 0x36f))
    return re.sub(r'[^a-z0-9]', '', s)


if len(sys.argv) < 2:
    print(__doc__)
    raise SystemExit('Falta la ruta del export.')
ORIGEN = sys.argv[1]
if not os.path.exists(ORIGEN):
    raise SystemExit('No existe: %s' % ORIGEN)

wb = openpyxl.load_workbook(ORIGEN, data_only=True, read_only=True)
ws = wb[wb.sheetnames[0]]
it = ws.iter_rows(values_only=True)
cab = [str(c).strip() if c is not None else '' for c in next(it)]
ix = {k: (cab.index(v) if v in cab else None) for k, v in COLS.items()}
faltan = [COLS[k] for k, v in ix.items() if v is None and k != 'mes']
if faltan:
    raise SystemExit('Al export le faltan columnas obligatorias: %s' % faltan)
crudas = [r for r in it]
wb.close()
g = lambda r, k: r[ix[k]] if ix[k] is not None else None

# ── La cartera, para poder distinguir una baja de un «aun no facturado» ────
env = {}
for l in io.open(r'D:\Windows\Projects\callpicker-cs\.env.local', encoding='utf-8'):
    l = l.strip()
    if '=' in l and not l.startswith('#'):
        k, v = l.split('=', 1); env[k.strip()] = v.strip().strip('"')
K = env['SUPABASE_SERVICE_ROLE_KEY']
rq = urllib.request.Request(
    env['NEXT_PUBLIC_SUPABASE_URL'] + '/rest/v1/cuentas?select=empresa,cid,consecutivo,asesor,estado',
    headers={'apikey': K, 'Authorization': 'Bearer ' + K})
CART = {norm(c['empresa']): c for c in json.loads(urllib.request.urlopen(rq, timeout=60).read().decode())}

VIVO = ('activo', 'en_riesgo')
filas, mov = [], collections.Counter()
for r in crudas:
    nom = str(g(r, 'cliente') or '').strip()
    if not nom:
        continue
    m = str(g(r, 'movimiento') or '').strip()
    mov[m or '(vacío)'] += 1
    k = norm(nom)
    c = CART.get(k)
    ini, fin = num(g(r, 'mrrIni')), num(g(r, 'mrrFin'))
    real = num(g(r, 'perdidaReal'))

    # PROVISIONAL: marcada como baja, pero la cuenta sigue viva en la base y la
    # «perdida» es exactamente su MRR inicio con fin en cero — la firma de un
    # contrato que aun no se factura, no la de una baja.
    provisional = bool(
        m == 'Churn confirmado' and fin == 0 and abs(real - ini) < 0.01
        and c and c['estado'] in VIVO
    )
    filas.append({
        'cliente': nom,
        'clasif': str(g(r, 'clasif') or '').strip() or None,
        'facturas': int(num(g(r, 'facturas'))),
        'meses': int(num(g(r, 'meses'))),
        'acumulado': num(g(r, 'acumulado')),
        'mrrIni': ini, 'mrrFin': fin,
        'ganado': num(g(r, 'ganado')),
        'movimiento': m or None,
        'perdidaReal': real,
        'perdidaFraude': num(g(r, 'perdidaFraude')),
        'rango': str(g(r, 'rango') or '').strip() or None,
        'consecutivo': c['consecutivo'] if c else None,
        'asesor': c['asesor'] if c else None,
        'cid': c['cid'] if c else None,
        'estadoBase': c['estado'] if c else None,
        'enCartera': bool(c),
        'provisional': provisional,
    })

ini = sum(f['mrrIni'] for f in filas)
fin = sum(f['mrrFin'] for f in filas)
real = sum(f['perdidaReal'] for f in filas)
fra = sum(f['perdidaFraude'] for f in filas)
gan = sum(f['ganado'] for f in filas)
prov = [f for f in filas if f['provisional']]

# La aritmetica del mes tiene que cerrar o el archivo no sirve.
desc = (ini - real - fra + gan) - fin
assert abs(desc) < 1.0, 'el mes no cierra: descuadre de %.2f' % desc

META = {
    'origen': os.path.basename(ORIGEN),
    'mes': str(crudas[0][ix['mes']]).strip() if ix['mes'] is not None and crudas else None,
    'filas': len(filas),
    'clientes': len({norm(f['cliente']) for f in filas}),
    'enCartera': sum(1 for f in filas if f['enCartera']),
    'mrrInicio': round(ini, 2), 'mrrFin': round(fin, 2),
    'perdidaReal': round(real, 2), 'perdidaFraude': round(fra, 2), 'ganado': round(gan, 2),
    'descuadre': round(desc, 2),
    'movimientos': dict(mov.most_common()),
    'provisionales': len(prov),
    'provisionalMonto': round(sum(f['perdidaReal'] for f in prov), 2),
    'advertencia': (
        'El corte se tomó con el mes en curso. Las filas marcadas `provisional` '
        'vienen como «Churn confirmado» pero su cuenta sigue activa o en riesgo '
        'en la base: MRR Fin en cero y pérdida exactamente igual al MRR inicio '
        'es la firma de un contrato que todavía no se factura, no la de una baja.'
    ),
}

os.makedirs(os.path.dirname(SALIDA), exist_ok=True)
with io.open(SALIDA, 'w', encoding='utf-8') as f:
    json.dump({'meta': META, 'filas': filas}, f, ensure_ascii=False, separators=(',', ':'))

print('=== %s ===' % META['mes'])
print('  filas %s · clientes %s · de la cartera %d'
      % (format(META['filas'], ','), format(META['clientes'], ','), META['enCartera']))
print()
print('  MRR inicio      $%14s' % format(round(ini), ','))
print('  − pérdida real  $%14s' % format(round(real), ','))
print('  − fraude/reestr $%14s' % format(round(fra), ','))
print('  + ganado        $%14s' % format(round(gan), ','))
print('  = MRR fin       $%14s   (declarado $%s · descuadre %.2f)'
      % (format(round(ini - real - fra + gan), ','), format(round(fin), ','), desc))
print()
print('  acumulado recurrente total: $%s' % format(round(sum(f['acumulado'] for f in filas)), ','))
print()
print('=== MOVIMIENTOS ===')
for m, k in mov.most_common():
    print('  %-44s %s' % (m[:44], format(k, ',')))
print()
print('=== LO PROVISIONAL (marcado como baja, pero la cuenta sigue viva) ===')
print('  filas: %d · monto que declararían perdido: $%s'
      % (len(prov), format(round(META['provisionalMonto']), ',')))
for f in sorted(prov, key=lambda x: -x['perdidaReal'])[:8]:
    print('    %-6s %-30s %3d meses · %-10s $%s'
          % (f['consecutivo'] or '—', f['cliente'][:30], f['meses'], f['estadoBase'],
             format(round(f['perdidaReal']), ',')))
print()
print('escrito %s (%.0f KB)' % (SALIDA, os.path.getsize(SALIDA) / 1024))
