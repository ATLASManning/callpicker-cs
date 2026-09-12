"""Concilia los archivos de llamadas con la cartera por CID Y por nombre.

   Instruccion de direccion (11 sep 2026): «conciliar de acuerdo a su CID o
   nombre de cliente y ahi detonar el modulo».

   Tres preguntas:
     1. De los CIDs del archivo que SI cruzan por CID, ¿el nombre coincide?
        Un nombre que no coincide es un CID mal capturado disfrazado de exito.
     2. ¿Hay cuentas de cartera sin CID (o con CID que no aparece) cuyo NOMBRE
        si este en el archivo? Esas son las que hoy se quedarian sin modulo.
     3. ¿Algun nombre del archivo cruza con MAS de una cuenta? Ahi no se detona
        nada: se reporta y lo decide una persona.

   El emparejamiento por nombre es conservador a proposito: normalizacion
   (minusculas, sin acentos, sin razon social) y coincidencia EXACTA. Lo que
   no case exacto se reporta como candidato, nunca se enlaza solo.
"""
import sys, io, json, re, unicodedata, collections
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', line_buffering=True)
import openpyxl

ENT = r"C:\Users\manni\OneDrive\Escritorio\Llamadas_entrantes_Clientes_AAA_Poco_Consumo_Actualizado.xlsx"
cu = json.load(io.open(r'D:\Proyectos\CP\cartera_asesor.json', encoding='utf-8'))

# Sufijos de razon social: se quitan para comparar, no para mostrar.
RAZON = r'\b(s\.?a\.?p\.?i\.?|s\.?a\.?|s\.?\s?de\s?r\.?l\.?|c\.?v\.?|de\s?c\.?v\.?|sc|sofom|e\.?n\.?r\.?|spr|rl)\b'


def norma(s):
    s = unicodedata.normalize('NFD', str(s or ''))
    s = ''.join(c for c in s if unicodedata.category(c) != 'Mn').lower()
    s = s.replace('&', ' y ')
    s = re.sub(r'[.,()\-_/]', ' ', s)
    s = re.sub(RAZON, ' ', s)
    s = re.sub(r'[^a-z0-9 ]', ' ', s)
    return re.sub(r'\s+', ' ', s).strip()


# ── nombres del archivo por CID ────────────────────────────────────────────
wb = openpyxl.load_workbook(ENT, data_only=True, read_only=True)
ws = wb[wb.sheetnames[0]]
it = ws.iter_rows(values_only=True)
cab = [str(c).strip() for c in next(it)]
ix = {c: k for k, c in enumerate(cab)}
nombresArchivo = collections.defaultdict(collections.Counter)
for r in it:
    cid = str(r[ix['customer_id']] or '').strip()
    e = str(r[ix['empresa']] or '').strip()
    if cid and e and e.upper() != 'NULL':
        nombresArchivo[cid][e] += 1
wb.close()
ARCH = {cid: c.most_common(1)[0][0] for cid, c in nombresArchivo.items()}
print('CIDs con nombre en el archivo: %d' % len(ARCH))
multi = {c: n for c, n in nombresArchivo.items() if len(n) > 1}
print('CIDs con MAS de un nombre en el archivo: %d %s' % (len(multi), list(multi)[:4]))

porCid = {str(c.get('cid') or '').strip(): c for c in cu if str(c.get('cid') or '').strip()}
porNom = collections.defaultdict(list)
for c in cu:
    porNom[norma(c['empresa'])].append(c)

print()
print('=' * 86)
print('1 · LOS QUE CRUZAN POR CID: ¿coincide el nombre?')
print('=' * 86)
igual, distinto = 0, []
for cid, nomArch in ARCH.items():
    c = porCid.get(cid)
    if not c:
        continue
    a, b = norma(nomArch), norma(c['empresa'])
    if a == b or a in b or b in a:
        igual += 1
    else:
        distinto.append((c['consecutivo'], cid, c['empresa'], nomArch, c['asesor']))
print('  nombre coincide o contiene : %d' % igual)
print('  nombre DISTINTO            : %d' % len(distinto))
for cons, cid, emp, arch, ase in sorted(distinto):
    print('    %-5s CID %-8s cartera «%s»' % (cons, cid, emp[:36]))
    print('          %14s archivo «%s»   [%s]' % ('', arch[:36], ase))

print()
print('=' * 86)
print('2 · CUENTAS DE CARTERA QUE HOY SE QUEDAN SIN MODULO')
print('=' * 86)
sinCruce = [c for c in cu if str(c.get('cid') or '').strip() not in ARCH]
print('  cuentas con asesor que NO cruzan por CID: %d' % len(sinCruce))
rescate = []
for c in sinCruce:
    n = norma(c['empresa'])
    hits = [(cid, a) for cid, a in ARCH.items() if norma(a) == n]
    if hits:
        rescate.append((c, hits))
print('  de esas, su NOMBRE si esta en el archivo: %d' % len(rescate))
for c, hits in rescate:
    print('    %-5s «%s» [%s] cid en cartera=%r -> archivo CID %s «%s»'
          % (c['consecutivo'], c['empresa'][:32], c['asesor'], c.get('cid'), hits[0][0], hits[0][1][:32]))
if not rescate:
    print('    (ninguna: el nombre no rescata cuentas nuevas)')

print()
print('  --- candidatos parciales (NO se enlazan solos, se reportan) ---')
parc = []
for c in sinCruce:
    n = norma(c['empresa'])
    if len(n) < 6:
        continue
    for cid, a in ARCH.items():
        na = norma(a)
        if len(na) < 6 or cid in porCid:
            continue
        if n.split()[0] == na.split()[0] and (n in na or na in n or n.split()[0] == na.split()[0]):
            parc.append((c['consecutivo'], c['empresa'], cid, a))
vistos = set()
for cons, emp, cid, a in parc:
    if (cons, cid) in vistos:
        continue
    vistos.add((cons, cid))
    print('    %-5s «%s»  ~?~  CID %s «%s»' % (cons, emp[:32], cid, a[:32]))
if not parc:
    print('    (ninguno)')

print()
print('=' * 86)
print('3 · ¿UN NOMBRE DEL ARCHIVO CRUZA CON MAS DE UNA CUENTA?')
print('=' * 86)
amb = []
for cid, a in ARCH.items():
    hits = porNom.get(norma(a), [])
    if len(hits) > 1:
        amb.append((cid, a, [h['consecutivo'] for h in hits]))
print('  nombres ambiguos: %d' % len(amb))
for cid, a, hits in amb:
    print('    CID %-8s «%s» -> %s' % (cid, a[:34], hits))
if not amb:
    print('    (ninguno)')

print()
print('=' * 86)
print('RESUMEN DE LA REGLA DE DETONACION')
print('=' * 86)
porCidOk = sum(1 for cid in ARCH if cid in porCid)
print('  detonan por CID                  : %d' % porCidOk)
print('  detonan por NOMBRE exacto (extra): %d' % len(rescate))
print('  TOTAL de fichas con modulo       : %d de %d cuentas con asesor' % (porCidOk + len(rescate), len(cu)))
print('  fichas con la tarjeta SIN LECTURA: %d' % (len(cu) - porCidOk - len(rescate)))
